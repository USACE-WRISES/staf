"""The state SQT curve registry (data/sqt/registry.json).

Built deterministically from the metric library's raw bins; every defect the audit named is
detected against the data; the two partial originals verify the records they cover; and the
loader searches, checks applicability and freezes records for projects.
"""
from __future__ import annotations

import csv
import importlib.util
import json
from pathlib import Path

import pytest

from streamcurves import sqt_registry as sr

APP_DIR = Path(__file__).resolve().parents[1]
REPO = APP_DIR.parents[1]
SCRIPT = APP_DIR / "scripts" / "build_sqt_registry.py"
TEMPLATES = "apps/stream-curves/data/templates/"
MN_LIST = TEMPLATES + "MN-List-of-Metricsv2.0.xlsx"
WI_CURVES = TEMPLATES + "WISQT_Reference_Curves.xlsx"


def _load_builder():
    spec = importlib.util.spec_from_file_location("build_sqt_registry", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def builder():
    return _load_builder()


@pytest.fixture(scope="module")
def no_originals(tmp_path_factory):
    return tmp_path_factory.mktemp("no-owner-originals")


@pytest.fixture(scope="module")
def doc(builder, no_originals):
    """A fresh build that sees only the two partial originals shipped in data/templates."""
    return builder.build(originals=no_originals)


@pytest.fixture(scope="module")
def recs(doc):
    return {r["key"]: r for r in doc["records"]}


def codes(rec, severity=None):
    return {i["code"] for i in rec["issues"] if severity is None or i["severity"] == severity}


# --------------------------------------------------------------------------- #
# determinism and the committed file
# --------------------------------------------------------------------------- #
def test_build_is_deterministic(builder, doc, no_originals):
    again = builder.build(originals=no_originals)
    assert builder.render(doc) == builder.render(again)
    text = builder.render(doc)
    assert "\r" not in text and text.endswith("\n")
    assert chr(0x2014) not in text          # no em dash in any user-visible string
    keys = [r["key"] for r in doc["records"]]
    assert keys == sorted(keys) and len(keys) == len(set(keys))


def _owner_originals_missing(committed: dict, builder) -> list[str]:
    root = builder.originals_root()
    missing = []
    for o in (committed.get("inputs") or {}).get("originals") or []:
        if o.get("origin") != "owner":
            continue
        rel = str(o["file"]).split("/", 1)[1]
        if not (root / rel).is_file():
            missing.append(o["file"])
    return missing


def test_check_passes_on_committed_registry(builder, capsys):
    committed = json.loads(sr.REGISTRY_PATH.read_text(encoding="utf-8"))
    missing = _owner_originals_missing(committed, builder)
    if missing:
        pytest.skip(f"the committed registry used owner originals not on this machine: {missing}")
    assert builder.main(["--check"]) == 0, capsys.readouterr().out
    raw = sr.REGISTRY_PATH.read_bytes()
    assert b"\r" not in raw


def test_check_fails_on_a_stale_file(builder, tmp_path, capsys):
    stale = tmp_path / "registry.json"
    doc = json.loads(sr.REGISTRY_PATH.read_text(encoding="utf-8"))
    doc["records"][0]["eligible"] = not doc["records"][0]["eligible"]
    stale.write_text(json.dumps(doc), encoding="utf-8")
    assert builder.main(["--check", "--out", str(stale)]) == 1
    assert doc["records"][0]["key"] in capsys.readouterr().out


# --------------------------------------------------------------------------- #
# record shape and fingerprints
# --------------------------------------------------------------------------- #
REQUIRED = {"key", "state", "stateName", "tool", "edition", "editionBasis", "citation",
            "originalMetricName", "stafMetricId", "function", "protocol", "units",
            "stratum", "stratumName", "streamTypes", "geography", "form", "direction",
            "source", "originalValues", "thresholds", "normalizedPoints", "extrapolation",
            "scoreScale", "sourceLimits", "fingerprints", "adaptedIn", "verification",
            "issues", "eligible"}


def test_record_shape(doc):
    assert doc["schemaVersion"] == sr.SCHEMA_VERSION
    assert len(doc["records"]) == doc["summary"]["records"] > 0
    for r in doc["records"]:
        assert REQUIRED <= set(r), r["key"]
        state, metric, stratum = r["key"].split(":")[1:]
        assert r["key"].startswith("sqt:") and state == r["state"].lower() and metric and stratum
        assert len(r["state"]) == 2 and r["tool"].endswith("Stream Quantification Tool")
        assert r["form"] in sr.FORMS and r["direction"] in sr.DIRECTIONS
        assert r["extrapolation"] in sr.EXTRAPOLATIONS
        assert r["scoreScale"] == sr.SQT_SCORE_SCALE
        assert r["verification"]["status"] in sr.VERIFICATION_STATUSES
        assert len(r["originalValues"]) == 6
        for b in r["originalValues"]:
            assert isinstance(b["field"], str) and isinstance(b["index"], str)
        xs = [p["x"] for p in r["normalizedPoints"]]
        assert xs == sorted(xs)
        assert all(0.0 <= p["y"] <= 1.0 for p in r["normalizedPoints"])
        assert all(i["severity"] in sr.SEVERITIES and i["detail"] for i in r["issues"])
        assert r["eligible"] == (not codes(r, "defect"))
        assert (r["verification"]["status"] == "defective") == (not r["eligible"])
        if r["verification"]["status"] == "unverified":
            assert r["verification"]["against"] is None
            assert "not verified against the original" in r["verification"]["reasons"][0]
        if r["units"] is None:
            assert "units-not-stated" in codes(r)


def test_fingerprints_are_reproducible(doc, builder):
    lib = builder.MetricLibrary(REPO / doc["inputs"]["metricLibraryCsv"]["file"])
    with open(REPO / doc["inputs"]["metricLibraryCsv"]["file"], encoding="utf-8-sig",
              newline="") as f:
        rows = list(csv.reader(f))
    header = rows[0]
    for r in doc["records"]:
        row = rows[r["source"]["csvRow"] - 1]
        start = header.index(f"Stratification {r['source']['stratification']}")
        cells = row[start:start + 22]
        ident = {k: row[header.index(k)] for k in builder.ROW_FIELDS}
        assert sr.fingerprint({"row": ident, "stratification": cells}) == \
            r["fingerprints"]["sourceRow"], r["key"]
        assert sr.points_fingerprint(r["normalizedPoints"]) == r["fingerprints"]["points"]
        assert cells[0] == r["stratum"]
        assert [b["field"] for b in r["originalValues"]] == cells[4:22:3]
        assert [b["index"] for b in r["originalValues"]] == cells[6:22:3]
    assert lib.groups[-1] == 28


# --------------------------------------------------------------------------- #
# the audit's defects, confirmed against the data
# --------------------------------------------------------------------------- #
def test_placeholder_curve_nc_sinuosity(recs):
    for key in ("sqt:nc:sinuosity:e5-streams-in-unconfined-alluvial-valleys",
                "sqt:nc:sinuosity:unconfined-alluvial-valleys"):
        r = recs[key]
        assert "insufficient-points" in codes(r, "defect") and not r["eligible"]
        assert len(r["normalizedPoints"]) == 1
        adapted = r["adaptedIn"][0]
        assert "placeholder" in adapted["differences"]
        assert [(p["x"], p["y"]) for p in adapted["points"]] == [(0, 0), (1, 0.4), (2, 0.7),
                                                                 (3, 1)]
        assert "adapted-placeholder" in codes(r)


def test_swapped_columns_and_clamp_sc_ecoli(recs):
    r = recs["sqt:sc:e-coli:default"]
    assert "swapped-columns" in codes(r, "defect") and not r["eligible"]
    assert [b["field"] for b in r["originalValues"]] == ["0", "0.29", "0.3", "0.69", "0.7", "1"]
    assert [(p["x"], p["y"]) for p in r["normalizedPoints"]] == [(35, 0.7), (121, 0.69),
                                                                 (349, 0.29)]
    adapted = r["adaptedIn"][0]
    assert {p["y"] for p in adapted["points"]} == {1.0}
    assert "clamped" in adapted["differences"] and "adapted-clamped" in codes(r)


def test_mislabelled_stratum_mn_dissolved_oxygen(recs):
    r = recs["sqt:mn:do:2b-2bd"]
    assert "mislabelled-stratum" in codes(r, "defect") and not r["eligible"]
    v = r["verification"]
    assert v["status"] == "defective"
    assert v["against"]["file"] == MN_LIST and v["against"]["cells"] == "F31:K31"
    detail = next(i["detail"] for i in r["issues"] if i["code"] == "mislabelled-stratum")
    assert "F32:K32" in detail and "Use Class: 7" in detail
    assert recs["sqt:mn:do:2a"]["verification"]["status"] == "verified"


def test_two_sided_curves_cut_to_one_limb(recs):
    confirmed = ["sqt:mn:pool-spacing-ratio:c-and-e-streams",
                 "sqt:mn:percent-riffle:a-and-b-streams",
                 "sqt:mn:percent-riffle:c-and-e-stream-types",
                 "sqt:wi:pool-spacing-ratio:c-and-e-stream-types",
                 "sqt:wi:percent-riffle:a-and-b-stream-types",
                 "sqt:wi:percent-riffle:c-and-e-stream-types",
                 "sqt:wi:width-depth-ratio-state:default"]
    for key in confirmed:
        r = recs[key]
        assert "two-sided-one-limb" in codes(r, "defect"), key
        assert r["verification"]["status"] == "defective"
        pts = r["verification"]["originalPoints"]
        signs = [b["y"] - a["y"] for a, b in zip(pts, pts[1:]) if b["y"] != a["y"]]
        assert any(s > 0 for s in signs) and any(s < 0 for s in signs), key
    for key in ("sqt:wy:width-depth-ratio-state:default",
                "sqt:sc:width-depth-ratio-state:default",
                "sqt:co:pool-spacing-ratio:c-stream-types",
                "sqt:nc:percent-riffle:streams-lt-3-pct-slope"):
        assert "two-sided-one-limb" in codes(recs[key], "defect"), key
    # every riffle record is one limb; B-type pool spacing is one-sided in the originals
    for r in recs.values():
        if r["stafMetricId"] == "habitat-provision-percent-riffle":
            assert "two-sided-one-limb" in codes(r)
    for key in ("sqt:mn:pool-spacing-ratio:a-and-b-stream-types",
                "sqt:wi:pool-spacing-ratio:bc-stream-types",
                "sqt:ak:pool-spacing-ratio:b-and-ba-stream-types"):
        assert "two-sided-one-limb" not in codes(recs[key]), key
    # the NC row that carries both limbs is two-sided and eligible; its plateau is kept
    nc = recs["sqt:nc:pool-spacing-ratio:slope-lt-4-pct-and-da-ge-10-and-c-or-e-stream-types"]
    assert nc["form"] == "two-sided" and nc["eligible"]
    assert {"x": 7.0, "y": 1.0} in nc["normalizedPoints"]
    assert "adapted-drops-breakpoint" in codes(nc)


def test_layers_that_never_reach_0_or_1(doc, recs, builder):
    adapted = builder.AdaptedBundles()
    never0 = sum(1 for L in adapted.all_layers if min(p["y"] for p in L["points"]) > 0)
    never1 = sum(1 for L in adapted.all_layers if max(p["y"] for p in L["points"]) < 1)
    s = doc["summary"]["adapted"]
    assert (s["layersNeverReach0"], s["layersNeverReach1"]) == (never0, never1)
    assert never0 > 0 and never1 > 0
    for r in recs.values():
        pts = r["normalizedPoints"]
        by_threshold = {t["index"] for t in r["thresholds"]}
        if len(pts) >= 2 and min(p["y"] for p in pts) > 0 and 0.0 not in by_threshold:
            assert "does-not-reach-0" in codes(r), r["key"]
        if len(pts) >= 2 and max(p["y"] for p in pts) < 1 and 1.0 not in by_threshold:
            assert "does-not-reach-1" in codes(r), r["key"]
        for a in r["adaptedIn"]:
            ends = (a["points"][0]["y"], a["points"][-1]["y"])
            if any(0 < y < 1 for y in ends):
                assert "adapted-held-flat" in codes(r), r["key"]
    mwat = recs["sqt:co:mwat:cs-i"]
    assert mwat["extrapolation"] == "unknown" and mwat["eligible"]


def test_rounding_drift_against_the_mn_original(recs):
    fish = recs["sqt:mn:fish-ibi:northern-rivers-northern"]
    assert "rounded-values" in codes(fish)
    assert fish["verification"]["status"] == "partially-verified"
    rows = {row["cell"]: row["result"] for row in fish["verification"]["comparison"]}
    assert rows["F46"] == "rounded"
    tss = recs["sqt:mn:tss:2b-2bd-south"]
    assert "value-drift" in codes(tss)
    macro = recs["sqt:mn:macroinvertebrate-ibi:northern-forest-rivers-northern"]
    assert "value-drift" in codes(macro)


def test_raw_and_adapted_differ_only_by_compile_rules(doc, recs):
    for r in recs.values():
        for a in r["adaptedIn"]:
            assert "differs-from-compile" not in a["differences"], r["key"]
            if a["pointsEqualRaw"]:
                assert a["points"] == r["normalizedPoints"]
            else:
                assert set(a["differences"]) & {"placeholder", "clamped", "read-swapped-columns",
                                                 "dropped-extra-breakpoint",
                                                 "threshold-read-as-point"}, r["key"]
    assert doc["summary"]["adapted"]["recordsCarried"] == len(recs)


def test_units_blank_everywhere_in_the_csv(recs):
    assert all(r["source"]["unit"] == "" for r in recs.values())
    for r in recs.values():
        if r["state"] == "MN":
            assert r["units"] and MN_LIST in r["unitsBasis"]
        else:
            assert r["units"] is None and "units-not-stated" in codes(r)
    assert recs["sqt:mn:do:2a"]["units"] == "mg/L"


def test_categorical_metrics_dropped_are_findings(doc):
    found = [(f["code"], f.get("metric"), f.get("file")) for f in doc["findings"]]
    assert ("categorical-not-carried", "Dominant BEHI/NBS", MN_LIST) in found
    assert ("categorical-not-carried", "Dominant BEHI/NBS", WI_CURVES) in found
    screening = [f for f in doc["findings"] if f["code"] == "screening-categorical-not-carried"]
    assert any(f["metric"] == "Aufeis / Icing Impacts" and f["states"] == ["AK"]
               for f in screening)
    class7 = next(f for f in doc["findings"] if f.get("cells") == "F32:K32")
    assert "sqt:mn:do:2b-2bd" in class7["detail"]


def test_alaska_interior_curves_labelled_statewide(recs):
    bundle = json.loads((REPO / "apps/library/assessments/ak-sqt-adapted/v1/"
                                "assessment.deep.json").read_text(encoding="utf-8"))
    assert bundle["applicability"] == "Alaska streams"
    mmi = recs["sqt:ak:alaskan-interior-mmi:default"]
    assert mmi["geography"]["region"] == "Interior Alaska"
    assert "interior-alaska-labelled-statewide" in codes(mmi)
    interior = [r for r in recs.values() if r["geography"]["region"] == "Interior Alaska"]
    assert len(interior) == 6 and all(r["state"] == "AK" for r in interior)
    for r in recs.values():
        if r["state"] == "AK" and r["geography"]["region"] is None:
            assert "alaska-scope-unconfirmed" in codes(r)
    assert all(r["edition"] is None and "edition-unknown" in codes(r)
               for r in recs.values() if r["state"] == "AK")


# --------------------------------------------------------------------------- #
# the two partial originals
# --------------------------------------------------------------------------- #
def test_partial_originals_verify_their_records(recs):
    mn = [r for r in recs.values() if r["state"] == "MN"]
    wi = [r for r in recs.values() if r["state"] == "WI"]
    assert mn and wi
    for r in mn:
        assert r["verification"]["against"]["file"] == MN_LIST, r["key"]
        assert r["verification"]["against"]["sheet"] == "Performance Standards"
    for r in wi:
        assert r["verification"]["against"]["file"] == WI_CURVES, r["key"]
        assert r["verification"]["against"]["sheet"] == "Reference_Curves"
    for r in recs.values():
        if r["state"] not in ("MN", "WI"):
            assert r["verification"]["against"] is None
            assert r["verification"]["status"] in ("unverified", "defective")
    for key, cells in (("sqt:mn:pool-depth-ratio:default", "F19:K19"),
                       ("sqt:mn:aggradation-ratio:default", "F22:K22"),
                       ("sqt:wi:pool-depth-ratio:default", "U401:Z401"),
                       ("sqt:wi:lwd-index:default", "U10:Z10")):
        v = recs[key]["verification"]
        assert v["status"] == "verified" and v["against"]["cells"] == cells, key
        assert all(row["result"] in ("exact", "consistent", "on-record-line")
                   for row in v["comparison"])
    # values the record lacks but that lie on its line still verify
    armor = recs["sqt:mn:percent-armoring:default"]["verification"]
    assert armor["status"] == "verified"
    assert {row["result"] for row in armor["comparison"]} == {"exact", "on-record-line"}
    # a missing breakpoint off the line is only a partial match
    ps = recs["sqt:mn:pool-spacing-ratio:a-and-b-stream-types"]
    assert ps["verification"]["status"] == "partially-verified"
    assert "missing-breakpoints" in codes(ps)
    # the WI DPI is logarithmic between its points
    dpi = recs["sqt:wi:diatom-phosphorus-index-dpi:default"]
    assert "original-log-form" in codes(dpi) and dpi["extrapolation"] == "unknown"
    # linear extension is read from the originals, never assumed
    assert recs["sqt:mn:bank-height-ratio-bhr:default"]["extrapolation"] == "linear"
    assert recs["sqt:wi:summer-mean-temperature:cold"]["extrapolation"] == "linear"
    assert recs["sqt:mn:do:2a"]["sourceLimits"] == ["Not applicable to ephemeral streams."]


def test_original_cells_out_of_order_are_not_compared(recs):
    riffle = recs["sqt:mn:percent-riffle:c-and-e-stream-types"]["verification"]
    results = {row["cell"]: row["result"] for row in riffle["comparison"]}
    assert results["H21"] == results["J21"] == "original-out-of-order"
    assert results["F21"] == results["K21"] == "exact"


def test_compare_values(builder):
    assert builder.compare_values("1.6", "1.60") == "exact"
    assert builder.compare_values("5.25", "5.3") == "consistent"
    assert builder.compare_values("29", "29.1") == "rounded"
    assert builder.compare_values("40", "40.5") == "drift"
    assert builder.compare_values("0.75", "0.80") == "mismatch"


# --------------------------------------------------------------------------- #
# owner originals under the originals root
# --------------------------------------------------------------------------- #
def _owner_workbook(path: Path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reference Curves"
    ws["B2"] = "Entrenchment Ratio (ER) C Stream Types"
    ws["B3"] = "Field Value"
    ws["E3"], ws["G3"], ws["H3"] = 2, 2.4, 4.2
    ws["B4"] = "Index Value"
    for col, lvl in zip("CDEFGH", (0, 0.29, 0.3, 0.69, 0.7, 1)):
        ws[f"{col}4"] = lvl
    ws["B6"] = "Coefficients - Y = a * X + b"
    ws["C7"] = "=SLOPE(C4:G4,C3:G3)"
    wb.save(path)


def test_owner_originals_are_read_from_sources_json(builder, tmp_path):
    co = tmp_path / "CO"
    co.mkdir()
    _owner_workbook(co / "co-curves.xlsx")
    (co / "manual.pdf").write_bytes(b"%PDF-1.4 test")
    (co / "sources.json").write_text(json.dumps({
        "state": "CO",
        "sources": [
            {"id": "co-test-curves", "file": "co-curves.xlsx", "edition": "test",
             "citation": "test workbook", "layout": "reference-curves",
             "sheet": "Reference Curves"},
            {"id": "co-test-manual", "file": "manual.pdf", "edition": "test",
             "citation": "test manual", "layout": "manual",
             "tables": [{"metric": "Pool Depth Ratio", "stratum": "Default", "page": 7,
                         "table": "Table 2", "values": [">= 1", "", "", "", "2.2", "3.2"],
                         "units": "ft/ft"}]},
        ]}), encoding="utf-8")
    wy = tmp_path / "WY"
    wy.mkdir()
    (wy / "unlisted.xlsx").write_bytes(b"not listed")
    doc = builder.build(originals=tmp_path)
    recs = {r["key"]: r for r in doc["records"]}
    er = recs["sqt:co:entrenchment-ratio-er:c-streams"]
    assert er["verification"]["status"] == "verified"
    assert er["verification"]["against"]["file"] == "sqt-originals/CO/co-curves.xlsx"
    assert er["extrapolation"] == "linear"
    pdr = recs["sqt:co:pool-depth-ratio:default"]
    assert pdr["verification"]["status"] == "verified"
    assert pdr["verification"]["against"]["page"] == 7 and pdr["units"] == "ft/ft"
    assert any(o["file"] == "sqt-originals/CO/co-curves.xlsx" and o["origin"] == "owner"
               for o in doc["inputs"]["originals"])
    assert any(f["code"] == "original-without-sources-json" and f["state"] == "WY"
               for f in doc["findings"])
    assert all("sqt-originals" not in json.dumps(r) for r in doc["records"]
               if r["state"] not in ("CO",))


# --------------------------------------------------------------------------- #
# loader: search, applicability, frozen copies
# --------------------------------------------------------------------------- #
def test_loader_search_and_filters(tmp_path):
    sr.clear_cache()
    assert sr.available() and not sr.available(tmp_path / "absent.json")
    everything = sr.records()
    assert len(everything) == sr.load()["summary"]["records"]
    mn = sr.records(state="MN")
    assert mn and all(r["state"] == "MN" for r in mn)
    assert len(sr.records(state="Minnesota")) == len(mn)
    hp = sr.records(function="habitat-provision")
    assert hp and all(r["function"]["id"] == "habitat-provision" for r in hp)
    assert len(sr.records(function="Habitat provision")) == len(hp)
    riffle = sr.records(metric="Percent Riffle")
    assert {r["state"] for r in riffle} == {"AK", "CO", "MI", "MN", "NC", "SC", "WI", "WY"}
    assert len(sr.records(metric="habitat-provision-percent-riffle")) == len(riffle)
    assert {r["state"] for r in sr.records(edition="2020")} == {"MN"}
    both = sr.records(state="WI", metric="Pool Spacing Ratio")
    assert len(both) == 3
    bad = sr.records(eligible=False)
    assert bad and all(not r["eligible"] for r in bad)
    assert len(bad) + len(sr.records(eligible=True)) == len(everything)
    ctx = {"states": ["MN", "WI"], "stafMetricId": "habitat-provision-pool-depth-ratio"}
    fit = sr.records(applicable_to=ctx)
    assert fit and all(r["state"] in ("MN", "WI") and r["eligible"] for r in fit)
    assert all(r["stafMetricId"] == "habitat-provision-pool-depth-ratio" for r in fit)


def test_record_returns_copies():
    sr.clear_cache()
    key = "sqt:wi:pool-depth-ratio:default"
    r = sr.record(key)
    assert r["key"] == key
    r["normalizedPoints"].clear()
    assert sr.record(key)["normalizedPoints"]
    assert sr.record("sqt:zz:none:none") is None


def test_applicability_checks():
    r = sr.record("sqt:mn:do:2a")
    checks = sr.applicability(r, {})
    assert [c["id"] for c in checks] == list(sr.CHECK_IDS)
    assert all(c["status"] in sr.CHECK_STATUSES and c["detail"] for c in checks)
    by = {c["id"]: c for c in checks}
    assert by["score-scale"]["status"] == "warn" and "0.39" in by["score-scale"]["detail"]
    assert by["protocol"]["status"] == "unknown"
    ctx = {"stafMetricId": "water-and-soil-quality-do", "units": "mg/L", "states": ["MN"],
           "direction": "increasing", "scoreScale": "sqt", "xRange": [6.0, 8.0]}
    by = {c["id"]: c for c in sr.applicability(r, ctx)}
    for cid in ("eligibility", "construct", "units", "direction", "score-scale", "geography",
                "stream-type", "extrapolation"):
        assert by[cid]["status"] == "pass", (cid, by[cid])
    assert by["source-limits"]["status"] == "warn"
    by = {c["id"]: c for c in sr.applicability(r, {"states": ["IA"], "units": "ppm",
                                                   "direction": "decreasing",
                                                   "stafMetricId": "other",
                                                   "flowType": "ephemeral"})}
    for cid in ("geography", "units", "direction", "construct", "source-limits"):
        assert by[cid]["status"] == "fail", cid
    er = sr.record("sqt:wi:entrenchment-ratio-er:c-stream-types")
    assert sr.applicability(er, {"streamType": "C"})[7]["status"] == "pass"
    assert sr.applicability(er, {"streamType": "Cb"})[7]["status"] == "warn"
    assert sr.applicability(er, {"streamType": "B"})[7]["status"] == "fail"
    assert sr.applicability(er, {})[7]["status"] == "unknown"
    mwat = sr.record("sqt:co:mwat:cs-i")
    assert sr.open_ends(mwat) == ["high"]
    assert {c["id"]: c for c in sr.applicability(mwat, {"xRange": [14, 16]})}[
        "extrapolation"]["status"] == "pass"
    assert {c["id"]: c for c in sr.applicability(mwat, {"xRange": [14, 25]})}[
        "extrapolation"]["status"] == "fail"
    region = sr.record("sqt:nc:ept-taxa-present:piedmont")
    assert {c["id"]: c for c in sr.applicability(region, {"states": ["NC"]})}[
        "geography"]["status"] == "warn"
    defective = sr.record("sqt:sc:e-coli:default")
    assert sr.applicability(defective, {})[0]["status"] == "fail"
    assert sr.overall(sr.applicability(defective, {})) == "fail"


def test_evaluate_and_bands():
    pts = [{"x": 1, "y": 0.0}, {"x": 2, "y": 0.7}, {"x": 3, "y": 1.0}]
    assert sr.evaluate(pts, 1.5) == pytest.approx(0.35)
    assert sr.evaluate(pts, 5, "flat") == 1.0 and sr.evaluate(pts, 0, "unknown") is None
    open_pts = [{"x": 1, "y": 0.3}, {"x": 2, "y": 0.7}]
    assert sr.evaluate(open_pts, 0, "linear") == 0.0
    assert sr.evaluate(open_pts, 0.5, "linear") == pytest.approx(0.1)
    assert sr.band(0.35) == "FAR" and sr.band(0.35, sr.STAF_SCORE_SCALE) == "NF"
    assert sr.band(0.69) == "FAR" and sr.band(0.70) == "F"


def test_frozen_copy_survives_registry_changes():
    r = sr.record("sqt:wi:pool-depth-ratio:default")
    frozen = sr.frozen_copy(r)
    assert frozen["frozen"]["registryKey"] == r["key"]
    assert frozen["frozen"]["contentFingerprint"] == sr.content_fingerprint(r)
    assert sr.frozen_intact(frozen)
    assert sr.frozen_copy(frozen)["frozen"] == frozen["frozen"]
    r["normalizedPoints"][0]["y"] = 0.5          # a later registry build changes the record
    assert frozen["normalizedPoints"][0]["y"] == 0.0
    assert sr.frozen_intact(frozen)
    edited = json.loads(json.dumps(frozen))
    edited["normalizedPoints"][0]["y"] = 0.5
    assert not sr.frozen_intact(edited)
