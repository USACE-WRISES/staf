"""The generated DEEP calculator as a file: what is in it, not what it computes
(``test_deep_calculator_parity.py`` proves the arithmetic).

Reproducible bytes, the sheets and their protection, the defined names the DEEP
filler reads, one entry cell for a metric that serves several functions, a
dropdown only where a metric has several curve sets, a formula vocabulary small
enough to audit, and no value cached in the file.
"""
from __future__ import annotations

import io
import re
import zipfile

import pytest

from streamcurves import deep_calculator as dc

openpyxl = pytest.importorskip("openpyxl")

ASC = [{"x": 0, "y": 0}, {"x": 10, "y": 1}]
ALLOWED_FUNCTIONS = {"IF", "AND", "MIN", "MAX", "INDEX", "MATCH", "IFERROR", "ISNUMBER",
                     "COUNT", "SUM", "NA"}


def _bundle() -> dict:
    embed = {"metricId": "spring-phab-xembed", "metricName": "Embeddedness",
             "xLabel": "Embeddedness (%)", "curve": {"points": ASC}, "activeStratum": "",
             "curveLayers": [{"stratum": "", "points": ASC},
                             {"stratum": "ge_2", "points": [{"x": 0, "y": 0}, {"x": 5, "y": 1}]}],
             "stratifier": {"variable": "nhd_slope", "breaks": [0.005, 0.02], "right": False,
                            "classes": [{"key": "lt_0.5", "label": "Low gradient"},
                                        {"key": "0.5_to_2", "label": "Moderate gradient"},
                                        {"key": "ge_2", "label": "Steep (2 percent and above)"}]}}
    fixed = {"metricId": "spring-pctimp2019ws", "metricName": "Impervious surface",
             "xLabel": "Impervious surface (%)", "criteriaBasis": "fixed",
             "curve": {"points": [{"x": 0, "y": 1}, {"x": 10, "y": 0.69},
                                  {"x": 25.005, "y": 0.39}, {"x": 44.5, "y": 0}]}}
    return {"assessmentId": "calc-test", "assessmentName": "Calculator test",
            "contentDigest": "sha256:abc", "sourceCitation": "Test citation",
            "library": {"version": 2, "status": "preliminary",
                        "updatedAt": "2026-09-19T12:34:56+00:00"},
            "functionCoverage": {"total": 20, "covered": 3, "excluded": 17},
            "referenceMethod": {"method": "pressure-screen", "statement": "Reference statement."},
            "metricsByFunction": [
                {"functionId": "hyporheic-connectivity", "functionName": "Hyporheic connectivity",
                 "discipline": "Hydraulics", "metrics": [embed]},
                {"functionId": "catchment-hydrology", "functionName": "Catchment hydrology",
                 "discipline": "Hydrology", "metrics": [fixed]},
                {"functionId": "bed-composition-bedform-dynamics",
                 "functionName": "Bed composition and bedform dynamics",
                 "discipline": "Geomorphology", "metrics": [embed]}],
            "insufficientReferenceSupport": [
                {"metricId": "spring-chem-ntl-diss", "metricName": "Dissolved nitrogen",
                 "functions": [{"functionId": "nutrient-cycling",
                                "functionName": "Nutrient cycling"}]}]}


@pytest.fixture(scope="module")
def data() -> bytes:
    return dc.build_calculator(_bundle())


@pytest.fixture(scope="module")
def book(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def test_the_bytes_are_reproducible(data):
    assert dc.build_calculator(_bundle()) == data
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert {i.date_time for i in z.infolist()} == {(1980, 1, 1, 0, 0, 0)}
        core = z.read("docProps/core.xml").decode("utf-8")
    # the version's own timestamp, never the build time
    assert core.count("2026-09-19T12:34:56Z") == 2
    other = _bundle()
    other["library"]["updatedAt"] = "2026-09-20T00:00:00+00:00"
    assert dc.build_calculator(other) != data


def test_the_sheets_and_their_protection(book):
    assert book.sheetnames == ["Instructions", dc.SCORE, "Metrics", "Results", "Reference",
                               "Metadata", "ChartData"]
    assert {ws.title for ws in book.worksheets if ws.sheet_state == "hidden"} \
        == {"Results", "ChartData"}
    assert all(ws.protection.sheet for ws in book.worksheets)
    assert book.calculation.fullCalcOnLoad is True
    assert book.active.title == dc.SCORE


def test_the_names_the_deep_filler_reads(book):
    names = set(book.defined_names)
    assert {"in_spring_phab_xembed", "st_spring_phab_xembed", "in_spring_pctimp2019ws",
            "idx_spring_phab_xembed", "idx_spring_pctimp2019ws", "fs_hyporheic_connectivity",
            "fs_catchment_hydrology", "fs_bed_composition_bedform_dynamics",
            "sub_index_physical", "sub_index_chemical", "sub_index_biological", "eci",
            "site_name", "site_reach", "site_coords", "site_date", "site_assessor",
            "site_notes", "curve_key", "strata_spring_phab_xembed", "meta_content_digest"} <= names
    assert {f"curve_x{i}" for i in range(1, 9)} | {f"curve_y{i}" for i in range(1, 9)} <= names
    # a single-curve metric has no curve-set entry
    assert "st_spring_pctimp2019ws" not in names
    for name in names:
        if name.startswith(("in_", "st_", "site_")):
            assert book.defined_names[name].attr_text.startswith(f"'{dc.SCORE}'!")


def test_a_shared_metric_is_entered_once_and_linked(book):
    ws = book[dc.SCORE]
    entry = book.defined_names["in_spring_phab_xembed"].attr_text.split("!")[1].replace("$", "")
    assert ws[entry].value is None and ws[entry].protection.locked is False
    links = [c for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and "$" + entry[0] + "$" + entry[1:] in c.value
             and c.column_letter == "E"]
    assert len(links) == 1 and links[0].protection.locked is True


def test_the_entry_cells_are_the_only_unlocked_cells(book):
    ws = book[dc.SCORE]
    named = {book.defined_names[n].attr_text.split("!")[1].replace("$", "")
             for n in book.defined_names if n.startswith(("in_", "st_", "site_"))}
    unlocked = {c.coordinate for row in ws.iter_rows() for c in row
                if c.protection.locked is False}
    assert named <= unlocked
    # everything else unlocked belongs to a merged entry block (site rows, notes)
    for coord in unlocked - named:
        cell = ws[coord]
        assert cell.value in (None, ""), coord


def test_the_curve_set_dropdown_lists_the_layers_in_words(book):
    ws = book[dc.SCORE]
    ref = book["Reference"]
    rng = book.defined_names["strata_spring_phab_xembed"].attr_text.split("!")[1].replace("$", "")
    labels = [c[0].value for c in ref[rng]]
    assert labels == [dc.POOLED_LABEL, "Steep (2 percent and above)"]
    cell = book.defined_names["st_spring_phab_xembed"].attr_text.split("!")[1].replace("$", "")
    assert ws[cell].value == dc.POOLED_LABEL                      # the default layer
    validations = [dv for dv in ws.data_validations.dataValidation if cell in dv.sqref]
    assert validations and validations[0].formula1 == "=strata_spring_phab_xembed"


def test_every_curve_is_a_row_of_eight_knots(book):
    ref = book["Reference"]
    rows = {r[0].value: r for r in ref.iter_rows(min_row=5) if r[0].value and "|" in str(r[0].value)}
    assert set(rows) == {"spring_phab_xembed|" + dc.POOLED_LABEL,
                         "spring_phab_xembed|Steep (2 percent and above)",
                         "spring_pctimp2019ws|"}
    fixed = rows["spring_pctimp2019ws|"]
    xs = [c.value for c in fixed[6:14]]
    ys = [c.value for c in fixed[14:22]]
    assert xs == [0, 10, 25.005, 44.5, 44.5, 44.5, 44.5, 44.5]     # padded with the last knot
    assert ys == [1, 0.69, 0.39, 0, 0, 0, 0, 0]
    assert fixed[4].value == 4 and fixed[5].value == "Fixed criteria, the same in every region"
    with pytest.raises(ValueError, match="exceeds the 8 knots"):
        many = _bundle()
        many["metricsByFunction"][1]["metrics"][0]["curve"]["points"] = [
            {"x": i, "y": i / 9} for i in range(10)]
        dc.build_calculator(many)


def test_the_formula_vocabulary_is_small_and_nothing_is_cached(data, book):
    used = set()
    for ws in book.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    used |= set(re.findall(r"([A-Z][A-Z0-9.]+)\(", c.value))
    assert used <= ALLOWED_FUNCTIONS, used - ALLOWED_FUNCTIONS
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/sheet"):
                xml = z.read(name).decode("utf-8")
                # openpyxl writes an empty <v></v> beside a formula: no cached result
                assert not re.search(r"<f>[^<]*</f><v>[^<]+</v>", xml), name
            if name.startswith("xl/charts/chart"):
                assert "dispNaAsBlank" in z.read(name).decode("utf-8")


def test_what_the_workbook_says_about_the_assessment(book):
    meta = {r[0].value: r[1].value for r in book["Metadata"].iter_rows(min_row=4) if r[0].value}
    assert meta["Assessment id"] == "calc-test" and meta["Version"] == 2
    assert meta["Content digest"] == "sha256:abc"
    assert meta["Reference method"] == "pressure-screen"
    assert meta["Generator version"] == dc.GENERATOR_VERSION
    score = " ".join(str(c.value) for row in book[dc.SCORE].iter_rows() for c in row if c.value)
    assert "Withheld for insufficient reference support (not scored)" in score
    assert "Dissolved nitrogen" in score and "Nutrient cycling" in score
    assert "This assessment covers 3 of 20 STAF functions (17 documented exclusions)." in score
    text = " ".join(str(c.value) for row in book["Instructions"].iter_rows() for c in row if c.value)
    assert "Reference condition. Reference statement." in text
    assert "—" not in score + text                           # no em dash in the copy


def test_the_shared_vocabulary():
    assert dc.metric_key("spring-phab-xembed") == "spring_phab_xembed"
    assert dc.metric_key("a.b c/d") == "a_b_c_d"
    assert dc.function_key("low-flow-baseflow-dynamics") == "low_flow_baseflow_dynamics"
    m = _bundle()["metricsByFunction"][0]["metrics"][0]
    assert dc.layer_label("", m) == dc.layer_label(None, m) == dc.POOLED_LABEL
    assert dc.layer_label("ge_2", m) == "Steep (2 percent and above)"
    assert dc.layer_label("an SQT stratum, with a comma", m) == "an SQT stratum, with a comma"
    assert dc.default_layer(m) == dc.POOLED_LABEL
    with pytest.raises(ValueError, match="scores no function"):
        dc.build_calculator({"assessmentId": "empty", "metricsByFunction": []})


# --------------------------------------------------------------------------- #
# The workbook restricts the claim too (2026-09-20)
# --------------------------------------------------------------------------- #
def test_a_partial_assessment_says_what_its_index_is_not():
    """The Results cells are ratios over the functions the assessment scores, so a
    partial one lands on the same 0 to 1 scale as a complete one. The workbook
    leaves the app, so the restriction goes with it rather than staying on screen."""
    from streamcurves import deep_calculator as dc
    partial = {"functionCoverage": {"total": 20, "covered": 16, "excluded": 4}}
    text = dc.claim_restriction_text(partial)
    assert "4 of the 20 STAF functions are not assessed" in text
    assert "not comparable with a full-framework index" in text
    assert "no condition class is claimed" in text


def test_a_complete_assessment_needs_no_restriction():
    from streamcurves import deep_calculator as dc
    assert dc.claim_restriction_text({"functionCoverage": {"total": 20, "covered": 20}}) == ""
    assert dc.claim_restriction_text({}) == ""            # an older bundle says nothing


def _latest_incomplete_bundles():
    """The latest published version of every assessment that leaves a function
    unassessed. Read from the library rather than pinned to a version, so a new
    publish cannot leave this test checking a superseded bundle. Once every latest
    version covers all twenty functions (2026-09-22), the newest version of each
    assessment that left one unassessed stands in, so the restriction stays tested."""
    import json
    from pathlib import Path
    root = Path(__file__).resolve().parents[2] / "library" / "assessments"
    latest, newest = {}, {}
    for adir in sorted(p for p in root.glob("*") if p.is_dir()):
        versions = sorted(int(v.name[1:]) for v in adir.glob("v*") if v.name[1:].isdigit())
        for v in reversed(versions):
            src = adir / f"v{v}" / "assessment.deep.json"
            if not src.exists():
                continue
            bundle = json.loads(src.read_text(encoding="utf-8"))
            fc = bundle.get("functionCoverage") or {}
            if fc.get("total") and fc.get("covered", 0) < fc["total"]:
                if v == versions[-1]:
                    latest[f"{adir.name}@v{v}"] = bundle
                newest[f"{adir.name}@v{v}"] = bundle
                break
    return latest or newest


def test_the_restriction_reaches_the_visible_sheet():
    import openpyxl
    from streamcurves import deep_calculator as dc
    bundles = _latest_incomplete_bundles()
    if not bundles:
        pytest.skip("no published assessment leaves a function unassessed")
    for key, bundle in bundles.items():
        ws = openpyxl.load_workbook(io.BytesIO(dc.build_calculator(bundle)))["DEEP Score"]
        said = [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and "not assessed" in c.value]
        assert said, f"{key}: the workbook's own sheet does not carry the restriction"



def test_the_reference_sheet_says_what_a_ladder_curve_rests_on():
    """The first 0.13 workbook read "3182 least-disturbed stations borrowed from
    None ecoregion 55" for a modeled curve and "Fixed criteria, the same in every
    region" for a regional nutrient benchmark, because the sentence rode only
    inside referenceSupport and the benchmark is marked fixed downstream."""
    from streamcurves import deep_calculator as dc
    modeled = {"basis": "modeled-reference", "criteriaBasis": "reference",
               "referenceSupport": {"status": "modeled", "nUsable": 3182,
                                    "basisStatement": "Reference curve from a modeled expectation."}}
    assert dc.support_text(modeled) == "Reference curve from a modeled expectation."
    benchmark = {"basis": "published-benchmark", "criteriaBasis": "fixed",
                 "referenceSupport": {"status": "published",
                                      "basisStatement": "Scored against a published criterion."}}
    assert dc.support_text(benchmark) == "Scored against a published criterion."
    easi = {"basis": "published-benchmark", "criteriaBasis": "fixed"}
    assert dc.support_text(easi) == "Fixed criteria, the same in every region"
    bare = {"basis": "modeled-reference", "basisLabel": "Modeled reference",
            "referenceSupport": {"status": "modeled", "nUsable": 3182}}
    assert "borrowed" not in dc.support_text(bare)
