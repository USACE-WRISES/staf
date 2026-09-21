"""The Excel calculator of one detailed assessment, generated at publish.

A detailed assessment differs by region and by version, so its calculator is
built from the published bundle itself: the bundle's metrics, grouped by
function, with every reference curve programmed in. It is generated where a
version is minted (``library.publish_version``) and recorded beside it, DEEP
ships a copy, and DEEP fills a copy in at run time without any spreadsheet
library of its own.

The workbook follows the EASI calculator's conventions (one entry sheet modelled
on the SFARI worksheet, orange entry cells, helper sheets that carry the
formulas, protected sheets, reproducible bytes) and scores exactly like DEEP:

* a value is read on its curve by piecewise-linear interpolation in DEEP's own
  order (left clamp, right clamp, the first segment whose upper knot the value
  does not exceed, a vertical step read at its later point), clamped to 0 to 1;
* a function score is the mean of its metrics' indices times 15, and its class
  is read from the unrounded score (5 and 10);
* an outcome sub-index is the weighted sum of function scores over the weighted
  maximum of the functions that have a score, zero when none has one;
* the Ecosystem Condition Index is (Physical + Chemical + Biological) / 3.

A blank value is Not Applicable and drops out. A metric that serves several
functions is entered once and linked. A metric with several curve sets gets a
dropdown, and the choice selects the curve.

``tests/test_deep_calculator_parity.py`` evaluates the workbook's formulas and
proves them equal to DEEP's scoring, knot by knot.

Pure: a bundle in, bytes out. openpyxl only (already a StreamCurves dependency).
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

GENERATOR = "streamcurves.deep_calculator"
GENERATOR_VERSION = "1.0"
FILE_NAME = "calculator.xlsx"
PASSWORD = "deep"            # a guard rail against accidental edits, not a secret
SCORE = "DEEP Score"         # the worksheet the user fills
KNOTS = 8                    # every published curve has at most eight points
POOLED_LABEL = "All streams (pooled)"
FALLBACK_STAMP = _dt.datetime(2026, 1, 1, 0, 0, 0)

OUTCOMES = ("physical", "chemical", "biological")
OUTCOME_LABEL = {"physical": "Physical", "chemical": "Chemical", "biological": "Biological"}
CATEGORY_ORDER = ("Hydrology", "Hydraulics", "Geomorphology", "Physicochemistry", "Biology")
CLASS_LABELS = ("Non-Functioning", "Functioning At-Risk", "Functioning")

# palette (the EASI and SFARI calculators')
C_INPUT = "F8CBAD"
C_LINK = "F2F2F2"
C_BAND = "F2F2F2"
C_GOOD = "8FAADC"
C_FAIR = "FFFF99"
C_POOR = "FF6969"
C_CALC = "EDEDED"
CATEGORY_FILL = {"Hydrology": "D9E1F2", "Hydraulics": "B4C6E7", "Geomorphology": "FCE4D6",
                 "Physicochemistry": "E2EFDA", "Biology": "FFF2CC"}

_VENDORED_MAPPING = (Path(__file__).resolve().parent / "_vendor" / "easi" / "data"
                     / "cwa-mapping.json")


def solid(hex_: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_, end_color=hex_)


FILL_INPUT, FILL_LINK, FILL_BAND, FILL_CALC = (solid(C_INPUT), solid(C_LINK), solid(C_BAND),
                                               solid(C_CALC))
FONT_BOLD = Font(bold=True)
FONT_NOTE = Font(italic=True, color="555555", size=9)
THIN = Side(style="thin", color="000000")
MED = Side(style="medium", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNLOCKED = Protection(locked=False)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_MID = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)


# --------------------------------------------------------------------------- #
# vocabulary shared with DEEP's runtime filler (apps/deep/deep/calculator.py);
# tests/test_deep_calculator.py keeps the two equal
# --------------------------------------------------------------------------- #
def metric_key(metric_id: str) -> str:
    """The defined-name stem of a metric: its id with every character outside
    ``A-Za-z0-9`` turned into an underscore."""
    return re.sub(r"[^A-Za-z0-9]", "_", str(metric_id))


def function_key(function_id: str) -> str:
    return re.sub(r"[^A-Za-z0-9]", "_", str(function_id))


def layer_label(stratum, metric: Optional[dict] = None) -> str:
    """What the curve-set dropdown shows for a layer: the pooled layer is unnamed
    in the bundle, and a class layer named by its key reads in the words of the
    metric's ``stratifier`` block."""
    if stratum in (None, ""):
        return POOLED_LABEL
    block = (metric or {}).get("stratifier")
    if isinstance(block, dict):
        for cls in block.get("classes") or []:
            if str((cls or {}).get("key")) == str(stratum) and cls.get("label"):
                return str(cls["label"])
    return str(stratum)


# --------------------------------------------------------------------------- #
# reading the bundle
# --------------------------------------------------------------------------- #
def outcome_mapping() -> dict[str, dict[str, str]]:
    rows = json.loads(_VENDORED_MAPPING.read_text(encoding="utf-8"))
    return {r["id"]: {k: r.get(k, "-") for k in OUTCOMES} for r in rows}


def _sorted_points(points) -> list[tuple[float, float]]:
    """DEEP's own reading of a curve: usable points, ascending x (stable), and
    each index clamped to 0 to 1."""
    pts = [(float(p["x"]), float(p["y"])) for p in points or []
           if p.get("x") is not None and p.get("y") is not None]
    pts.sort(key=lambda p: p[0])
    return [(x, min(1.0, max(0.0, y))) for x, y in pts]


def metric_layers(metric: dict) -> list[tuple[str, list[tuple[float, float]]]]:
    """``[(label, points)]`` for a metric: its curve layers in the bundle's order,
    or its single curve under an empty label. The order is what DEEP falls back
    on, so it is kept."""
    layers = metric.get("curveLayers") or []
    if layers:
        out = []
        for layer in layers:
            pts = _sorted_points(layer.get("points"))
            if pts:
                out.append((layer_label(layer.get("stratum"), metric), pts))
        if out:
            return out
    return [("", _sorted_points((metric.get("curve") or {}).get("points")))]


def default_layer(metric: dict) -> str:
    """The layer DEEP scores on when no curve set is chosen: the declared
    ``activeStratum`` when it names a layer, else the first layer."""
    layers = metric.get("curveLayers") or []
    if not layers:
        return ""
    names = [str(layer.get("stratum", "")) for layer in layers if layer.get("points")]
    active = metric.get("activeStratum")
    if active and str(active) in names:
        return layer_label(active, metric)
    return layer_label(names[0] if names else "", metric)


def functions_of(bundle: dict) -> list[dict]:
    """Function blocks with metrics, in category order then bundle order."""
    blocks = [fn for fn in bundle.get("metricsByFunction") or [] if fn.get("metrics")]
    rank = {c: i for i, c in enumerate(CATEGORY_ORDER)}
    return sorted(blocks, key=lambda fn: rank.get(fn.get("discipline"), len(rank)))


# --------------------------------------------------------------------------- #
# formulas
# --------------------------------------------------------------------------- #
def interpolation_formula(x: str, xs: list[str], ys: list[str]) -> str:
    """DEEP's ``interp_curve`` over ``KNOTS`` knot cells, as one formula.

    ``xs`` and ``ys`` are cell references in ascending x. A curve with fewer
    knots is padded by repeating its last point, so the right clamp fires
    before any padded segment can.
    """
    n = len(xs)

    def seg(a: int, b: int) -> str:
        return (f"IF({xs[b]}-{xs[a]}<=0,{ys[b]},"
                f"{ys[a]}+({x}-{xs[a]})/({xs[b]}-{xs[a]})*({ys[b]}-{ys[a]}))")
    chain = seg(n - 2, n - 1)
    for k in range(n - 3, -1, -1):
        chain = f"IF({x}<={xs[k + 1]},{seg(k, k + 1)},{chain})"
    body = f"IF({x}<={xs[0]},{ys[0]},IF({x}>={xs[n - 1]},{ys[n - 1]},{chain}))"
    return f"MAX(0,MIN(1,{body}))"


class Sheet:
    """Row-by-row writer with absolute references and defined names."""

    def __init__(self, wb: Workbook, ws, name: str):
        self.wb, self.ws, self.name = wb, ws, name
        self.row = 1

    def ref(self, col: int, row: int) -> str:
        return f"'{self.name}'!${get_column_letter(col)}${row}"

    def rng(self, col: int, r1: int, r2: int) -> str:
        c = get_column_letter(col)
        return f"'{self.name}'!${c}${r1}:${c}${r2}"

    def name_cell(self, name: str, col: int, row: int) -> str:
        ref = self.ref(col, row)
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
        return ref

    def name_range(self, name: str, col: int, r1: int, r2: int) -> str:
        ref = self.rng(col, r1, r2)
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
        return ref

    def cell(self, col: int, value=None, *, row: Optional[int] = None, fill=None, font=None,
             fmt=None, unlocked=False, border=False, align=None):
        c = self.ws.cell(row=row or self.row, column=col)
        if value is not None:
            c.value = value
        if fill is not None:
            c.fill = fill
        if font is not None:
            c.font = font
        if fmt is not None:
            c.number_format = fmt
        if unlocked:
            c.protection = UNLOCKED
        if align is not None:
            c.alignment = align
        if border is True:
            c.border = BOX
        elif border:
            c.border = border
        return c

    def merge(self, r1: int, c1: int, r2: int, c2: int):
        if (r1, c1) != (r2, c2):
            self.ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def box(ws, r1: int, c1: int, r2: int, c2: int, outer: Side = MED, inner: Optional[Side] = THIN):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            old = cell.border
            cell.border = Border(
                left=outer if c == c1 else (inner or old.left),
                right=outer if c == c2 else (inner or old.right),
                top=outer if r == r1 else (inner or old.top),
                bottom=outer if r == r2 else (inner or old.bottom))


def link_formula(ref: str) -> str:
    return f'=IF({ref}="","",{ref})'


# --------------------------------------------------------------------------- #
# the generator
# --------------------------------------------------------------------------- #
class Builder:
    def __init__(self, bundle: dict):
        self.bundle = bundle
        self.wb = Workbook()
        self.functions = functions_of(bundle)
        self.mapping = outcome_mapping()
        self.weights = {"D": 1.0, "i": float(
            (bundle.get("scoringContract") or {}).get("indirectWeight", 0.10)), "-": 0.0}
        self.score_max = float((bundle.get("scoringContract") or {}).get("functionScoreMax", 15))
        # one record per distinct metric, in first-seen order
        self.metrics: dict[str, dict] = {}
        for fn in self.functions:
            for m in fn["metrics"]:
                self.metrics.setdefault(m["metricId"], m)
        self.curve_rows: dict[tuple[str, str], int] = {}     # (metric key, layer label) -> row
        self.strata_range: dict[str, str] = {}               # metric key -> list range
        self.metric_row: dict[str, int] = {}                 # metric key -> Metrics sheet row
        self.entry_cell: dict[str, str] = {}                 # metric key -> entry ref
        self.stratum_cell: dict[str, str] = {}
        self.index_cell: dict[str, str] = {}
        self.fn_score_cell: dict[str, str] = {}               # function id -> ref on Results

    # ---- Reference: every curve as a row of knots -------------------------
    def build_reference(self):
        ws = self.wb.active
        ws.title = "Reference"
        S = Sheet(self.wb, ws, "Reference")
        S.cell(1, "Reference curves, bands and weights", font=Font(bold=True, size=13))
        S.cell(1, "One row per metric and curve set. Knots in ascending x, the last knot repeated "
                  "to fill eight. The DEEP Score sheet reads these rows.", row=2, font=FONT_NOTE)
        heads = (["Curve key", "Metric id", "Metric", "Curve set", "Knots", "Scored against"]
                 + [f"x{i}" for i in range(1, KNOTS + 1)] + [f"y{i}" for i in range(1, KNOTS + 1)])
        S.row = 4
        for c, text in enumerate(heads, 1):
            S.cell(c, text, fill=FILL_BAND, font=FONT_BOLD, border=True)
        first = S.row + 1
        S.row = first
        for mid, m in self.metrics.items():
            key = metric_key(mid)
            for label, pts in metric_layers(m):
                if not pts:
                    continue
                if len(pts) > KNOTS:
                    raise ValueError(f"{mid}: a curve with {len(pts)} points exceeds the "
                                     f"{KNOTS} knots the calculator carries")
                padded = pts + [pts[-1]] * (KNOTS - len(pts))
                S.cell(1, f"{key}|{label}", border=True)
                S.cell(2, mid, border=True)
                S.cell(3, m.get("metricName") or mid, border=True)
                S.cell(4, label, border=True)
                S.cell(5, len(pts), border=True)
                S.cell(6, support_text(m), border=True)
                for i, (x, y) in enumerate(padded):
                    S.cell(7 + i, x, border=True)
                    S.cell(7 + KNOTS + i, y, border=True)
                self.curve_rows[(key, label)] = S.row
                S.row += 1
        last = S.row - 1
        S.name_range("curve_key", 1, first, last)
        for i in range(KNOTS):
            S.name_range(f"curve_x{i + 1}", 7 + i, first, last)
            S.name_range(f"curve_y{i + 1}", 7 + KNOTS + i, first, last)

        # curve-set lists (a list range per stratified metric: labels can be long
        # and carry commas, which an inline validation list cannot)
        S.row += 2
        S.cell(1, "Curve sets by metric", font=FONT_BOLD)
        S.row += 1
        for mid, m in self.metrics.items():
            labels = [label for label, pts in metric_layers(m) if pts]
            if len(labels) < 2:
                continue
            key = metric_key(mid)
            r1 = S.row
            for label in labels:
                S.cell(1, label, border=True)
                S.cell(2, m.get("metricName") or mid, font=FONT_NOTE)
                S.row += 1
            self.strata_range[key] = S.name_range(f"strata_{key}", 1, r1, S.row - 1)

        S.row += 2
        S.cell(1, "Condition classes", font=FONT_BOLD)
        S.row += 1
        for text in ("Metric index: at or below 0.39 Non-Functioning, at or below 0.69 "
                     "Functioning At-Risk, above Functioning.",
                     "Function score (0 to 15): at or below 5 Non-Functioning, at or below 10 "
                     "Functioning At-Risk, above Functioning.",
                     "Outcome weights: a direct contribution counts 1, an indirect one "
                     f"{self.weights['i']:g}."):
            S.cell(1, text)
            S.row += 1
        ws.column_dimensions["A"].width = 58
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 34
        ws.column_dimensions["D"].width = 40
        ws.column_dimensions["F"].width = 60
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    # ---- Metrics: one row per distinct metric, the interpolation -----------
    def build_metrics(self):
        ws = self.wb.create_sheet("Metrics")
        S = Sheet(self.wb, ws, "Metrics")
        S.cell(1, "Metric indices", font=Font(bold=True, size=13))
        S.cell(1, "One row per metric. The value and the curve set come from the DEEP Score "
                  "sheet, the knots from the Reference sheet, and the index is DEEP's "
                  "piecewise-linear interpolation.", row=2, font=FONT_NOTE)
        heads = (["Metric", "Value", "Curve set in use", "Curve key", "Curve row"]
                 + [f"x{i}" for i in range(1, KNOTS + 1)] + [f"y{i}" for i in range(1, KNOTS + 1)]
                 + ["Index"])
        S.row = 4
        for c, text in enumerate(heads, 1):
            S.cell(c, text, fill=FILL_BAND, font=FONT_BOLD, border=True)
        S.row = 5
        for mid, m in self.metrics.items():
            key = metric_key(mid)
            r = S.row
            self.metric_row[key] = r
            S.cell(1, m.get("metricName") or mid, border=True)
            # value and curve set are wired by build_score (the entry cells live there)
            default = default_layer(m)
            value, chosen, ckey, crow = (S.ref(2, r), S.ref(3, r), S.ref(4, r), S.ref(5, r))
            S.cell(4, f'={q(key + "|")}&{chosen}', fill=FILL_CALC, border=True)
            S.cell(5, f"=IFERROR(MATCH({ckey},curve_key,0),"
                      f"MATCH({q(key + '|' + default)},curve_key,0))",
                   fill=FILL_CALC, border=True)
            xs, ys = [], []
            for i in range(KNOTS):
                S.cell(6 + i, f"=INDEX(curve_x{i + 1},{crow})", fill=FILL_CALC, border=True)
                S.cell(6 + KNOTS + i, f"=INDEX(curve_y{i + 1},{crow})", fill=FILL_CALC,
                       border=True)
                xs.append(S.ref(6 + i, r))
                ys.append(S.ref(6 + KNOTS + i, r))
            idx_col = 6 + 2 * KNOTS
            S.cell(idx_col, f'=IF(ISNUMBER({value}),{interpolation_formula(value, xs, ys)},"")',
                   fill=FILL_CALC, border=True, fmt="0.000")
            self.index_cell[key] = S.name_cell(f"idx_{key}", idx_col, r)
            S.row += 1
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["C"].width = 34
        ws.column_dimensions["D"].width = 50
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    # ---- DEEP Score: the sheet the assessor fills ---------------------------
    def build_score(self):
        ws = self.wb.create_sheet(SCORE)
        S = Sheet(self.wb, ws, SCORE)
        self.score_ws = ws
        name = self.bundle.get("assessmentName") or self.bundle.get("assessmentId") or ""
        S.cell(1, "Detailed Evaluation of Ecosystem Processes (DEEP) Calculator",
               font=Font(bold=True, size=14))
        S.cell(1, name, row=2, font=Font(bold=True, size=11))
        S.cell(1, version_line(self.bundle), row=3, font=FONT_NOTE)

        # the site block
        site = [("Stream or site", "site_name"), ("Reach ID", "site_reach"),
                ("Coordinates", "site_coords"), ("Date", "site_date"),
                ("Assessor(s)", "site_assessor")]
        S.row = 5
        for label, nm in site:
            S.cell(1, label, fill=FILL_BAND, font=FONT_BOLD, border=True)
            S.merge(S.row, 2, S.row, 4)
            c = S.cell(2, fill=FILL_INPUT, unlocked=True, border=True, align=LEFT_MID)
            if nm == "site_date":
                c.number_format = "yyyy-mm-dd"
            for col in (3, 4):
                S.cell(col, border=True, fill=FILL_INPUT, unlocked=True)
            S.name_cell(nm, 2, S.row)
            S.row += 1

        S.row += 1
        heads = ["Category", "Function", "Metric", "Units", "Value", "Curve set", "Index",
                 "Function score (0 to 15)", "Condition"]
        head_row = S.row
        for c, text in enumerate(heads, 1):
            S.cell(c, text, fill=FILL_BAND, font=FONT_BOLD, border=True, align=CENTER)
        S.row += 1

        M = Sheet(self.wb, self.wb["Metrics"], "Metrics")
        R = Sheet(self.wb, self.wb["Results"], "Results")
        first_table_row = S.row
        seen: set[str] = set()
        for fn in self.functions:
            fid = fn["functionId"]
            cat = fn.get("discipline") or ""
            r1 = S.row
            idx_refs = []
            for m in fn["metrics"]:
                mid = m["metricId"]
                key = metric_key(mid)
                S.cell(3, m.get("metricName") or mid, border=True, align=LEFT_MID)
                S.cell(4, units_of(m), border=True, align=CENTER)
                labels = [label for label, pts in metric_layers(m) if pts]
                if key not in seen:
                    seen.add(key)
                    S.cell(5, fill=FILL_INPUT, unlocked=True, border=True, align=CENTER)
                    self.entry_cell[key] = S.name_cell(f"in_{key}", 5, S.row)
                    if len(labels) > 1:
                        c = S.cell(6, default_layer(m), fill=FILL_INPUT, unlocked=True,
                                   border=True, align=LEFT_MID)
                        dv = DataValidation(type="list", formula1=f"=strata_{key}",
                                            allow_blank=True, showErrorMessage=True,
                                            errorTitle="Curve set",
                                            error="Choose a curve set from the list.")
                        ws.add_data_validation(dv)
                        dv.add(c)
                        self.stratum_cell[key] = S.name_cell(f"st_{key}", 6, S.row)
                    else:
                        S.cell(6, "", fill=FILL_LINK, border=True)
                    # wire the Metrics sheet row to the entry cells
                    mr = self.metric_row[key]
                    M.cell(2, link_formula(self.entry_cell[key]), row=mr, fill=FILL_CALC,
                           border=True)
                    chosen = (f'=IF({self.stratum_cell[key]}="",{q(default_layer(m))},'
                              f'{self.stratum_cell[key]})' if key in self.stratum_cell
                              else f"={q(labels[0] if labels else '')}")
                    M.cell(3, chosen, row=mr, fill=FILL_CALC, border=True)
                else:
                    # entered once, linked here
                    S.cell(5, link_formula(self.entry_cell[key]), fill=FILL_LINK, border=True,
                           align=CENTER)
                    S.cell(6, link_formula(self.stratum_cell[key]) if key in self.stratum_cell
                           else "", fill=FILL_LINK, border=True, align=LEFT_MID)
                S.cell(7, link_formula(self.index_cell[key]), fill=FILL_CALC, border=True,
                       fmt="0.00", align=CENTER)
                idx_refs.append(S.ref(7, S.row))
                S.row += 1
            r2 = S.row - 1
            # category, function, score and condition span the function's rows
            S.cell(1, cat, row=r1, fill=solid(CATEGORY_FILL.get(cat, C_BAND)), border=True,
                   align=CENTER, font=FONT_BOLD)
            S.cell(2, fn.get("functionName") or fid, row=r1, border=True, align=LEFT_MID,
                   font=FONT_BOLD)
            rng = f"{get_column_letter(7)}{r1}:{get_column_letter(7)}{r2}"
            rr = self.fn_result_row[fid]
            R.cell(2, f"=IF(COUNT('{SCORE}'!{rng})=0,\"\",SUM('{SCORE}'!{rng})/"
                      f"COUNT('{SCORE}'!{rng})*{self.score_max:g})", row=rr, fill=FILL_CALC,
                   border=True)
            S.cell(8, link_formula(self.fn_score_cell[fid]), row=r1, border=True, align=CENTER,
                   fmt="0.0", font=FONT_BOLD)
            score = S.ref(8, r1)
            S.cell(9, f'=IF({score}="","",IF({score}<=5,"{CLASS_LABELS[0]}",'
                      f'IF({score}<=10,"{CLASS_LABELS[1]}","{CLASS_LABELS[2]}")))',
                   row=r1, border=True, align=CENTER)
            for col in (1, 2, 8, 9):
                S.merge(r1, col, r2, col)
            box(ws, r1, 1, r2, 9, outer=MED, inner=THIN)
        last_table_row = S.row - 1

        # class colours on the index, the function score and the condition
        for col, lo, hi in ((7, 0.39, 0.69), (8, 5, 10)):
            L = get_column_letter(col)
            rng = f"{L}{first_table_row}:{L}{last_table_row}"
            top = f"{L}{first_table_row}"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(ISNUMBER({top}),{top}<={lo})'], fill=solid(C_POOR)))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(ISNUMBER({top}),{top}>{lo},{top}<={hi})'], fill=solid(C_FAIR)))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(ISNUMBER({top}),{top}>{hi})'], fill=solid(C_GOOD)))
        rng = f"I{first_table_row}:I{last_table_row}"
        for label, colour in zip(CLASS_LABELS, (C_POOR, C_FAIR, C_GOOD)):
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'I{first_table_row}="{label}"'], fill=solid(colour)))

        # the summary: outcomes and the Ecosystem Condition Index
        S.row += 1
        S.cell(1, "Summary", font=Font(bold=True, size=12))
        S.row += 1
        for c, text in enumerate(["Outcome", "Sub-index (0 to 1)", "Condition"], 1):
            S.cell(c, text, fill=FILL_BAND, font=FONT_BOLD, border=True, align=CENTER)
        S.row += 1
        for outcome in OUTCOMES:
            ref = self.sub_index_cell[outcome]
            S.cell(1, OUTCOME_LABEL[outcome], border=True, font=FONT_BOLD)
            S.cell(2, f"={ref}", border=True, fmt="0.00", align=CENTER)
            v = S.ref(2, S.row)
            S.cell(3, f'=IF({v}<=0.39,"{CLASS_LABELS[0]}",IF({v}<=0.69,"{CLASS_LABELS[1]}",'
                      f'"{CLASS_LABELS[2]}"))', border=True, align=CENTER)
            S.row += 1
        S.cell(1, "Ecosystem Condition Index", border=True, font=FONT_BOLD, align=LEFT_MID)
        S.ws.row_dimensions[S.row].height = 30
        S.cell(2, f"={self.eci_cell}", border=True, fmt="0.00", align=CENTER, font=FONT_BOLD)
        v = S.ref(2, S.row)
        S.cell(3, f'=IF({v}<=0.39,"{CLASS_LABELS[0]}",IF({v}<=0.69,"{CLASS_LABELS[1]}",'
                  f'"{CLASS_LABELS[2]}"))', border=True, align=CENTER)
        self.summary_last = S.row
        S.row += 1
        restriction = claim_restriction_text(self.bundle)
        if restriction:
            # Under the number, not in the notes. The cell above is a ratio over the
            # functions this assessment can score, so on a partial assessment it is a
            # running total and reads as a full-framework index unless told otherwise.
            S.cell(1, restriction, font=FONT_NOTE)
            S.ws.merge_cells(start_row=S.row, start_column=1, end_row=S.row, end_column=3)
            S.row += 1
        S.cell(1, f"Functions with a score: ", font=FONT_NOTE)
        S.cell(2, f"=COUNT({self.fn_score_range})&\" of {len(self.functions)}\"",
               font=FONT_NOTE)
        S.row += 1
        S.cell(1, coverage_text(self.bundle), font=FONT_NOTE)
        S.row += 2

        withheld = self.bundle.get("insufficientReferenceSupport") or []
        if withheld:
            S.cell(1, "Withheld for insufficient reference support (not scored)", font=FONT_BOLD)
            S.row += 1
            for w in withheld:
                fns = ", ".join(str(f.get("functionName") or f.get("functionId"))
                                for f in w.get("functions") or [])
                S.cell(1, str(w.get("metricName") or w.get("metricId") or ""))
                S.cell(3, fns, font=FONT_NOTE)
                S.row += 1
            S.row += 1

        S.cell(1, "Notes", fill=FILL_BAND, font=FONT_BOLD, border=True)
        S.row += 1
        notes_row = S.row
        S.merge(notes_row, 1, notes_row + 5, 9)
        S.cell(1, fill=FILL_INPUT, unlocked=True, align=LEFT_TOP)
        for r in range(notes_row, notes_row + 6):
            for c in range(1, 10):
                cell = ws.cell(row=r, column=c)
                cell.fill, cell.protection = FILL_INPUT, UNLOCKED
        box(ws, notes_row, 1, notes_row + 5, 9, outer=MED, inner=None)
        S.name_cell("site_notes", 1, notes_row)

        for col, width in zip("ABCDEFGHI", (16, 30, 40, 12, 14, 40, 10, 16, 22)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
        ws.sheet_view.showGridLines = False
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    # ---- Results: function scores and the rollup (hidden) -------------------
    def build_results(self):
        ws = self.wb.create_sheet("Results")
        S = Sheet(self.wb, ws, "Results")
        S.cell(1, "Function scores and the outcome rollup", font=Font(bold=True, size=13))
        heads = ["Function", "Score"] + [f"{OUTCOME_LABEL[o]} weight" for o in OUTCOMES] \
            + [f"{OUTCOME_LABEL[o]} weighted" for o in OUTCOMES] \
            + [f"{OUTCOME_LABEL[o]} maximum" for o in OUTCOMES]
        S.row = 3
        for c, text in enumerate(heads, 1):
            S.cell(c, text, fill=FILL_BAND, font=FONT_BOLD, border=True)
        S.row = 4
        first = S.row
        self.fn_result_row: dict[str, int] = {}
        for fn in self.functions:
            fid = fn["functionId"]
            self.fn_result_row[fid] = S.row
            S.cell(1, fn.get("functionName") or fid, border=True)
            S.cell(2, "", fill=FILL_CALC, border=True)          # filled by build_score
            self.fn_score_cell[fid] = S.name_cell(f"fs_{function_key(fid)}", 2, S.row)
            score = S.ref(2, S.row)
            codes = self.mapping.get(fid) or {}
            for i, outcome in enumerate(OUTCOMES):
                w = self.weights.get(codes.get(outcome, "-"), 0.0)
                S.cell(3 + i, w, border=True)
                wref = S.ref(3 + i, S.row)
                S.cell(6 + i, f"=IF(ISNUMBER({score}),{score}*{wref},0)", fill=FILL_CALC,
                       border=True)
                S.cell(9 + i, f"=IF(ISNUMBER({score}),{self.score_max:g}*{wref},0)",
                       fill=FILL_CALC, border=True)
            S.row += 1
        last = S.row - 1
        self.fn_score_range = S.rng(2, first, last)
        S.row += 1
        S.cell(1, "Outcome", fill=FILL_BAND, font=FONT_BOLD, border=True)
        S.cell(2, "Sub-index", fill=FILL_BAND, font=FONT_BOLD, border=True)
        S.row += 1
        self.sub_index_cell: dict[str, str] = {}
        for i, outcome in enumerate(OUTCOMES):
            S.cell(1, OUTCOME_LABEL[outcome], border=True)
            wsum, wmax = S.rng(6 + i, first, last), S.rng(9 + i, first, last)
            S.cell(2, f"=IF(SUM({wmax})=0,0,SUM({wsum})/SUM({wmax}))", fill=FILL_CALC,
                   border=True)
            self.sub_index_cell[outcome] = S.name_cell(f"sub_index_{outcome}", 2, S.row)
            S.row += 1
        S.cell(1, "Ecosystem Condition Index", border=True, font=FONT_BOLD)
        S.cell(2, "=(" + "+".join(self.sub_index_cell[o] for o in OUTCOMES) + ")/3",
               fill=FILL_CALC, border=True)
        self.eci_cell = S.name_cell("eci", 2, S.row)
        ws.column_dimensions["A"].width = 40
        ws.sheet_state = "hidden"
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    # ---- ChartData (hidden) and the two bar charts --------------------------
    def build_charts(self):
        cd = self.wb.create_sheet("ChartData")
        C = Sheet(self.wb, cd, "ChartData")
        C.cell(1, "Chart series (hidden). One column per class so each bar takes the class "
                  "colour.", font=FONT_NOTE)
        for c, text in enumerate(["Function", "Score", "Non-Functioning (0-5)",
                                  "Functioning At-Risk (6-10)", "Functioning (11-15)"], 1):
            C.cell(c, text, row=2, font=FONT_BOLD)
        r = 3
        for fn in self.functions:
            C.cell(1, fn.get("functionName") or fn["functionId"], row=r)
            C.cell(2, link_formula(self.fn_score_cell[fn["functionId"]]), row=r)
            s = C.ref(2, r)
            # chart labels take the number format of these cells, so format them as the sheet prints them
            C.cell(3, f"=IF(ISNUMBER({s}),IF({s}<=5,{s},NA()),NA())", row=r, fmt="0.0")
            C.cell(4, f"=IF(ISNUMBER({s}),IF(AND({s}>5,{s}<=10),{s},NA()),NA())", row=r, fmt="0.0")
            C.cell(5, f"=IF(ISNUMBER({s}),IF({s}>10,{s},NA()),NA())", row=r, fmt="0.0")
            r += 1
        f_first, f_last = 3, r - 1
        r += 1
        for c, text in enumerate(["Outcome", "Score", "Non-Functioning (0.00-0.39)",
                                  "Functioning At-Risk (0.40-0.69)", "Functioning (0.70-1.00)"], 1):
            C.cell(c, text, row=r, font=FONT_BOLD)
        o_head = r
        r += 1
        for outcome in OUTCOMES:
            C.cell(1, OUTCOME_LABEL[outcome], row=r)
            C.cell(2, f"={self.sub_index_cell[outcome]}", row=r)
            s = C.ref(2, r)
            C.cell(3, f"=IF({s}<=0.39,{s},NA())", row=r, fmt="0.00")
            C.cell(4, f"=IF(AND({s}>0.39,{s}<=0.69),{s},NA())", row=r, fmt="0.00")
            C.cell(5, f"=IF({s}>0.69,{s},NA())", row=r, fmt="0.00")
            r += 1
        o_first, o_last = o_head + 1, r - 1
        cd.column_dimensions["A"].width = 38
        cd.sheet_state = "hidden"
        cd.protection.sheet = True
        cd.protection.password = PASSWORD

        def bar(title, head_row, first, last, vmax, unit, fmt, height):
            ch = BarChart()
            ch.type = "bar"
            ch.grouping = "clustered"
            ch.overlap = 100
            ch.gapWidth = 40
            ch.title = title
            ch.title.overlay = False
            ch.style = 2
            ch.add_data(Reference(cd, min_col=3, max_col=5, min_row=head_row, max_row=last),
                        titles_from_data=True)
            ch.set_categories(Reference(cd, min_col=1, min_row=first, max_row=last))
            for s, colour in zip(ch.series, (C_POOR, C_FAIR, C_GOOD)):
                s.graphicalProperties = GraphicalProperties(solidFill=colour)
                s.graphicalProperties.line = LineProperties(solidFill="000000")
            ch.x_axis.scaling.orientation = "maxMin"
            ch.x_axis.delete = False
            ch.y_axis.delete = False
            ch.y_axis.scaling.min = 0
            ch.y_axis.scaling.max = vmax
            ch.y_axis.majorUnit = unit
            ch.y_axis.numFmt = NumFmt(formatCode=fmt, sourceLinked=False)
            ch.y_axis.majorGridlines = None
            ch.dataLabels = DataLabelList()
            ch.dataLabels.showVal = True
            ch.dataLabels.showSerName = False
            ch.dataLabels.showCatName = False
            ch.dataLabels.showLegendKey = False
            ch.legend.position = "b"
            ch.legend.overlay = False
            ch.width, ch.height = 22, height
            return ch

        self.score_ws.add_chart(bar("Function Score", 2, f_first, f_last, 15, 1, "0",
                                    max(8.0, 0.62 * len(self.functions))), "K5")
        self.score_ws.add_chart(bar("Outcome Score", o_head, o_first, o_last, 1, 0.1, "0.00", 6.5),
                                f"K{5 + int(2.2 * max(13, len(self.functions)))}")

    # ---- Instructions and Metadata -----------------------------------------
    def build_instructions(self):
        ws = self.wb.create_sheet("Instructions", 0)
        S = Sheet(self.wb, ws, "Instructions")
        S.cell(1, "Detailed Evaluation of Ecosystem Processes (DEEP) Calculator",
               font=Font(bold=True, size=14))
        S.cell(1, self.bundle.get("assessmentName") or "", row=2, font=Font(bold=True, size=11))
        S.row = 4
        lines = [
            "This workbook scores one detailed assessment. It carries the assessment's metrics "
            "and their reference curves, and it scores exactly as the DEEP web application does.",
            "1. On the DEEP Score sheet, fill the site block and type each metric's measured "
            "value in its orange cell, in the units shown.",
            "2. Where a metric has several curve sets, choose the set that fits the reach from "
            "the list beside the value. Left alone, the default set applies.",
            "3. Leave a value blank when the metric does not apply. It drops out of the score.",
            "4. A metric that serves more than one function is entered once. The grey cells "
            "repeat it.",
            "5. The index of each metric (0 to 1) is read from its reference curve. The function "
            "score is the mean index of the function's metrics times 15.",
            "6. The outcome sub-indices weigh each function's score by its direct or indirect "
            "contribution, and the Ecosystem Condition Index is the mean of the three.",
            "Condition classes. Index at or below 0.39 Non-Functioning, at or below 0.69 "
            "Functioning At-Risk, above Functioning. Function score at or below 5 "
            "Non-Functioning, at or below 10 Functioning At-Risk, above Functioning.",
            "Sheets. DEEP Score (entries and results), Metrics (the interpolation of every "
            "metric), Reference (curves, curve sets, classes), Metadata (assessment identity).",
            "The sheets are protected against accidental edits. The password is deep.",
        ]
        statement = (self.bundle.get("referenceMethod") or {}).get("statement")
        if statement:
            lines.insert(1, "Reference condition. " + str(statement))
        for text in lines:
            S.cell(1, text, align=LEFT_TOP)
            ws.row_dimensions[S.row].height = 15 * max(1, -(-len(text) // 110))
            S.row += 1
        ws.column_dimensions["A"].width = 120
        ws.sheet_view.showGridLines = False
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    def build_metadata(self):
        ws = self.wb.create_sheet("Metadata")
        S = Sheet(self.wb, ws, "Metadata")
        S.cell(1, "Metadata and version information", font=Font(bold=True, size=13))
        S.row = 3
        for c, text in enumerate(["Item", "Value"], 1):
            S.cell(c, text, fill=FILL_BAND, font=FONT_BOLD, border=True)
        S.row += 1
        lib = self.bundle.get("library") or {}
        region = self.bundle.get("region") or {}
        ref = self.bundle.get("referenceMethod") or {}
        contract = self.bundle.get("scoringContract") or {}
        rows = [
            ("Assessment id", self.bundle.get("assessmentId")),
            ("Assessment name", self.bundle.get("assessmentName")),
            ("Version", lib.get("version")),
            ("Lifecycle status", lib.get("status") or lib.get("lifecycle")),
            ("Content digest", self.bundle.get("contentDigest")),
            ("Region", " ".join(str(x) for x in (region.get("kind"), region.get("code"),
                                                 region.get("name")) if x)),
            ("Source", self.bundle.get("sourceCitation")),
            ("Reference method", ref.get("method")),
            ("Reference screen", ref.get("screenLabel")),
            ("Scoring method", contract.get("methodVersion")),
            ("Functions scored", len(self.functions)),
            ("Metrics", len(self.metrics)),
            ("Function coverage", coverage_text(self.bundle)),
            ("Generator", GENERATOR),
            ("Generator version", GENERATOR_VERSION),
            ("Formula vocabulary", "IF AND MIN MAX INDEX MATCH IFERROR ISNUMBER COUNT SUM NA"),
        ]
        for k, v in rows:
            S.cell(1, k, border=True)
            S.cell(2, "" if v is None else v, border=True)
            name = "meta_" + re.sub(r"[^a-z0-9]+", "_", k.lower()).strip("_")
            S.name_cell(name, 2, S.row)
            S.row += 1
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 100
        ws.sheet_view.showGridLines = False
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    def build(self) -> bytes:
        if not self.functions:
            raise ValueError("the bundle scores no function; nothing to build a calculator from")
        self.build_reference()
        self.build_metrics()
        self.build_results()
        self.build_score()
        self.build_charts()
        self.build_instructions()
        self.build_metadata()
        order = ["Instructions", SCORE, "Metrics", "Results", "Reference", "Metadata", "ChartData"]
        self.wb._sheets = [self.wb[name] for name in order]
        self.wb.active = 1
        for ws in self.wb.worksheets:
            ws.sheet_view.tabSelected = ws.title == SCORE
        stamp = bundle_stamp(self.bundle)
        self.wb.properties.creator = "DEEP"
        self.wb.properties.lastModifiedBy = "DEEP"
        self.wb.properties.title = "DEEP calculator"
        self.wb.properties.created = stamp
        self.wb.properties.modified = stamp
        self.wb.calculation.fullCalcOnLoad = True
        buf = io.BytesIO()
        self.wb.save(buf)
        return repack(buf.getvalue(), stamp)


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def q(text: str) -> str:
    """A string literal for a formula."""
    return '"' + str(text).replace('"', '""') + '"'


def units_of(metric: dict) -> str:
    label = str(metric.get("xLabel") or "")
    if label.endswith(")") and "(" in label:
        return label[label.rfind("(") + 1:-1].strip()
    return ""


def support_text(metric: dict) -> str:
    """One line on what the metric is scored against, for the Reference sheet."""
    # A curve that rests on no station of this ecoregion says what it does rest
    # on. The bundle carries the sentence, so the workbook and the app agree.
    basis = str(metric.get("basis") or "")
    sup_ = metric.get("referenceSupport") if isinstance(metric.get("referenceSupport"),
                                                        dict) else {}
    # a curve from a rung above the ecoregion hierarchy, including a published
    # benchmark, which is marked fixed downstream but is regional, not universal
    ladder = str(sup_.get("status") or "") in ("national", "modeled", "published")
    if basis in ("national-reference", "modeled-reference", "published-benchmark") and (
            ladder or str(metric.get("criteriaBasis") or "") != "fixed"):
        stated = str(metric.get("basisStatement") or sup_.get("basisStatement") or "").strip()
        if stated:
            return stated
    if str(metric.get("criteriaBasis") or "") == "fixed" and not ladder:
        return "Fixed criteria, the same in every region"
    if ladder:
        # never the station-pool sentence: its counts would describe a model fit
        return str(metric.get("basisLabel") or sup_.get("basisLabel") or basis)
    sup = metric.get("referenceSupport")
    carried = metric.get("carriedForward") if isinstance(metric.get("carriedForward"), dict) else {}
    tail = (f", carried forward from version {carried.get('fromVersion')}"
            if carried.get("fromVersion") else "")
    if isinstance(sup, dict) and sup.get("status"):
        # methodology 0.14: a pool admitted under the regional screen says so
        lim = sup.get("agricultureLimit")
        screen = (f" under the regional screen (agriculture at most {float(lim):g} percent)"
                  if sup.get("screenId") and lim is not None else "")
        if sup.get("status") == "local":
            return f"{sup.get('nUsable')} least-disturbed stations of this ecoregion{tail}"
        if sup.get("status") == "local_relaxed":
            return (f"{sup.get('nUsable')} least-disturbed streams of this ecoregion{screen}"
                    f"{tail}")
        where = (f"NARS-9 region {sup.get('regionCode')}" if sup.get("level") == "nars9" else
                 f"{sup.get('levelLabel') or sup.get('level')} ecoregion {sup.get('regionCode')}"
                 + (f" ({sup.get('regionName')})" if sup.get("regionName") else ""))
        return (f"{sup.get('nUsable')} least-disturbed stations borrowed from {where}{screen}"
                f", transfer risk {sup.get('transferRisk')}{tail}")
    n = metric.get("referenceN")
    return f"{int(n)} reference sites" if isinstance(n, (int, float)) else ""


def version_line(bundle: dict) -> str:
    lib = bundle.get("library") or {}
    parts = []
    if lib.get("version") is not None:
        parts.append(f"Version {lib.get('version')}")
    status = lib.get("status") or lib.get("lifecycle")
    if status:
        parts.append(str(status).replace("_", " ").title())
    if bundle.get("sourceCitation"):
        parts.append(str(bundle["sourceCitation"]))
    return " | ".join(parts)


def coverage_text(bundle: dict) -> str:
    cov = bundle.get("functionCoverage") or {}
    if not cov.get("total"):
        return ""
    text = f"This assessment covers {cov.get('covered')} of {cov.get('total')} STAF functions"
    if cov.get("excluded"):
        text += f" ({cov.get('excluded')} documented exclusions)"
    return text + "."


def claim_restriction_text(bundle: dict) -> str:
    """What the index above the line may and may not be claimed as.

    A sub-index is a ratio over the functions the assessment can score, so an
    assessment missing some of the framework lands on the same 0 to 1 scale as a
    complete one and reads as comparable. DEEP restricts the claim to an interval
    in that case; the workbook cannot compute an interval from entries alone, so it
    says plainly what the number is and is not.
    """
    cov = bundle.get("functionCoverage") or {}
    total, covered = cov.get("total"), cov.get("covered")
    if not total or covered is None or covered >= total:
        return ""
    missing = int(total) - int(covered)
    return (f"This index is computed over the {covered} functions this assessment scores. "
            f"{missing} of the {total} STAF functions are not assessed, so it is not "
            f"comparable with a full-framework index and no condition class is claimed "
            f"for the assessment as a whole.")


def bundle_stamp(bundle: dict) -> _dt.datetime:
    """The workbook's created and modified time: the version's own
    ``library.updatedAt``, so a rebuild of the same version reproduces the bytes."""
    raw = str((bundle.get("library") or {}).get("updatedAt") or "")
    try:
        stamp = _dt.datetime.fromisoformat(raw.replace("Z", "+00:00"))
        return stamp.astimezone(_dt.timezone.utc).replace(tzinfo=None, microsecond=0)
    except ValueError:
        return FALLBACK_STAMP


NA_AS_BLANK = ('<extLst><ext uri="{56B9EC1D-385E-4148-901F-78D8002777C0}" '
               'xmlns:c16r3="http://schemas.microsoft.com/office/drawing/2017/03/chart">'
               '<c16r3:dataDisplayOptions16><c16r3:dispNaAsBlank val="1"/>'
               '</c16r3:dataDisplayOptions16></ext></extLst>')


def repack(source: bytes, stamp: _dt.datetime) -> bytes:
    """Rewrite the package with fixed entry timestamps so the bytes are
    reproducible, pin the two times openpyxl stamps on save, and give each chart
    the "#N/A as blank" option (without it every unscored bar prints "#N/A")."""
    text_stamp = stamp.strftime("%Y-%m-%dT%H:%M:%SZ")
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as zin, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                for tag in ("modified", "created"):
                    text = re.sub(rf"(<dcterms:{tag}[^>]*>)[^<]*(</dcterms:{tag}>)",
                                  rf"\g<1>{text_stamp}\g<2>", text)
                data = text.encode("utf-8")
            elif info.filename.startswith("xl/charts/chart") and info.filename.endswith(".xml"):
                text = data.decode("utf-8")
                if text.count("</chart>") == 1 and "dispNaAsBlank" not in text:
                    data = text.replace("</chart>", NA_AS_BLANK + "</chart>").encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0
            zout.writestr(zi, data)
    return out.getvalue()


def build_calculator(bundle: dict) -> bytes:
    """The calculator workbook for a published (or about to be published) bundle."""
    return Builder(bundle).build()


def sha256_of(data: bytes) -> str:
    return "sha256:" + hashlib.sha256(data).hexdigest()
