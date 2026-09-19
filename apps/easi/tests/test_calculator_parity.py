"""Same inputs, same EASI results: the Excel calculator against the application.

The retained case set (``tests/data/calculator_cases.json``, built by
``tests/calculator_cases.py``) holds evidence records scored by the engine,
the calculator entries those scores rest on, and the expected results. Three
gates run on it:

1. the engine still reproduces the stored expectations (a drift gate on the
   frozen method);
2. the committed workbook is byte-identical to what the generator produces
   from the scoring definitions, and its reference tables and metadata equal
   the live catalog;
3. the workbook, evaluated with the ``formulas`` package, gives the expected
   rating, score, index, completeness and route per metric and the expected
   sub-indices, ECI, rated count and provisional flag per case.

The Excel COM gate (``EASI_EXCEL_PARITY=1``) repeats gate 3 in the installed
Excel, the target runtime.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
from pathlib import Path

import openpyxl
import pytest

import calculator_cases as cc
from calculator_eval import ExcelComBackend, FormulasBackend

ALLOWED_FUNCTIONS = {"IF", "AND", "OR", "NOT", "MIN", "MAX", "ROUND", "INDEX", "MATCH", "IFERROR",
                     "ISNUMBER", "COUNT", "SUM", "SUMPRODUCT", "AVERAGE", "STDEV.P", "PRODUCT",
                     "UPPER", "TRIM", "SUBSTITUTE", "LEN", "NA"}
SHEETS = ["Instructions", "EASI Score", "Metrics", "Results", "Reference", "Metadata", "ChartData"]
HIDDEN = {"Results", "ChartData"}
OUTPUT_NAMES = ([f"m{n:02d}_{field}" for n in range(1, 21)
                 for field in ("rating", "score", "index", "status", "computed")]
                + [f"m{n:02d}_route" for n in range(1, 21)]
                + ["sub_index_physical", "sub_index_chemical", "sub_index_biological", "eci",
                   "sub_index_physical_display", "sub_index_chemical_display", "sub_index_biological_display",
                   "eci_display", "functions_rated", "provisional"])


@pytest.fixture(scope="module")
def fixture() -> dict:
    if os.environ.get("EASI_WRITE_GOLDEN") == "1":
        data = cc.build_fixture()
        cc.FIXTURE.write_text(json.dumps(data, indent=1) + "\n", encoding="utf-8")
        return data
    if not cc.FIXTURE.exists():
        pytest.skip("run with EASI_WRITE_GOLDEN=1 to build tests/data/calculator_cases.json")
    return cc.load_fixture()


@pytest.fixture(scope="module")
def backend():
    pytest.importorskip("formulas")
    return FormulasBackend(cc.WORKBOOK)


def _same(a, b, tol=1e-9) -> bool:
    if a is None or b is None:
        return a is None and b is None
    if isinstance(a, str) or isinstance(b, str):
        return str(a) == str(b)
    return abs(float(a) - float(b)) <= tol


def _display_ok(raw, shown, expected_shown) -> bool:
    """Two-decimal display, tolerant only at an exact half-hundredth tie."""
    if raw is None or expected_shown is None:
        return shown is None and expected_shown is None
    if _same(shown, expected_shown, 1e-9):
        return True
    scaled = float(raw) * 100
    tie = abs(scaled - round(scaled)) > 0.5 - 1e-9
    return tie and abs(float(shown) - float(expected_shown)) <= 0.01 + 1e-9


def compare(case: dict, got: dict) -> list[str]:
    """Mismatches between the workbook outputs and the case expectations."""
    exp = case["expected"]
    problems = []
    for n in range(1, 21):
        key = f"m{n:02d}"
        e = exp["metrics"][key]
        if not _same(got.get(f"{key}_rating"), e["rating"]):
            problems.append(f"{key} rating {got.get(f'{key}_rating')!r} != {e['rating']!r}")
        if not _same(got.get(f"{key}_score"), e["score"]):
            problems.append(f"{key} score {got.get(f'{key}_score')!r} != {e['score']!r}")
        idx = got.get(f"{key}_index")
        if e["index"] is None:
            if idx is not None:
                problems.append(f"{key} index {idx!r} != None")
        elif idx is None or abs(round(float(idx), 3) - float(e["index"])) > 1e-9:
            problems.append(f"{key} index {idx!r} != {e['index']!r}")
        status = got.get(f"{key}_status")
        if status != cc.COMPLETENESS.get(e["completeness"], "not rated"):
            problems.append(f"{key} status {status!r} != {e['completeness']!r}")
        if "route" in e and got.get(f"{key}_route") != e["route"]:
            problems.append(f"{key} route {got.get(f'{key}_route')!r} != {e['route']!r}")
        # under an override score the workbook still shows what the evidence computed
        if "computed" in e and not _same(got.get(f"{key}_computed"), e["computed"]):
            problems.append(f"{key} computed {got.get(f'{key}_computed')!r} != {e['computed']!r}")
    for outcome in ("physical", "chemical", "biological"):
        raw = (exp["subIndicesRaw"] or {}).get(outcome)
        if not _same(got.get(f"sub_index_{outcome}"), raw):
            problems.append(f"{outcome} sub-index {got.get(f'sub_index_{outcome}')!r} != {raw!r}")
        if not _display_ok(raw, got.get(f"sub_index_{outcome}_display"), (exp["subIndices"] or {}).get(outcome)):
            problems.append(f"{outcome} display {got.get(f'sub_index_{outcome}_display')!r} != "
                            f"{(exp['subIndices'] or {}).get(outcome)!r}")
    if not _same(got.get("eci"), exp["eciRaw"]):
        problems.append(f"eci {got.get('eci')!r} != {exp['eciRaw']!r}")
    if not _display_ok(exp["eciRaw"], got.get("eci_display"), exp["eci"]):
        problems.append(f"eci display {got.get('eci_display')!r} != {exp['eci']!r}")
    if not _same(got.get("functions_rated"), exp["rated"]):
        problems.append(f"rated {got.get('functions_rated')!r} != {exp['rated']!r}")
    if got.get("provisional") != ("yes" if exp["provisional"] else "no"):
        problems.append(f"provisional {got.get('provisional')!r} != {exp['provisional']!r}")
    return problems


# --------------------------------------------------------------------------- #
# gate 1: the engine still reproduces the stored expectations
# --------------------------------------------------------------------------- #
def test_engine_reproduces_the_stored_case_expectations(fixture):
    from easi.national import method_version
    assert fixture["method_version"] == method_version()
    drift = []
    for case in fixture["cases"]:
        report = cc.score_case(case)
        if cc.expected_from(report) != case["expected"]:
            drift.append(case["id"])
        if cc.entries_from(report, case) != case["entries"]:
            drift.append(case["id"] + " (entries)")
    assert not drift, f"{len(drift)} cases drifted from the stored expectations: {drift[:10]}"


# --------------------------------------------------------------------------- #
# gate 2: the committed workbook is what the generator produces
# --------------------------------------------------------------------------- #
def test_committed_workbook_equals_the_generated_bytes():
    assert cc.WORKBOOK.exists(), "run scripts/build_calculator.py"
    generated = cc.bc.generate()
    assert hashlib.sha256(generated).hexdigest() == hashlib.sha256(cc.WORKBOOK.read_bytes()).hexdigest()


def test_workbook_structure_and_metadata():
    from easi import config
    from easi.national import method_version
    wb = openpyxl.load_workbook(cc.WORKBOOK)
    assert wb.sheetnames == SHEETS
    for ws in wb.worksheets:
        assert ws.protection.sheet, f"{ws.title} is not protected"
        assert (ws.sheet_state == "hidden") == (ws.title in HIDDEN), ws.title
    score = wb["EASI Score"]
    assert score.freeze_panes == "A8" and score.sheet_view.showGridLines is False
    assert len(score._charts) == 2, "the Function Score and Outcome Score charts"
    rules = sum(len(cf.rules) for cf in score.conditional_formatting)
    assert rules >= 21, "three class colours on ratings, input ratings, scores, indices and the summary"
    assert score.merged_cells.ranges, "the outcome, category and function cells are merged"
    functions, external = set(), []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    functions.update(re.findall(r"([A-Z][A-Z0-9.]+)\(", cell.value))
                    if "[" in cell.value:
                        external.append(cell.coordinate)
                    assert len(cell.value) <= 8192, cell.coordinate
    assert functions <= ALLOWED_FUNCTIONS, functions - ALLOWED_FUNCTIONS
    assert not external, external
    names = {k: v.attr_text for k, v in wb.defined_names.items()}
    for name in OUTPUT_NAMES + ["in_impervious", "ctx_region", "ov_stageClass", "curve_keys", "anchor_good"]:
        assert name in names, name
    # one Override Score under every function, as on the Assessment page, and the computed
    # rating kept in a cell of its own beside the final one
    overrides = {n for n in names if re.fullmatch(r"ov_m\d\d_rating", n)}
    assert overrides == {cc.bc.score_override_name(mkey) for mkey in cc.FUNCTIONS} and len(overrides) == 20
    for mkey in cc.FUNCTIONS:
        assert names[f"{mkey}_computed"] != names[f"{mkey}_rating"], mkey
    # each is the last row of its function, labelled as the owner asked, and set apart from the
    # function's metrics: a band of its own, a double rule above it, a bold prompt beside the entry
    for mkey in cc.FUNCTIONS:
        row = int(names[cc.bc.score_override_name(mkey)].rsplit("$", 1)[1])
        label, entry = score.cell(row=row, column=9), score.cell(row=row, column=10)
        assert label.value == "Override Score (Optional)" and label.font.b and label.alignment.horizontal == "right"
        assert label.fill.fgColor.rgb.endswith(cc.bc.C_OVERRIDE) and entry.fill.fgColor.rgb.endswith(cc.bc.C_INPUT)
        assert label.border.top.style == "double" and score.cell(row=row - 1, column=9).border.bottom.style == "double"
        assert label.border.bottom.style == "medium", "the override closes its function's block"
        assert score.row_dimensions[row].height >= cc.bc.OVERRIDE_ROW_HEIGHT
        assert score.cell(row=row, column=6).value is None, "inside the function's merged name, rating and score"
    unlocked = [ws.cell(*openpyxl.utils.cell.coordinate_to_tuple(names[n].split("!")[1].replace("$", "")))
                for n in names if n.startswith(("in_", "ctx_", "ov_", "site_"))
                for ws in [wb[names[n].split("!")[0].strip("'")]]]
    assert unlocked and all(not c.protection.locked for c in unlocked)
    meta = {r[0].value: r[1].value for r in wb["Metadata"].iter_rows(min_row=3) if r[0].value}
    identity = config.scoring_identity()
    assert meta["Scoring method digest"] == method_version()
    assert meta["Catalog sha256"] == identity["catalog_sha256"]
    assert meta["Curves sha256"] == identity["curves_sha256"]
    assert meta["Calculator version"] == cc.bc.TEMPLATE_VERSION
    # one version: the generator reads it from the module that serves the workbook
    from easi import calculator
    assert cc.bc.TEMPLATE_VERSION == calculator.TEMPLATE_VERSION == "1.0"
    assert cc.bc.OUT_NAME == calculator.blank_filename() == cc.WORKBOOK.name
    ref = wb["Reference"]
    keys = [c.value for c in ref["A"] if isinstance(c.value, str) and "|" in c.value]
    curves = cc.sm.curve_sets()
    assert len(keys) == sum(len(s["curves"]) for s in curves.values()) == identity["curve_count"]


# --------------------------------------------------------------------------- #
# gate 3: the workbook scores every case like the engine
# --------------------------------------------------------------------------- #
GROUPS = ["base", "band-edges", "curves", "counts", "missing", "routes", "overrides", "strata", "rollup",
          "override-score"]


@pytest.mark.parametrize("group", GROUPS)
def test_workbook_matches_the_engine(fixture, backend, group):
    cases = [c for c in fixture["cases"] if c["group"] == group]
    assert cases, group
    failures = []
    for case in cases:
        got = backend.evaluate(case["entries"], OUTPUT_NAMES)
        problems = compare(case, got)
        if problems:
            failures.append(f"{case['id']}: " + "; ".join(problems[:6]))
    assert not failures, f"{len(failures)} of {len(cases)} cases differ:\n" + "\n".join(failures[:25])


# --------------------------------------------------------------------------- #
# the monthly flow helper (not linked to any rating, so gate 3 never reaches it)
# --------------------------------------------------------------------------- #
HELPER_FLOWS = [14.178, 13.866, 53.395, 70.182, 34.141, 14.33, 7.674, 5.962, 5.026, 14.334, 20.306, 18.268]


def _helper_entries(flows):
    return {f"in_m{month:02d}": value for month, value in enumerate(flows, 1)}


def test_the_monthly_flow_helper_computes_the_engines_variability(backend):
    from easi.metrics import hydraulics
    want = hydraulics.monthly_flow_cv({f"qe_{m:02d}": v for m, v in enumerate(HELPER_FLOWS, 1)})
    assert want == 0.851775
    assert backend.evaluate(_helper_entries(HELPER_FLOWS), ["flow_cv_helper"])["flow_cv_helper"] == want
    # a full year or nothing, and never a division by a zero mean
    assert backend.evaluate(_helper_entries(HELPER_FLOWS[:11]), ["flow_cv_helper"])["flow_cv_helper"] is None
    assert backend.evaluate(_helper_entries([0.0] * 12), ["flow_cv_helper"])["flow_cv_helper"] is None


def test_the_helper_names_its_newer_function_the_way_the_file_format_requires():
    # STDEV.P postdates the original file format: written bare, Excel reads #NAME?, while the
    # formulas package accepts either spelling, so only the file text can catch it. The error
    # stayed hidden behind the IF until the application began to enter the twelve flows.
    import zipfile
    with zipfile.ZipFile(cc.WORKBOOK) as archive:
        sheets = "".join(archive.read(n).decode("utf-8") for n in archive.namelist()
                         if n.startswith("xl/worksheets/"))
    assert "_xlfn.STDEV.P(" in sheets
    assert not re.search(r"(?<!_xlfn\.)STDEV\.P\(", sheets)


@pytest.mark.excel
def test_excel_computes_the_monthly_flow_helper():
    if not ExcelComBackend.available():
        pytest.skip("set EASI_EXCEL_PARITY=1 on a Windows box with Excel and pywin32 under py -3.12")
    got = ExcelComBackend(cc.WORKBOOK).evaluate(_helper_entries(HELPER_FLOWS), ["flow_cv_helper"])
    assert got["flow_cv_helper"] == 0.851775


@pytest.mark.excel
def test_excel_matches_the_engine(fixture):
    if not ExcelComBackend.available():
        pytest.skip("set EASI_EXCEL_PARITY=1 on a Windows box with Excel and pywin32 under py -3.12")
    backend = ExcelComBackend(cc.WORKBOOK)
    cases = fixture["cases"]
    results = backend.evaluate_many([(c["entries"], OUTPUT_NAMES) for c in cases])
    failures = []
    for case, got in zip(cases, results):
        problems = compare(case, got)
        if problems:
            failures.append(f"{case['id']}: " + "; ".join(problems[:6]))
    assert not failures, f"{len(failures)} of {len(cases)} cases differ in Excel:\n" + "\n".join(failures[:25])
