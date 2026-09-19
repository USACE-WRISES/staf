"""Observed evidence on the Assessment cards (2026-09-18).

The technical report gives Channel evolution and Bank condition a documented
observation as their top evidence level, and the Excel calculator has the cells
for it. The application offered no way to enter one: only the engine function
existed (``assessment.apply_observed_evidence``). The two cards now carry the
entries, the scoring pipeline applies them in the engine's own order (rescore,
then the observations), and a completed workbook carries them into those cells.

The card blocks are module-level functions, so they render offline; the wiring
inside the server function is pinned as source text.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

import openpyxl
import pytest

import calculator_cases as cc
from easi import assessment, calculator, report, screening_methods
from easi.metrics import geomorphology as g

app = pytest.importorskip("app")

WWW = Path(app.__file__).parent / "www"
SRC = Path(app.__file__).read_text(encoding="utf-8")
JS = (WWW / "worksheet.js").read_text(encoding="utf-8")
CSS = (WWW / "styles.css").read_text(encoding="utf-8")
CHAN, BANK = g.CHANNEL_EVOL_ID, g.BANK_EROSION_ID
DELINEATION = {"comid": 9327042, "gnis_name": "Mink Brook", "snapped_lat": 43.6858, "snapped_lon": -72.2367}


def _scored(observed=None, ratings=None) -> dict:
    """What the application's scored() builds: rescore, then the observations."""
    rep = cc.score_case({"record": {}, "observed": None})
    rep = assessment.rescore(rep, ratings or {})
    return assessment.apply_observed_evidence(rep, observed) if observed else rep


def _rows(rep) -> dict:
    return {row["metricId"]: row for row in rep["metricRows"]}


# --------------------------------------------------------------------------- #
# the two cards, and only those two
# --------------------------------------------------------------------------- #
def test_only_the_two_metrics_with_an_observed_route_carry_the_block():
    assert set(app.OBSERVED_INPUTS) == {CHAN, BANK}
    assert app.OBSERVED_INPUTS[CHAN] == ("stageClass", "indicators")
    assert app.OBSERVED_INPUTS[BANK] == ("erodingBankPct", "armoredBankPct")
    others = [mid for mid in cc.FUNCTIONS.values() if mid not in app.OBSERVED_INPUTS]
    assert len(others) == 18 and all(app._observed_editor(mid) is None for mid in others)
    # every entry the card offers is an input of the catalog's observed variant
    for mid, variant in ((CHAN, "observed-channel-adjustment"), (BANK, "observed-bank-condition")):
        method = next(v for v in screening_methods.method_for(mid)["variants"] if v["methodKey"] == variant)
        assert set(app.OBSERVED_INPUTS[mid]) <= {i["key"] for i in method["inputs"]}


def test_the_channel_block_shows_the_catalogs_stage_wording():
    html = str(app._observed_editor(CHAN))
    assert "Observed channel condition" in html and "(optional)" in html
    criteria = screening_methods.method_for(CHAN)["fieldReference"]["criteria"]
    for rating in ("Good", "Fair", "Poor"):
        assert f"{rating}: {criteria[rating]}" in html
    assert "Stage I/VI" in html and "Stages III-V" in html            # what the ratings mean
    options = re.findall(r'<option[^>]*value="([^"]*)"', html)
    assert options == ["", "Good", "Fair", "Poor"] and "No observation" in html
    assert 'data-key="stageClass"' in html and 'data-key="indicators"' in html
    assert html.count(f'data-mid="{CHAN}"') == 2 and html.count("easi-obs-in") == 2
    assert "Both entries are needed" in html and "—" not in html


def test_the_bank_block_shows_both_percentages_and_the_class_limits():
    html = str(app._observed_editor(BANK))
    assert "Observed bank condition" in html and "the worse one governs" in html
    assert 'data-key="erodingBankPct"' in html and 'data-key="armoredBankPct"' in html
    assert html.count('type="number"') == 2 and 'min="0"' in html and 'max="100"' in html
    criteria = screening_methods.method_for(BANK)["fieldReference"]["criteria"]
    for rating in ("Good", "Fair", "Poor"):
        assert criteria[rating].replace("<", "&lt;").replace(">", "&gt;") in html
    assert "—" not in html


def test_the_block_is_seeded_from_what_was_entered():
    html = str(app._observed_editor(CHAN, {"stageClass": "Fair", "indicators": 'headcut <b>up</b> & "bars"'}))
    assert re.search(r'<option[^>]*value="Fair"[^>]*selected|<option[^>]*selected[^>]*value="Fair"', html)
    assert "headcut &lt;b&gt;up&lt;/b&gt; &amp;" in html and "<b>up</b>" not in html      # escaped
    html = str(app._observed_editor(BANK, {"erodingBankPct": 30.0, "armoredBankPct": 12.5}))
    assert 'value="30"' in html and 'value="12.5"' in html


def test_posted_values_are_cleaned():
    value = app._observed_value
    assert [value(CHAN, "stageClass", v) for v in ("Good", "Fair", "Poor", "good", "", None, 5)] == \
        ["Good", "Fair", "Poor", None, None, None, None]
    assert value(CHAN, "indicators", "  headcut   upstream \n bars ") == "headcut upstream bars"
    assert value(CHAN, "indicators", "   ") is None and len(value(CHAN, "indicators", "x" * 900)) == 500
    assert [value(BANK, "erodingBankPct", v) for v in ("30", 30, 0, "100", "100.5", "-1", "x", "nan", "inf", None)] == \
        [30.0, 30.0, 0.0, 100.0, None, None, None, None, None, None]
    # an entry of another metric, or a key the metric does not have, is ignored
    assert value(cc.FUNCTIONS["m03"], "stageClass", "Good") is None
    assert value(CHAN, "erodingBankPct", 30) is None and value(BANK, "stageClass", "Good") is None


# --------------------------------------------------------------------------- #
# the scoring order, and what the card says about it
# --------------------------------------------------------------------------- #
def test_a_complete_observation_governs_and_an_incomplete_one_does_nothing():
    base = _rows(_scored())
    rows = _rows(_scored({CHAN: {"stageClass": "Poor", "indicators": "headcut upstream, 2026-09 visit"},
                          BANK: {"erodingBankPct": 70.0, "armoredBankPct": 10.0}}))
    assert (rows[CHAN]["status"], rows[CHAN]["rating"], rows[CHAN]["confidence"]) == ("observed", "Poor", "H")
    assert (rows[BANK]["status"], rows[BANK]["rating"]) == ("observed", "Poor")            # 70% eroding governs
    assert rows[CHAN]["proxyResult"]["rating"] == base[CHAN]["rating"]                     # the proxy is kept
    # a class without its indicators, or one percentage alone, leaves the automatic rating
    rows = _rows(_scored({CHAN: {"stageClass": "Poor"}, BANK: {"erodingBankPct": 70.0}}))
    assert rows[CHAN]["status"] != "observed" and rows[CHAN]["rating"] == base[CHAN]["rating"]
    assert rows[BANK]["status"] != "observed" and rows[BANK]["rating"] == base[BANK]["rating"]


def test_an_observation_outranks_the_rating_select_and_the_card_says_so():
    rep = _scored({CHAN: {"stageClass": "Poor", "indicators": "headcut upstream"}}, ratings={CHAN: "Good"})
    row = _rows(rep)[CHAN]
    assert (row["status"], row["rating"]) == ("observed", "Poor")
    select = str(app._rate_select(CHAN, row))
    assert "disabled" in select and "Rated from your observation" in select
    assert "disabled" not in str(app._rate_select(CHAN, _rows(_scored())[CHAN]))
    proxy = row["proxyResult"]["rating"]
    assert app._observed_hint(CHAN, {"stageClass": "Poor", "indicators": "x"}, row) == \
        f"Rated from your observation. The automatic rating was {proxy}."


def test_the_hint_names_what_is_missing():
    row = _rows(_scored())[CHAN]
    assert app._observed_hint(CHAN, {}, row) is None
    assert app._observed_hint(CHAN, {"stageClass": "Poor"}, row).startswith("Add the observed indicators.")
    assert app._observed_hint(CHAN, {"indicators": "bars"}, row).startswith("Pick the observed class.")
    assert app._observed_hint(BANK, {"armoredBankPct": 5.0}, _rows(_scored())[BANK]).startswith(
        "Enter both percentages.")
    assert app._observed_hint(cc.FUNCTIONS["m03"], {"stageClass": "Poor"}, row) is None


def test_the_server_applies_the_observations_after_the_rescore():
    scored = SRC.split("def scored():", 1)[1].split("@reactive.calc", 1)[0]
    assert scored.index("assessment.rescore(") < scored.index("assessment.apply_observed_evidence(sc, observed)")
    assert scored.index('row["status"] = "xs-derived"') < scored.index("apply_observed_evidence")
    handler = SRC.split("def _apply_observed():", 1)[1].split("@reactive", 1)[0]
    assert "@reactive.event(input.observed_set)\n    def _apply_observed():" in SRC
    assert "if key not in OBSERVED_INPUTS.get(mid, ()):" in handler and "_observed_value(mid, key" in handler
    # a fresh screening and a new analysis both start without observations
    assert SRC.count("_overrides.set({}); _notes.set({}); _observed.set({})") == 2
    # the card: the block sits in the static skeleton, the hint in the live slot
    panel = SRC.split("def fn_panel():", 1)[1].split("def _cur_row(", 1)[0]
    assert "_observed_editor(mid, observed0)" in panel and "observed0 = (_observed() or {}).get(mid)" in panel
    live = SRC.split("def fn_metric_live():", 1)[1].split("def _cur_method", 1)[0]
    assert '"observed" if is_observed else "desktop"' in live and "_observed_hint(" in live
    assert "and not is_observed else None" in live                 # no restore chip under an observation


def test_the_client_posts_each_entry():
    assert 'send("observed_set", { mid: el.getAttribute("data-mid"), key: el.getAttribute("data-key"),' in JS
    assert 'document.addEventListener("change", function (e) {\n    if (!isObserved(e.target)) return;' in JS
    assert 'if (!isObserved(el) || el.tagName === "SELECT") return;' in JS          # typed text is debounced
    assert 'src="worksheet.js?v=10"' in SRC and 'href="styles.css?v=61"' in SRC
    for rule in (".easi-obs {", ".easi-obs-in {", ".easi-obs-hint.applied", ".easi-rate-sel[disabled]"):
        assert rule in CSS, rule


# --------------------------------------------------------------------------- #
# into the exports and the completed workbook
# --------------------------------------------------------------------------- #
def test_observations_reach_the_completed_workbook_and_the_list():
    observed = {CHAN: {"stageClass": "Poor", "indicators": "headcut upstream, 2026-09 visit"},
                BANK: {"erodingBankPct": 70.0, "armoredBankPct": 10.0}}
    res = {"delineation": dict(DELINEATION), "report": _scored(observed)}
    entries, _ = calculator.entries_from_result(res)
    assert entries["ov_stageClass"] == "Poor" and entries["ov_indicators"] == "headcut upstream, 2026-09 visit"
    assert entries["ov_erodingBankPct"] == 70.0 and entries["ov_armoredBankPct"] == 10.0
    assert "ov_m09_rating" not in entries and "ov_m10_rating" not in entries          # observed, not overridden
    assert entries["in_bhr"] == 1.1 and entries["in_er"] == 2.5                       # still from the other rows
    ws = openpyxl.load_workbook(io.BytesIO(calculator.build_filled(res)))["EASI Score"]
    cells = calculator.entry_cells()
    assert ws[cells["ov_stageClass"]].value == "Poor" and ws[cells["ov_erodingBankPct"]].value == 70
    rows = {r["metricId"]: r for r in report.desktop_metric_rows(res)}
    assert rows[CHAN]["inputs"] == [{"label": "Documented adjustment class", "value": "Poor"},
                                    {"label": "Observed indicators", "value": "headcut upstream, 2026-09 visit"}]
    assert [i["value"] for i in rows[BANK]["inputs"]] == ["70%", "10%"]
    assert report.build_pdf({**res, "watershed_geojson": None, "reach_geojson": None})[:4] == b"%PDF"


def test_the_workbook_scores_the_observations_like_the_application():
    pytest.importorskip("formulas")
    from calculator_eval import FormulasBackend
    from test_calculator_parity import OUTPUT_NAMES, compare
    observed = {CHAN: {"stageClass": "Fair", "indicators": "widening, bars forming"},
                BANK: {"erodingBankPct": 70.0, "armoredBankPct": 10.0}}
    rep = _scored(observed, ratings={cc.FUNCTIONS["m03"]: "Poor", CHAN: "Good"})
    entries, _ = calculator.entries_from_result({"delineation": dict(DELINEATION), "report": rep})
    got = FormulasBackend(cc.WORKBOOK).evaluate({k: v for k, v in entries.items() if not k.startswith("site_")},
                                                OUTPUT_NAMES)
    assert not compare({"expected": cc.expected_from(rep)}, got)
    assert (got["m09_rating"], got["m09_route"]) == ("Fair", "observed channel class")
    assert (got["m10_rating"], got["m10_route"]) == ("Poor", "observed bank condition")
    assert (got["m03_rating"], got["m03_route"]) == ("Poor", "override score")


# --------------------------------------------------------------------------- #
# the monthly flow helper
# --------------------------------------------------------------------------- #
def test_monthly_flows_are_a_full_year_or_nothing():
    erom = {f"qe_{m:02d}": 10.0 + m for m in range(1, 13)}
    assert calculator.monthly_flows({**erom, "qe_ma": 16.5}) == [10.0 + m for m in range(1, 13)]
    for broken in (None, {}, [], {**erom, "qe_07": None}, {**erom, "qe_07": "x"}, {**erom, "qe_07": float("nan")},
                   {k: v for k, v in erom.items() if k != "qe_12"}):
        assert calculator.monthly_flows(broken) is None


def test_the_completed_workbook_carries_the_twelve_flows_behind_the_variability():
    flows = [100.0 + 5 * m for m in range(12)]
    res = {"delineation": dict(DELINEATION), "report": _scored(), "eromMonthly": flows}
    entries, _ = calculator.entries_from_result(res)
    assert [entries[f"in_m{m:02d}"] for m in range(1, 13)] == flows
    ws = openpyxl.load_workbook(io.BytesIO(calculator.build_filled(res)))["EASI Score"]
    cells = calculator.entry_cells()
    assert [ws[cells[f"in_m{m:02d}"]].value for m in range(1, 13)] == flows
    # without a full year, or without the variability entry itself, the helper stays blank
    for result in ({**res, "eromMonthly": flows[:11]}, {**res, "eromMonthly": None},
                   {**res, "eromMonthly": flows[:11] + [float("nan")]}):
        assert not [k for k in calculator.entries_from_result(result)[0] if re.fullmatch(r"in_m\d\d", k)]
    # the application and the batch engine both keep the flows on the result
    assert 'merged["eromMonthly"] = calculator.monthly_flows((d.get("ctx_inputs") or {}).get("erom"))' in SRC
    api = (Path(app.__file__).parent / "easi" / "batch" / "api.py").read_text(encoding="utf-8")
    assert '"eromMonthly": calculator.monthly_flows(ctx_inputs.get("erom")),' in api
