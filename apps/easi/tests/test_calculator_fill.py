"""The completed calculator: a screening typed into the workbook's entry cells.

``easi.calculator.entries_from_result`` reads the entries from a scored report
alone (the live application has no stored evidence record), and
``build_filled`` writes them into the committed workbook at the zip level. Three
things are proven here on the parity case set of ``tests/calculator_cases.py``:

1. the report-driven entries never disagree with the record-driven oracle
   (``calculator_cases.entries_from``): where both hold a value it is the same
   value, and the runtime never holds one the oracle lacks. It may be blank where
   the oracle is not, because only the method that rated a row traces its inputs;
2. those blanks are neutral: the workbook, evaluated with the runtime entries,
   still gives the engine's ratings, scores, indices and rollup. A stratified
   sample runs by default, the full sweep with ``EASI_FILL_PARITY_FULL=1``;
3. the completed workbook is the blank plus the intended cells and nothing else.
"""
from __future__ import annotations

import datetime as dt
import io
import os
import re
import zipfile

import openpyxl
import pytest

import calculator_cases as cc
from calculator_eval import FormulasBackend
from easi import calculator
from test_calculator_parity import OUTPUT_NAMES, compare

TODAY = dt.date(2026, 9, 18)
CHAN = "channel-evolution-channel-evolution-stage-and-trends"
DELINEATION = {"comid": 9327042, "gnis_name": "Mink Brook", "snapped_lat": 43.6858, "snapped_lon": -72.2367}


def result_for(case: dict, **extra) -> dict:
    return {"delineation": dict(DELINEATION), "report": cc.score_case(case), **extra}


def present(entries: dict) -> dict:
    """Entries that hold something, without the site block (the oracle has none)."""
    return {k: v for k, v in entries.items() if v not in (None, "") and not k.startswith("site_")}


def same(a, b) -> bool:
    if isinstance(a, str) or isinstance(b, str):
        return str(a).strip().upper() == str(b).strip().upper()
    return abs(float(a) - float(b)) <= 1e-12


@pytest.fixture(scope="module")
def fixture() -> dict:
    if not cc.FIXTURE.exists():
        pytest.skip("run test_calculator_parity.py with EASI_WRITE_GOLDEN=1 first")
    return cc.load_fixture()


@pytest.fixture(scope="module")
def scored(fixture) -> list[tuple[dict, dict, dict]]:
    """(case, runtime entries, oracle entries) for every parity case."""
    out = []
    for case in fixture["cases"]:
        runtime, _ = calculator.entries_from_result(result_for(case))
        out.append((case, present(runtime), present(case["entries"])))
    return out


# --------------------------------------------------------------------------- #
# 1. the report-driven entries against the record-driven oracle
# --------------------------------------------------------------------------- #
def test_runtime_entries_never_disagree_with_the_oracle(scored):
    wrong, invented = [], []
    for case, runtime, oracle in scored:
        for name, value in runtime.items():
            if name not in oracle:
                invented.append(f"{case['id']}: {name}={value!r}")
            elif not same(value, oracle[name]):
                wrong.append(f"{case['id']}: {name} {value!r} != {oracle[name]!r}")
    assert not wrong, wrong[:10]
    assert not invented, invented[:10]


def test_every_entry_name_is_a_cell_of_the_workbook(scored):
    cells = calculator.entry_cells()
    assert len(cells) == 61 + 20                      # the 1.0 workbook: an Override Score per function
    for _case, runtime, oracle in scored:
        assert set(runtime) <= set(cells) and set(oracle) <= set(cells)


def test_the_feature_code_is_a_number(scored):
    # as text the workbook's canal test never matches and Channel evolution loses its route
    for case, runtime, _oracle in scored:
        if "ctx_fcode" in runtime:
            assert isinstance(runtime["ctx_fcode"], int), case["id"]
    canal = [runtime for case, runtime, _ in scored if case["id"].startswith("canal:33600")]
    assert canal and all(r["ctx_fcode"] == 33600 for r in canal)


# --------------------------------------------------------------------------- #
# 2. the blanks are neutral: the workbook still scores like the engine
# --------------------------------------------------------------------------- #
def _sample(scored):
    """Up to three cases per (group, which entries differ), plus the cases whose route
    rests on a context entry (canals) or on the order of expert inputs (override scores)."""
    if os.environ.get("EASI_FILL_PARITY_FULL") == "1":
        return list(scored)
    taken: dict = {}
    picked = []
    for item in scored:
        case, runtime, oracle = item
        signature = (case["group"], frozenset(set(oracle) ^ set(runtime)))
        forced = case["id"].startswith(("canal:", "override-score:over-", "override-score:under-",
                                        "override-score:beside-", "override-score:all:", "route:population:"))
        if forced or taken.get(signature, 0) < 3:
            taken[signature] = taken.get(signature, 0) + 1
            picked.append(item)
    return picked


def test_workbook_scores_like_the_engine_from_the_runtime_entries(scored):
    pytest.importorskip("formulas")
    backend = FormulasBackend(cc.WORKBOOK)
    sample = _sample(scored)
    assert len(sample) >= 40 and {case["group"] for case, *_ in sample} == {c["group"] for c, *_ in scored}
    failures = []
    for case, runtime, _oracle in sample:
        problems = compare(case, backend.evaluate(runtime, OUTPUT_NAMES))
        if problems:
            failures.append(f"{case['id']}: " + "; ".join(problems[:6]))
    assert not failures, f"{len(failures)} of {len(sample)} cases differ:\n" + "\n".join(failures[:25])


# --------------------------------------------------------------------------- #
# 3. the completed workbook is the blank plus the intended cells
# --------------------------------------------------------------------------- #
_CELL = re.compile(r'<c r="([A-Z]+[0-9]+)"[^>]*?(?:/>|>.*?</c>)', re.S)


def _cells_of(sheet_xml: str) -> dict:
    return {m.group(1): m.group(0) for m in _CELL.finditer(sheet_xml)}


def _assessed_result() -> dict:
    case = {"record": {}, "observed": None, "ratings": {cc.FUNCTIONS["m03"]: "Poor", CHAN: "Fair"}}
    res = result_for(case)
    res["report"]["metricRows"] = [
        {**row, "userNote": "=HYPERLINK(\"x\") culvert <b>outfall</b> & riprap"
         if row["metricId"] == cc.FUNCTIONS["m03"] else ""} for row in res["report"]["metricRows"]]
    return res


def test_only_the_score_sheet_changes_and_only_the_intended_cells():
    res = _assessed_result()
    blank, filled = calculator.blank_bytes(), calculator.build_filled(res, today=TODAY)
    assert filled == calculator.build_filled(res, today=TODAY)          # deterministic for a date
    with zipfile.ZipFile(io.BytesIO(blank)) as a, zipfile.ZipFile(io.BytesIO(filled)) as b:
        assert a.namelist() == b.namelist()
        for ia, ib in zip(a.infolist(), b.infolist()):
            assert (ia.date_time, ia.compress_type, ia.create_system, ia.external_attr) == \
                   (ib.date_time, ib.compress_type, ib.create_system, ib.external_attr), ia.filename
        changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
        assert len(changed) == 1
        before, after = (_cells_of(z.read(changed[0]).decode("utf-8")) for z in (a, b))
        book = a.read("xl/workbook.xml").decode("utf-8")
    assert 'fullCalcOnLoad="1"' in book, "the blank must recalculate on open: a completed copy carries no results"
    entries, _ = calculator.entries_from_result(res)
    cells = calculator.entry_cells()
    intended = {cells[n] for n in present(entries)} | {cells[n] for n in entries if n.startswith("site_")}
    intended |= {cells["site_date"], cells["site_notes"]}
    assert set(before) == set(after)
    assert {ref for ref in before if before[ref] != after[ref]} == intended
    # the template's other parts are untouched, so the charts, validations and formats survive
    wb = openpyxl.load_workbook(io.BytesIO(filled))
    ws = wb["EASI Score"]
    assert len(ws._charts) == 2 and len(ws.data_validations.dataValidation) == 56 + 20
    assert ws.protection.sheet and ws.merged_cells.ranges


def test_round_trip_by_defined_name_and_the_workbook_calculates_as_filled(tmp_path):
    pytest.importorskip("formulas")
    res = _assessed_result()
    entries, _ = calculator.entries_from_result(res)
    path = tmp_path / calculator.filled_filename(res)
    path.write_bytes(calculator.build_filled(res, today=TODAY))
    assert path.name == "easi-calculator-comid-9327042.xlsx"
    ws = openpyxl.load_workbook(path)["EASI Score"]
    cells = calculator.entry_cells()
    for name, value in present(entries).items():
        got = ws[cells[name]].value
        assert same(got, value) and isinstance(got, str) == isinstance(value, str), (name, got, value)
    assert ws[cells["ctx_fcode"]].data_type == "n" and ws[cells["ov_m03_rating"]].value == "Poor"
    assert ws[cells["site_name"]].value == "Mink Brook, COMID 9327042"
    assert ws[cells["site_coords"]].value == "43.68580, -72.23670"
    assert ws[cells["site_date"]].value == dt.datetime(2026, 9, 18)
    assert ws[cells["site_assessor"]].value is None
    notes = ws[cells["site_notes"]]
    assert notes.data_type == "s" and notes.value.startswith(calculator.NOTES_LABEL + "\n")
    # the assessor's text comes back as typed: markup, an ampersand and a would-be formula
    # are all plain text
    assert '=HYPERLINK("x") culvert <b>outfall</b> & riprap' in notes.value
    assert "Override scores: Reach inflow Poor (computed " in notes.value
    # the file itself calculates to the application's results, with nothing set from outside
    backend = FormulasBackend(path)
    wanted = [backend.key(name) for name in OUTPUT_NAMES]
    solution = backend.model.calculate(outputs=wanted)
    from calculator_eval import _plain
    got = {name: _plain(getattr(solution.get(key), "value", solution.get(key)))
           for name, key in zip(OUTPUT_NAMES, wanted)}
    case = {"expected": cc.expected_from(res["report"])}
    assert not compare(case, got), compare(case, got)[:6]
    assert got["m03_rating"] == "Poor" and got["m03_route"] == "override score"


@pytest.mark.parametrize("empty", ['<c r="J8" s="14" t="n"></c>', '<c r="J8" s="14"/>', '<c r="J8" s="14" />'])
def test_both_serializations_of_an_empty_cell_fill(empty):
    # openpyxl writes <c ...></c> with lxml installed and <c .../> without it
    sheet = f'<sheetData><row r="8">{empty}<c r="K8" s="3" t="inlineStr"><is><t>%</t></is></c></row></sheetData>'
    assert '<c r="J8" s="14" t="n"><v>12.5</v></c>' in calculator._fill_cells(sheet, {"J8": 12.5})
    text = calculator._fill_cells(sheet, {"J8": "4A"})
    assert '<c r="J8" s="14" t="inlineStr"><is><t xml:space="preserve">4A</t></is></c>' in text
    assert text.count("<c ") == 2 and "K8" in text


def test_text_is_escaped_for_xml_and_for_excel():
    sheet = '<row r="2"><c r="B2" s="14" t="n"></c></row>'
    text = calculator._fill_cells(sheet, {"B2": "a < b & c > d _x000D_ \x07bell"})
    # Excel decodes _xHHHH_ on read, so a literal one is escaped; control characters are dropped
    assert "a &lt; b &amp; c &gt; d _x005F_x000D_ bell" in text


def test_a_moved_or_occupied_cell_fails_loudly():
    sheet = '<row r="8"><c r="J8" s="14" t="n"></c><c r="J9" s="14" t="n"><v>3</v></c></row>'
    with pytest.raises(ValueError, match="J99"):
        calculator._fill_cells(sheet, {"J99": 1})
    with pytest.raises(ValueError, match="J9 is not empty"):
        calculator._fill_cells(sheet, {"J9": 1})


def test_numbers_are_written_the_way_excel_reads_them():
    np = pytest.importorskip("numpy")
    assert calculator._number_text(33600.0) == "33600" and calculator._number_text(np.int64(3)) == "3"
    assert calculator._number_text(np.float64(0.621233)) == "0.621233"      # not "np.float64(...)"
    assert calculator._number_text(0.1 + 0.2) == repr(0.1 + 0.2)
    for value in (float("nan"), float("inf"), True, "5", None):
        assert calculator._number_text(value) is None
    assert calculator._excel_serial(dt.date(2026, 9, 18)) == 46283


def test_the_aliases_are_the_generators():
    assert calculator.ALIASES == cc.bc.ALIASES


# --------------------------------------------------------------------------- #
# what the assessor did on the Assessment page
# --------------------------------------------------------------------------- #
def _rows(res):
    return {row["metricId"]: row for row in res["report"]["metricRows"]}


def test_override_scores_and_notes_are_carried():
    res = _assessed_result()
    entries, disclosures = calculator.entries_from_result(res)
    assert entries["ov_m03_rating"] == "Poor" and entries["ov_m09_rating"] == "Fair"
    assert [n for n in entries if re.fullmatch(r"ov_m\d\d_rating", n)] == ["ov_m03_rating", "ov_m09_rating"]
    assert disclosures[0].startswith("Override scores: Reach inflow Poor (computed ")
    # an overridden row keeps the trace of the rating it replaced, so its values still fill
    assert entries["in_roadDensity"] == 1.2
    text = calculator.notes_text(res, disclosures, TODAY)
    assert text.splitlines()[1] == ("Completed by the EASI web application on 2026-09-18 from the screening of "
                                    "Mink Brook, COMID 9327042.")
    assert "Assessor notes. Reach inflow: =HYPERLINK" in text and "—" not in text


def test_every_function_carries_its_override_including_the_three_the_registry_does_not_flag():
    # the Assessment page offers the rating select on all 20 function cards (the registry's
    # ``overrideable`` flag is not enforced there), so Catchment hydrology, Surface water
    # storage and Watershed connectivity can be overridden and the workbook must carry it
    from easi import config
    unflagged = [mkey for mkey, mid in cc.FUNCTIONS.items() if not config.METRIC_REGISTRY[mid].get("overrideable")]
    assert unflagged == ["m01", "m02", "m20"]
    ratings = {cc.FUNCTIONS[mkey]: "Poor" for mkey in unflagged}
    res = result_for({"record": {}, "observed": None, "ratings": ratings})
    entries, disclosures = calculator.entries_from_result(res)
    assert {name: entries[name] for name in entries if name.startswith("ov_m")} == {
        "ov_m01_rating": "Poor", "ov_m02_rating": "Poor", "ov_m20_rating": "Poor"}
    assert not any("no Override Score entry" in line for line in disclosures)
    assert disclosures[0] == ("Override scores: Catchment hydrology Poor (computed Good); Surface water storage "
                              "Poor (computed Fair); Watershed connectivity Poor (computed Good).")
    ws = openpyxl.load_workbook(io.BytesIO(calculator.build_filled(res, today=TODAY)))["EASI Score"]
    cells = calculator.entry_cells()
    assert [ws[cells[f"ov_{mkey}_rating"]].value for mkey in cc.FUNCTIONS] == \
        ["Poor" if mkey in unflagged else None for mkey in cc.FUNCTIONS]


def test_a_cross_section_edit_fills_its_values_and_no_override_score():
    # the app re-rates the geometry metrics from the edited section and relabels the rows
    # "xs-derived": their traced BHR and ER reproduce the rating, so no rating is typed in.
    # A manual pick afterwards hands that row back its base trace (the old BHR), which
    # must not win the shared cell.
    res = result_for({"record": {}, "observed": None})
    rows = _rows(res)
    m06, m10 = cc.FUNCTIONS["m06"], cc.FUNCTIONS["m10"]

    def with_bhr(row, value, status):
        trace = {**row["scoring"], "inputs": [{**i, "value": value} if i["key"] == "bhr" else i
                                              for i in row["scoring"]["inputs"]]}
        return {**row, "scoring": trace, "status": status}

    rows[m10] = with_bhr(rows[m10], 1.62, "xs-derived")
    rows[m06] = {**rows[m06], "status": "override", "rating": "Good",
                 "effectiveOverride": {"rating": "Good", "generatedRating": "Good"}}
    res["report"]["metricRows"] = list(rows.values())
    entries, disclosures = calculator.entries_from_result(res)
    assert entries["in_bhr"] == 1.62                      # the edited section, not m06's base 1.1
    assert "ov_m10_rating" not in entries and entries["ov_m06_rating"] == "Good"
    assert any("Bank-height ratio" in line and "1.62" in line and "1.1" in line for line in disclosures)


def test_a_value_the_application_could_not_use_everywhere_is_disclosed():
    # the NLCD outage fallback reaches Catchment hydrology only: the one impervious cell then
    # rates Light and thermal regime in the workbook although the application left it unrated
    res = result_for({"record": {}, "observed": None})
    rows = _rows(res)
    m13 = "light-and-thermal-regime-stream-temperature"
    trace = {**rows[m13]["scoring"], "inputs": [{**i, "value": None} if i["key"] == "impervious" else i
                                                for i in rows[m13]["scoring"]["inputs"]]}
    rows[m13] = {**rows[m13], "scoring": trace, "rating": None, "status": "unavailable"}
    res["report"]["metricRows"] = list(rows.values())
    entries, disclosures = calculator.entries_from_result(res)
    assert entries["in_impervious"] == 8.0
    assert any(line.startswith("Not rated by the application: Light & thermal regime.") for line in disclosures)


def test_unknown_inputs_and_broken_reports_never_raise():
    res = result_for({"record": {}, "observed": None})
    row = res["report"]["metricRows"][0]
    row["scoring"] = {**row["scoring"], "inputs": row["scoring"]["inputs"] + [
        {"key": "legacyOnly", "label": "A legacy input", "value": 4.2, "contextOnly": False}]}
    entries, disclosures = calculator.entries_from_result(res)
    assert "in_legacyOnly" not in entries
    assert "Not entered (no cell in this workbook): A legacy input." in disclosures
    for broken in (None, {}, {"report": {}}, {"report": {"metricRows": [{"metricId": "nope"}, {}]}},
                   {"report": {"metricRows": [{**row, "scoring": None}]}, "delineation": None}):
        entries, _ = calculator.entries_from_result(broken)
        assert calculator.build_filled(broken, today=TODAY)[:2] == b"PK"


def test_the_slope_class_is_forced_only_when_the_workbook_would_derive_another():
    res = result_for({"record": {"slope": 0.004}, "observed": None})
    assert res["report"]["strata"]["slope_class"] == "lt_0.5"
    assert "ctx_slope_override" not in calculator.entries_from_result(res)[0]
    res["report"]["strata"] = {**res["report"]["strata"], "slope_class": "ge_2"}
    assert calculator.entries_from_result(res)[0]["ctx_slope_override"] == "ge_2"
    res["report"]["strata"] = {**res["report"]["strata"], "slope_class": None}
    assert calculator.entries_from_result(res)[0]["ctx_slope_override"] == "national"


def test_the_notes_stay_inside_the_box():
    res = _assessed_result()
    res["report"]["metricRows"] = [{**row, "userNote": "long note " * 60} for row in res["report"]["metricRows"]]
    entries, disclosures = calculator.entries_from_result(res)
    text = calculator.notes_text(res, disclosures, TODAY)
    assert len(text) <= calculator.NOTES_LIMIT and text.endswith(calculator.NOTES_CONTINUED)


def test_the_site_block_and_the_download_name():
    routed = {"delineation": {**DELINEATION, "gnis_name": "Mink Brook"},
              "siteAnchor": {"anchorKind": "hrSurrogate",
                             "clickedStream": {"gnisName": "", "nhdplusId": 10000900012345}}, "report": {}}
    site = calculator.site_identity(routed)
    assert site["label"] == "Unnamed stream, NHDPlusID 10000900012345" and site["comid"] == "9327042"
    assert calculator.filled_filename(routed) == "easi-calculator-nhdplusid-10000900012345.xlsx"
    assert calculator.filled_filename({"delineation": DELINEATION}) == "easi-calculator-comid-9327042.xlsx"
    assert calculator.filled_filename({"delineation": {"snapped_lat": -12.5, "snapped_lon": 30.25}}) == \
        "easi-calculator-s12.50000-e30.25000.xlsx"
    assert calculator.filled_filename({}) == "easi-calculator.xlsx"
