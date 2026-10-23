"""Generate the EASI Excel calculator from the application's own scoring definitions.

The workbook is never edited by hand. Every band, curve, regional threshold,
lookup, weight and anchor is read from the same files the application scores
with (``data/screening-methods.json``, ``data/reference-curves.json``,
``data/cwa-mapping.json``, ``data/easi-metrics.json``, ``data/scoring-identity.json``)
and written as worksheet formulas, so the calculator and the web application
implement one methodology. ``tests/test_calculator_parity.py`` proves that on a
retained case set scored by the application's own engine.

Sheets: Instructions (laid out as SFARI's), Scenario (the worksheet the user fills,
laid out like the SFARI calculator: outcomes, functional categories, functions and
metrics side by side, orange entry cells, a scoring summary table, legends and two bar
charts), Metrics (every rule and route beside the Good, Fair and Poor ranges of each
rated value, read only), Results (the rollup, hidden engine sheet), Reference (the
reference curves drawn, then every table the formulas read) and ChartData (hidden chart
series). The workbook's identity (scoring method digest, package digest of the method
files, generator sha256) sits in its document properties, out of sight (owner,
2026-10-23: no Metadata sheet). The apps serve it as the template every EASI workbook
is assembled from (``easi/workbook.py``), named by its date.

Formulas use direct cell references and a small function vocabulary (IF AND OR
NOT MIN MAX ROUND INDEX MATCH IFERROR ISNUMBER COUNT SUM SUMPRODUCT AVERAGE
STDEV.P PRODUCT UPPER TRIM SUBSTITUTE LEN NA) so Excel 2016 and later,
LibreOffice and the Python ``formulas`` evaluator all agree. Defined names are
added for every entry cell and reference column for readers and for the parity
harness, but no formula depends on them.

Route logic that lives in the adapters rather than the catalog (fallbacks and
observed overrides) is written out in ROUTES below and mirrors
``easi/metrics/*.py`` and ``easi/assessment.py``:

    nutrients          region known and at least one analyte, else min(CHEM) fallback
    impairment         conclusive ATTAINS category, else min(CHEM) fallback
    population support prg_bmmi0809 in [0, 1], else the ICI/IWI products (all twelve)
    channel evolution  observed class with a note > canal FCODE > worst of BHR and ER
    bank condition     observed erosion and armoring (both) > BHR proxy

Every function ends with one more entry, the Override Score: the rating select
the Assessment page shows on all 20 function cards (the registry's
``overrideable`` flag is not enforced there, so neither is it here). It replaces
the computed rating exactly as ``assessment.rescore`` does. On the two functions
with observed entries the order is the one the engine composes,
``apply_observed_evidence(rescore(...))``: complete observed entries, then the
override, then the automatic result. The row is set apart from the function's
metrics (a double rule, a band of its own, a taller row), because it is a
user-defined input and not one of the desktop quantities.

Run:  .venv/Scripts/python.exe scripts/build_calculator.py [--out PATH] [--check]
"""
from __future__ import annotations

import argparse
import contextlib
import datetime as _dt
import hashlib
import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, Reference, ScatterChart, Series  # noqa: E402
from openpyxl.chart.data_source import NumFmt  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.chart.marker import Marker  # noqa: E402
from openpyxl.chart.shapes import GraphicalProperties  # noqa: E402
from openpyxl.drawing.line import LineProperties  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.packaging.custom import StringProperty  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from easi import calculator as served  # noqa: E402
from easi import config, scoring  # noqa: E402
from easi import screening_methods as sm  # noqa: E402
from easi.national import method_version  # noqa: E402

TEMPLATE_VERSION = served.TEMPLATE_VERSION   # one source: the module that serves the workbook
TEMPLATE_DATE = TEMPLATE_VERSION              # the calculator is named by its date (owner, 2026-10-23)
MODEL_VERSION = "EASI v1.0"                   # as SFARI's "SFARI v1.0"
OUT_DIR = os.path.join(ROOT, "www", "calculator")
OUT_NAME = served.TEMPLATE_PATH.name
PASSWORD = "easi"      # a guard rail against accidental edits, not a secret
FIXED_STAMP = _dt.datetime.strptime(TEMPLATE_DATE, "%Y-%m-%d")
SCORE = "Scenario"     # the worksheet the user fills
#: The identity a workbook carries in its document properties (File, Info, Properties):
#: what ``check`` regenerates with and what StreamCurves verifies before a calculator rides in
#: a method package. Out of sight, so no sheet of the workbook holds it (owner, 2026-10-23).
STAMP_METHOD, STAMP_PACKAGE, STAMP_GENERATOR = "Scoring method digest", "Package digest", "Generator sha256"

#: The namespace prefixes openpyxl registers when it is imported (``openpyxl/xml/functions.py``).
#: lxml and ElementTree keep one prefix registry per process, and a library imported later may
#: register another prefix for the same namespace: owslib, which pynhd imports, registers ``dct``
#: for Dublin Core terms, and ``docProps/core.xml`` would then say ``dct:created`` for
#: ``dcterms:created`` (the same XML, other bytes). ``Builder.build`` puts these back for the save.
OPENPYXL_NAMESPACES = {
    "dcterms": "http://purl.org/dc/terms/",
    "dcmitype": "http://purl.org/dc/dcmitype/",
    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "vt": "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "cdr": "http://schemas.openxmlformats.org/drawingml/2006/chartDrawing",
    "xml": "http://www.w3.org/XML/1998/namespace",
    "cust": "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties",
}


def registered_prefix(uri: str) -> str | None:
    """The prefix the XML library openpyxl writes with gives a new element in ``uri`` now
    (None when no prefix is registered for it)."""
    from openpyxl import LXML
    if LXML:
        from lxml import etree
        prefix = etree.Element("{%s}probe" % uri).prefix
        return None if prefix is None or re.fullmatch(r"ns\d+", prefix) else prefix
    import xml.etree.ElementTree as ElementTree
    return ElementTree._namespace_map.get(uri)


@contextlib.contextmanager
def openpyxl_namespaces():
    """openpyxl's own prefixes (``OPENPYXL_NAMESPACES``) while the workbook is written, and the
    registry as it was afterwards, so building leaves no trace on the process either."""
    from openpyxl.xml.functions import register_namespace
    before = dict((uri, registered_prefix(uri)) for uri in OPENPYXL_NAMESPACES.values())
    for prefix, uri in OPENPYXL_NAMESPACES.items():
        register_namespace(prefix, uri)
    try:
        yield
    finally:
        for prefix, uri in OPENPYXL_NAMESPACES.items():
            if before[uri] is not None and before[uri] != prefix:
                register_namespace(before[uri], uri)


RATINGS = ("Good", "Fair", "Poor")
INDEX_EDGES = (0.39, 0.69)
COVERAGE_THRESHOLD = 0.70

# Keys that name the same quantity under two catalog names (the adapters read
# the same StreamCat field for both).
ALIASES = {"chemCat": "chemCatchment", "chemWs": "chemWatershed"}

# ---------------------------------------------------------------------------
# Palette: the SFARI calculator's worksheet look (Office theme greys for the
# bands, the accent-2 tint for entries, the same three class colours as its
# legends and charts, the same five functional-category tints).
# ---------------------------------------------------------------------------
C_INPUT = "F8CBAD"       # user entry
C_LINK = "F2F2F2"        # repeats an entry made elsewhere
C_BAND = "F2F2F2"        # band headers and the Physical block labels
C_CHEM = "D9D9D9"        # Chemical block labels
C_BIO = "BFBFBF"         # Biological block labels
C_GOOD = "8FAADC"        # Functioning
C_FAIR = "FFFF99"        # Functioning At-Risk
C_POOR = "FF6969"        # Non-Functioning
C_CALC = "EDEDED"        # computed cells on the detail sheets
CATEGORY_FILL = {"Hydrology": "D9E1F2", "Hydraulics": "B4C6E7", "Geomorphology": "FCE4D6",
                 "Physicochemistry": "FFF2CC", "Biology": "E2EFDA"}
OUTCOME_OF = {"Hydrology": "physical", "Hydraulics": "physical", "Geomorphology": "physical",
              "Physicochemistry": "chemical", "Biology": "biological"}
OUTCOME_FILL = {"physical": C_BAND, "chemical": C_CHEM, "biological": C_BIO}
OUTCOME_LABEL = {"physical": "Physical", "chemical": "Chemical", "biological": "Biological"}
CLASS_LABELS = {"Good": "Functioning", "Fair": "Functioning At-Risk", "Poor": "Non-Functioning"}
#: Metrics columns: the item, its value, the Good / Fair / Poor ranges it is rated against (the
#: criteria beside every rated value, owner 2026-10-23), the rating, the index anchor and a note.
M_ITEM, M_VALUE, M_GOOD, M_FAIR, M_POOR, M_RATING, M_INDEX, M_NOTE = range(1, 9)
#: the ranges' tints: the web application's Good / Fair / Poor band colours (its Scoring method panel)
BIN_FILL = {"Good": "C8D9F2", "Fair": "F5E7A6", "Poor": "F5B5B5"}
CLASS_FILL = {"Good": C_GOOD, "Fair": C_FAIR, "Poor": C_POOR}
#: SFARI's line on the sheets that take no entries, and their tab colour
NO_INPUT = "**DO NOT INPUT VALUES HERE**"
C_TAB_NO_INPUT = "A6A6A6"
#: How each curve set's value reads in a range: decimals and unit (the reference curves' axes)
CURVE_DISPLAY = {"flow-variability": (2, ""), "entrenchment": (2, ""),
                 "corridor-woody": (1, "%"), "corridor-natural": (1, "%")}
CURVE_AXIS = {"flow-variability": "Monthly flow variability (CV of the twelve EROM monthly flows)",
              "entrenchment": "Entrenchment ratio", "corridor-woody": "Woody riparian cover in the 100 m corridor (%)",
              "corridor-natural": "Natural cover in the 100 m corridor (%)"}
CURVE_TITLE = {"flow-variability": "Low flow and baseflow dynamics",
               "entrenchment": "Floodplain connectivity, Channel evolution",
               "corridor-woody": "Light and thermal regime, Habitat provision",
               "corridor-natural": "Carbon processing"}
SLOPE_LABEL = {"lt_0.5": "Slope under 0.5%", "0.5_to_2": "Slope 0.5 to 2%", "ge_2": "Slope 2% or more",
               "national": "National"}
#: Reference's curve table: the knot values from column G, their indices from L, the ranges from Q
CURVE_X, CURVE_BINS = 7, 17
#: the curve charts, two across and two down above the table (rows per chart, size in cm)
CHART_ROWS, CHART_WIDTH_CM, CHART_HEIGHT_CM = 25, 16.5, 12.5


def num_text(x: float, digits: int) -> str:
    """A rounded value as a range shows it (Excel's general format: no trailing zeros)."""
    text = f"{round(float(x), digits):.{digits}f}"
    return text.rstrip("0").rstrip(".") if "." in text else text


def band_bins(bands: list[dict]) -> dict:
    """Rating -> the catalog's own label of its band (the text the application shows)."""
    out: dict = {}
    for b in bands:
        assert b.get("label"), b
        out[b["rating"]] = f"{out[b['rating']]} or {b['label']}" if b["rating"] in out else b["label"]
    return out


def curve_bins(higher_is_better: bool, x39: str, x69: str, unit: str, *, formula: bool = False) -> dict:
    """Rating -> its range on a curve, in the catalog's notation, from the curve's 0.39 and 0.69
    crossings (index 0.69 and above is Good, 0.39 and above Fair). ``x39`` and ``x69`` are text, or
    with ``formula`` expressions, and the range is then an expression joining them with ``&``."""
    lit = q if formula else (lambda s: s)
    tail = [lit(unit)] if unit else []
    if higher_is_better:
        spec = {"Good": [lit("≥"), x69], "Fair": [x39, lit("-<"), x69], "Poor": [lit("<"), x39]}
    else:
        spec = {"Good": [lit("≤"), x69], "Fair": [lit(">"), x69, lit("-≤"), x39], "Poor": [lit(">"), x39]}
    return dict((k, ("&" if formula else "").join(v + tail)) for k, v in spec.items())


def solid(hex_: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_)


def cf_fill(hex_: str) -> PatternFill:
    """A fill for conditional formatting (Excel reads the dxf fill from start and end colours)."""
    return PatternFill(start_color=hex_, end_color=hex_, fill_type="solid")


FILL_INPUT = solid(C_INPUT)
FILL_LINK = solid(C_LINK)
FILL_BAND = solid(C_BAND)
FILL_CALC = solid(C_CALC)
FILL_HEAD = solid(C_BAND)
FONT_HEAD = Font(bold=True)
FONT_BOLD = Font(bold=True)
FONT_NOTE = Font(italic=True, color="555555", size=9)
THIN = Side(style="thin", color="000000")
MED = Side(style="medium", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNLOCKED = Protection(locked=False)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER_NOWRAP = Alignment(horizontal="center", vertical="center")
LEFT_MID = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
RIGHT_MID = Alignment(horizontal="right", vertical="center")

# ---------------------------------------------------------------------------
# Inputs: one entry per quantity. ``keys`` are the catalog input keys the entry
# serves. The worksheet places each entry under the first function that reads
# it (FRONT_ROWS) and repeats it as a linked cell under the others.
# ---------------------------------------------------------------------------
CONTEXT_INPUTS = [
    # name, label, kind, options/min/max, note
    ("ctx_region", "NARS-9 aggregate ecoregion of the site", "list", None,
     "Code of the EPA NARS nine-region ecoregion containing the site (CPL, NAP, NPL, SAP, SPL, TPL, UMW, WMT, XER). "
     "It selects the regional reference curves and the nutrient thresholds. Leave blank if unknown: the national "
     "curves apply and nutrients fall back to the CHEM index."),
    ("ctx_fcode", "NHD feature code (FCODE) of the reach", "whole", (0, 99999),
     "From NHDPlus V2. Codes 33600, 33601 and 33603 are canals and ditches, which rate Channel evolution Poor."),
    ("ctx_slope_override", "Slope class override", "list", None,
     "Normally derived from the channel slope below. Set only to force a class (lt_0.5, 0.5_to_2, ge_2 or national)."),
]

QUANTITY_INPUTS = [
    # group, key, label, units, kind, (min, max), source, note
    ("Watershed land cover (StreamCat, watershed area of interest)", "impervious",
     "Impervious cover of the watershed", "%", "decimal", (0, 100), "StreamCat pctimp2019ws",
     "Used by Catchment hydrology and Light and thermal regime."),
    ("Watershed land cover (StreamCat, watershed area of interest)", "agriculture",
     "Agricultural cover of the watershed (crop plus hay and pasture)", "%", "decimal", (0, 100),
     "StreamCat pctcrop2019ws + pcthay2019ws (the application rounds the sum to two decimals)",
     "Used by Catchment hydrology, Sediment continuity and Bed composition."),
    ("Watershed land cover (StreamCat, watershed area of interest)", "woodyWetland",
     "Woody wetland cover of the watershed", "%", "decimal", (0, 100), "StreamCat pctwdwet2019ws", ""),
    ("Watershed land cover (StreamCat, watershed area of interest)", "herbaceousWetland",
     "Herbaceous wetland cover of the watershed", "%", "decimal", (0, 100), "StreamCat pcthbwet2019ws", ""),
    ("Watershed land cover (StreamCat, watershed area of interest)", "roadDensity",
     "Road density of the watershed", "km/km2", "decimal", (0, 1e6), "StreamCat rddensws",
     "Used by Reach inflow and Sediment continuity."),
    ("Watershed land cover (StreamCat, watershed area of interest)", "kFactor",
     "Soil erodibility K factor of the watershed", "", "decimal", (0, 1), "StreamCat kffactws", ""),
    ("Flow regulation and flow variability", "storage",
     "Normalized dam storage upstream", "m3/km2", "decimal", (0, 1e12), "StreamCat damnrmstorws",
     "A true zero is valid evidence. Leave blank when unknown."),
    ("Flow regulation and flow variability", "runoff",
     "Mean annual runoff of the watershed", "mm", "decimal", (0, 1e6), "StreamCat runoffws",
     "Must be positive for the degree of regulation."),
    ("Flow regulation and flow variability", "flowCv",
     "Monthly flow variability (coefficient of variation of the twelve EROM mean monthly flows)", "ratio",
     "decimal", (0, 1e6), "NHDPlus V2 EROM QE_01 to QE_12: population standard deviation over the mean, rounded to six decimals",
     "Enter the value from the EASI report, or compute it below from the twelve monthly flows."),
    ("Riparian corridor (StreamCat, 100 m watershed corridor)", "woodyRiparian",
     "Woody riparian cover in the 100 m corridor (forest, shrub and woody wetland)", "%", "decimal", (0, 100),
     "StreamCat conifer + deciduous + mixed forest + shrub + woody wetland at wsrp100 (rounded to one decimal)",
     "Used by Light and thermal regime and Habitat provision."),
    ("Riparian corridor (StreamCat, 100 m watershed corridor)", "forest",
     "Forest cover in the corridor", "%", "decimal", (0, 100), "StreamCat pctconif + pctdecid + pctmxfst at wsrp100", ""),
    ("Riparian corridor (StreamCat, 100 m watershed corridor)", "shrub",
     "Shrub cover in the corridor", "%", "decimal", (0, 100), "StreamCat pctshrb2019wsrp100", ""),
    ("Riparian corridor (StreamCat, 100 m watershed corridor)", "grassland",
     "Grassland cover in the corridor", "%", "decimal", (0, 100), "StreamCat pctgrs2019wsrp100", ""),
    ("Riparian corridor (StreamCat, 100 m watershed corridor)", "wetland",
     "Wetland cover in the corridor (woody and herbaceous)", "%", "decimal", (0, 100),
     "StreamCat pctwdwet + pcthbwet at wsrp100", ""),
    ("Channel geometry (3DEP cross-sections and NHDPlus attributes)", "bhr",
     "Bank-height ratio (low-bank height over maximum bankfull depth), reach median", "ratio", "decimal", (0, 1e3),
     "EASI report cross-section block or a field survey", "Used by High flow dynamics, Channel evolution and Channel and floodplain dynamics."),
    ("Channel geometry (3DEP cross-sections and NHDPlus attributes)", "er",
     "Entrenchment ratio (flood-prone width over bankfull width), reach median", "ratio", "decimal", (0, 1e3),
     "EASI report cross-section block or a field survey", "Used by Floodplain connectivity and Channel evolution."),
    ("Channel geometry (3DEP cross-sections and NHDPlus attributes)", "slope",
     "Channel slope", "m/m", "decimal", (0, 10), "NHDPlus V2 VAA slope",
     "Used by Hyporheic connectivity and to pick the entrenchment slope class."),
    ("Channel geometry (3DEP cross-sections and NHDPlus attributes)", "sinuosity",
     "Reach sinuosity", "ratio", "decimal", (1, 100), "reach geometry (channel length over straight-line length)", ""),
    ("Water quality (Water Quality Portal and ATTAINS)", "tn",
     "Total nitrogen, median of station medians within 5 miles and 10 years", "mg/L", "decimal", (0, 1e4),
     "WQP total-fraction observations", "Leave blank when there are no qualifying observations."),
    ("Water quality (Water Quality Portal and ATTAINS)", "tp",
     "Total phosphorus, median of station medians within 5 miles and 10 years", "mg/L", "decimal", (0, 1e4),
     "WQP total-fraction observations", ""),
    ("Water quality (Water Quality Portal and ATTAINS)", "category",
     "ATTAINS integrated-report category of the assessment unit at the reach or within 2 km", "", "category", None,
     "EPA ATTAINS", "1, 2, 3, 4A, 4B, 4C or 5. Category 3 and a blank fall back to the CHEM index."),
    ("Landscape integrity indices (StreamCat)", "chemCatchment",
     "CHEM integrity index, catchment", "index", "decimal", (0, 1), "StreamCat chemcat",
     "Fallback for Nutrient cycling and Water and soil quality; a component of the ICI product."),
    ("Landscape integrity indices (StreamCat)", "chemWatershed",
     "CHEM integrity index, watershed", "index", "decimal", (0, 1), "StreamCat chemws",
     "Fallback for Nutrient cycling and Water and soil quality; a component of the IWI product."),
    ("Landscape integrity indices (StreamCat)", "prGBmmi",
     "Predicted probability of Good benthic condition (prg_bmmi0809)", "probability", "decimal", (0, 1),
     "StreamCat prg_bmmi0809 at the other area of interest", "Population support uses it directly; blank falls back to the ICI/IWI products."),
    ("Landscape integrity indices (StreamCat)", "hydCat", "HYD integrity, catchment", "index", "decimal", (0, 1), "StreamCat hydcat", ""),
    ("Landscape integrity indices (StreamCat)", "sedCat", "SED integrity, catchment", "index", "decimal", (0, 1), "StreamCat sedcat", ""),
    ("Landscape integrity indices (StreamCat)", "connCat", "CONN integrity, catchment", "index", "decimal", (0, 1), "StreamCat conncat", ""),
    ("Landscape integrity indices (StreamCat)", "tempCat", "TEMP integrity, catchment", "index", "decimal", (0, 1), "StreamCat tempcat", ""),
    ("Landscape integrity indices (StreamCat)", "habtCat", "HABT integrity, catchment", "index", "decimal", (0, 1), "StreamCat habtcat", ""),
    ("Landscape integrity indices (StreamCat)", "hydWs", "HYD integrity, watershed", "index", "decimal", (0, 1), "StreamCat hydws", ""),
    ("Landscape integrity indices (StreamCat)", "sedWs", "SED integrity, watershed", "index", "decimal", (0, 1), "StreamCat sedws", ""),
    ("Landscape integrity indices (StreamCat)", "connWs", "CONN integrity, watershed", "index", "decimal", (0, 1), "StreamCat connws", ""),
    ("Landscape integrity indices (StreamCat)", "tempWs", "TEMP integrity, watershed", "index", "decimal", (0, 1), "StreamCat tempws", ""),
    ("Landscape integrity indices (StreamCat)", "habtWs", "HABT integrity, watershed", "index", "decimal", (0, 1), "StreamCat habtws", ""),
    ("Counts (USGS NAS and USACE NID)", "taxaCount",
     "Established non-native taxa recorded in the HUC12", "count", "whole", (0, 100000), "USGS NAS",
     "A successful empty query is zero. Leave blank when the query was unavailable."),
    ("Counts (USGS NAS and USACE NID)", "damCount",
     "Mapped dams within one mile of the site", "count", "whole", (0, 100000), "USACE NID", ""),
]

OVERRIDE_INPUTS = [
    ("ov_stageClass", "Observed channel-evolution class", "list_rating", None,
     "Good, Fair or Poor from a documented field assessment. Applies only with the note below."),
    ("ov_indicators", "Observed channel-evolution indicators (note)", "text", None,
     "Required for the observed class to replace the automatic proxy."),
    ("ov_erodingBankPct", "Observed eroding bank", "decimal", (0, 100),
     "Percent of bank length eroding, from field or verified imagery. Both bank entries are required."),
    ("ov_armoredBankPct", "Observed armored bank", "decimal", (0, 100),
     "Percent of bank length armored."),
]
OBSERVED_SOURCE = "Field assessment or verified imagery"

# The Override Score: the last row of every function, the rating select of the
# Assessment page. It takes Good, Fair or Poor (the application overrides the
# rating, and the score follows: 13, 8 or 3), and it is drawn apart from the
# function's metrics because it is a user-defined input, not a desktop quantity.
SCORE_OVERRIDE_LABEL = "Override Score (Optional)"
SCORE_OVERRIDE_SOURCE = "User-defined override"     # one line, so the row keeps its own height
SCORE_OVERRIDE_NOTE = ("Optional. Pick Good, Fair or Poor to override this function's computed rating and score, as "
                       "a rating changed on the EASI Assessment page does. Leave blank to keep the computed rating.")
SCORE_OVERRIDE_NOTE_OBSERVED = ("Optional. Pick Good, Fair or Poor to override this function's computed rating and "
                                "score, as a rating changed on the EASI Assessment page does. Complete observed "
                                "entries above take precedence. Leave blank to keep the computed rating.")
SCORE_OVERRIDE_ROUTE = "override score"
C_OVERRIDE = "FDE9D9"    # the override row's band: a lighter tint of the entry orange
FILL_OVERRIDE = solid(C_OVERRIDE)
DOUBLE = Side(style="double", color="000000")
OVERRIDE_ROW_HEIGHT = 24.0


def score_override_name(mkey: str) -> str:
    """Defined name of a function's Override Score entry (``ov_m03_rating``: it holds a rating)."""
    return f"ov_{mkey}_rating"


MONTH_INPUTS = [f"m{i:02d}" for i in range(1, 13)]

# ---------------------------------------------------------------------------
# The Scenario worksheet: which rows each function shows in the Metrics
# columns. ("in", key) is a quantity (an entry the first time it appears, a
# linked cell afterwards), ("ov", name) an observed-evidence entry, and
# ("calc", field, label, units) a computed quantity taken from the Metrics
# sheet (the value the catalog actually bands).
# ---------------------------------------------------------------------------
FRONT_ROWS = {
    "m01": [("in", "impervious"), ("in", "agriculture")],
    "m02": [("in", "woodyWetland"), ("in", "herbaceousWetland"),
            ("calc", "combined", "Wetland extent (sum, capped at 100)", "%")],
    "m03": [("in", "roadDensity")],
    "m04": [("in", "storage"), ("in", "runoff"), ("calc", "combined", "Degree of regulation", "%")],
    "m05": [("in", "flowCv")],
    "m06": [("in", "bhr")],
    "m07": [("in", "er"), ("in", "slope")],
    "m08": [("in", "slope"), ("in", "sinuosity")],
    "m09": [("in", "bhr"), ("in", "er"), ("ov", "ov_stageClass"), ("ov", "ov_indicators")],
    "m10": [("in", "bhr"), ("ov", "ov_erodingBankPct"), ("ov", "ov_armoredBankPct")],
    "m11": [("in", "agriculture"), ("in", "kFactor"), ("in", "roadDensity")],
    "m12": [("in", "agriculture")],
    "m13": [("in", "woodyRiparian"), ("in", "impervious")],
    "m14": [("in", "forest"), ("in", "shrub"), ("in", "grassland"), ("in", "wetland"),
            ("calc", "combined", "Natural corridor cover (sum, capped at 100)", "%")],
    "m15": [("in", "tn"), ("in", "tp"), ("in", "chemCatchment"), ("in", "chemWatershed")],
    "m16": [("in", "category"), ("in", "chemCatchment"), ("in", "chemWatershed")],
    "m17": [("in", "woodyRiparian")],
    "m18": [("in", "prGBmmi"), ("in", "hydCat"), ("in", "sedCat"), ("in", "connCat"), ("in", "tempCat"),
            ("in", "habtCat"), ("in", "chemCatchment"), ("in", "hydWs"), ("in", "sedWs"), ("in", "connWs"),
            ("in", "tempWs"), ("in", "habtWs"), ("in", "chemWatershed"),
            ("calc", "ici", "ICI (product of the six catchment indices)", "index"),
            ("calc", "iwi", "IWI (product of the six watershed indices)", "index")],
    "m19": [("in", "taxaCount")],
    "m20": [("in", "damCount")],
}

# Shorter labels and sources for the worksheet; the guidance note of each entry
# (shown when the cell is selected) carries the rest.
FRONT_LABELS = {
    "agriculture": "Agricultural cover of the watershed (crop, hay and pasture)",
    "kFactor": "Soil erodibility (K factor) of the watershed",
    "flowCv": "Monthly flow variability (CV of the twelve EROM monthly flows)",
    "woodyRiparian": "Woody riparian cover in the 100 m corridor",
    "forest": "Forest cover in the 100 m corridor",
    "shrub": "Shrub cover in the 100 m corridor",
    "grassland": "Grassland cover in the 100 m corridor",
    "wetland": "Wetland cover in the 100 m corridor (woody and herbaceous)",
    "bhr": "Bank-height ratio (reach median)",
    "er": "Entrenchment ratio (reach median)",
    "tn": "Total nitrogen (median of station medians)",
    "tp": "Total phosphorus (median of station medians)",
    "category": "ATTAINS integrated-report category",
    "prGBmmi": "Predicted probability of Good benthic condition",
    "hydCat": "HYD integrity index, catchment", "sedCat": "SED integrity index, catchment",
    "connCat": "CONN integrity index, catchment", "tempCat": "TEMP integrity index, catchment",
    "habtCat": "HABT integrity index, catchment", "hydWs": "HYD integrity index, watershed",
    "sedWs": "SED integrity index, watershed", "connWs": "CONN integrity index, watershed",
    "tempWs": "TEMP integrity index, watershed", "habtWs": "HABT integrity index, watershed",
    "taxaCount": "Established non-native taxa recorded in the HUC12",
    "damCount": "Mapped dams within one mile of the site",
}
FRONT_SOURCES = {
    "agriculture": "StreamCat pctcrop2019ws + pcthay2019ws",
    "flowCv": "NHDPlus V2 EROM QE_01 to QE_12, or the helper below",
    "woodyRiparian": "StreamCat forest + shrub + woody wetland at wsrp100",
    "forest": "StreamCat pctconif + pctdecid + pctmxfst at wsrp100",
    "wetland": "StreamCat pctwdwet + pcthbwet at wsrp100",
    "bhr": "EASI report cross-section block or a field survey",
    "er": "EASI report cross-section block or a field survey",
    "sinuosity": "Reach geometry (channel length over straight-line length)",
    "tn": "WQP total-fraction observations within 5 miles and 10 years",
    "tp": "WQP total-fraction observations within 5 miles and 10 years",
    "category": "EPA ATTAINS (at the reach or within 2 km)",
    "prGBmmi": "StreamCat prg_bmmi0809",
}
FRONT_NOTES = {
    "agriculture": "The application rounds the sum to two decimals. Used by Catchment hydrology, Sediment continuity "
                   "and Bed composition.",
    "woodyRiparian": "Conifer, deciduous and mixed forest, shrub and woody wetland at wsrp100, rounded to one decimal. "
                     "Used by Light and thermal regime and Habitat provision.",
    "flowCv": "Population standard deviation over the mean of the twelve EROM mean monthly flows, rounded to six "
              "decimals. Enter the value from the EASI report or compute it in the helper below the table.",
    "prGBmmi": "Population support uses the model probability directly. Leave it blank to fall back to the twelve "
               "integrity indices (ICI and IWI products).",
}

#: The Instructions sheet's text, in SFARI's words where the two calculators work alike (owner, 2026-10-23)
INSTRUCTIONS = [
    "Add the site information into the orange cells at the top of each scenario worksheet.",
    "Enter each metric value in the orange cells of the Value column (column J). The Source column names "
    "where each value comes from, and selecting a cell shows its guidance note. Grey cells repeat a value "
    "entered higher in the table.",
    "Leave a cell blank when a value is unavailable, and never enter zero for an unknown value. A blank "
    "metric leaves its function unrated, and unrated functions drop out of the sub-indices and the EASI Index.",
    "Optional. Pick an Override Score (Good, Fair or Poor) in the last row of a function to replace its "
    "computed rating, and enter observed channel and bank conditions in their rows when you have field "
    "observations.",
    "Once the metric values are entered, the ratings and Function Scores (columns G and H), the Functional "
    "Categories, the Outcomes and the overall EASI Index (column A) will auto-populate. The scoring summary "
    "table and the bar charts to the right of the table will also auto-populate.",
    "Record other metrics, notes, or sources in the grey box at the bottom of the worksheet.",
]
COMMENTS = [
    "The calculator scores exactly as the EASI web application does, with the same criteria, reference "
    "curves, weights and rounding. The same inputs give the same ratings, Function Scores, sub-indices and "
    "EASI Index.",
    "The Metrics sheet shows how each rating was reached, beside the Good, Fair and Poor ranges of every "
    "metric. The Reference sheet plots the reference curves and holds the tables the calculator reads. "
    "Nothing is entered on either sheet.",
    "The optional monthly flow helper below the table computes the monthly flow variability from the twelve "
    "EROM mean monthly flows. Copy its result into the Low flow and baseflow dynamics row.",
    "Derived values are rounded to twelve decimals before they are rated, and displayed indices to two "
    "decimals, as in the web application. When a sub-index lands exactly on a half-hundredth, Excel rounds "
    "away from zero and the web application rounds to the even digit, so a displayed value can differ by "
    "0.01 while the class stays the same.",
    "The calculator does not retrieve data, delineate a watershed or draw cross-sections. An Override Score "
    "replaces the computed rating and score of its function, as a rating changed in the web application "
    "does, and the Metrics sheet keeps the computed rating beside it.",
    "EASI is a screening-level desktop estimate, not a field-validated assessment. The reference curves "
    "express regional expectations derived from least-disturbed reaches, and several proxies remain "
    "unvalidated. The technical report states each metric's limitations.",
]
DEVELOPERS = "Leanne M. Stepchinski, Gabrielle C. David, Samantha R. Wiest, and Garrett T. Menichino"
EMAILS = ("Leanne.M.Stepchinski@usace.army.mil, Gabrielle.C.David@usace.army.mil, "
          "Samantha.R.Wiest@erdc.dren.mil, Garrett.T.Menichino@usace.army.mil")
WAIVER = ('This model is provided for free as part of the USACE Technical Report "Ecosystem Assessment '
          'Screening Index (EASI)" by ' + DEVELOPERS + ".")
NO_LIABILITY = ("None of the authors nor the US Army Corps of Engineers accepts responsibility or liability for "
                "the model's use by third parties.")


# ---------------------------------------------------------------------------
# Small formula helpers
# ---------------------------------------------------------------------------
def fnum(x) -> str:
    """A number literal Excel reads back exactly (repr keeps 17 digits)."""
    if isinstance(x, bool):
        return "TRUE" if x else "FALSE"
    if isinstance(x, int):
        return str(x)
    return repr(float(x))


def q(s: str) -> str:
    return '"' + str(s).replace('"', '""') + '"'


def band_predicate(b: dict, x: str) -> str:
    """The exact interval test of one catalog band (matches rating_for_value)."""
    parts = []
    if b.get("min") is not None:
        parts.append(f"{x}>={fnum(b['min'])}" if b.get("minInclusive") else f"{x}>{fnum(b['min'])}")
    if b.get("max") is not None:
        parts.append(f"{x}<={fnum(b['max'])}" if b.get("maxInclusive") else f"{x}<{fnum(b['max'])}")
    if not parts:
        return "TRUE"
    return parts[0] if len(parts) == 1 else "AND(" + ",".join(parts) + ")"


def bands_formula(bands: list[dict], x: str) -> str:
    """Rating text from bands, "" when unrated or not a number."""
    expr = '""'
    for b in reversed(bands):
        expr = f"IF({band_predicate(b, x)},{q(b['rating'])},{expr})"
    return f"IF(NOT(ISNUMBER({x})),\"\",{expr})"


def index_from_rating(r: str, anchors: dict[str, str]) -> str:
    """Index anchor from a rating cell (text), "" when blank."""
    return (f"IF({r}={q('Good')},{anchors['Good']},IF({r}={q('Fair')},{anchors['Fair']},"
            f"IF({r}={q('Poor')},{anchors['Poor']},\"\")))")


def rating_from_index_edges(idx: str) -> str:
    lo, hi = INDEX_EDGES
    return f"IF(NOT(ISNUMBER({idx})),\"\",IF({idx}>={fnum(hi)},\"Good\",IF({idx}>={fnum(lo)},\"Fair\",\"Poor\")))"


def rating_from_anchor(idx: str, anchors: dict[str, str]) -> str:
    """Rating whose anchor equals a combined worst/best index."""
    return (f"IF(NOT(ISNUMBER({idx})),\"\",IF({idx}={anchors['Good']},\"Good\","
            f"IF({idx}={anchors['Fair']},\"Fair\",IF({idx}={anchors['Poor']},\"Poor\",\"\"))))")


def link_formula(ref: str) -> str:
    """Show an entry made elsewhere, blank when it is blank."""
    return f'=IF({ref}="","",{ref})'


class Sheet:
    """Row-by-row writer with absolute references and defined names."""

    def __init__(self, wb: Workbook, ws, name: str):
        self.wb, self.ws, self.name = wb, ws, name
        self.row = 1

    def ref(self, col: int, row: int) -> str:
        return f"'{self.name}'!${get_column_letter(col)}${row}"

    def rng(self, col: int, r1: int, r2: int) -> str:
        c = get_column_letter(col)
        return f"'{self.name}'!${c}${r1}:${c}${r2}"

    def name_cell(self, name: str, col: int, row: int) -> str:
        ref = self.ref(col, row)
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
        return ref

    def name_range(self, name: str, col: int, r1: int, r2: int) -> str:
        ref = self.rng(col, r1, r2)
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
        return ref

    def cell(self, col: int, value=None, *, row: int | None = None, fill=None, font=None,
             fmt=None, unlocked=False, wrap=False, border=False, align=None):
        c = self.ws.cell(row=row or self.row, column=col)
        if value is not None:
            c.value = value
        if fill is not None:
            c.fill = fill
        if font is not None:
            c.font = font
        if fmt is not None:
            c.number_format = fmt
        if unlocked:
            c.protection = UNLOCKED
        if wrap:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if align is not None:
            c.alignment = align
        if border is True:
            c.border = BOX
        elif border:
            c.border = border
        return c

    def header(self, *labels: str, widths: dict[int, int] | None = None):
        for i, label in enumerate(labels, 1):
            self.cell(i, label, fill=FILL_HEAD, font=FONT_HEAD, border=True)
        self.row += 1

    def title(self, text: str, size: int = 13):
        self.cell(1, text, font=Font(bold=True, size=size))
        self.row += 1

    def blank(self, n: int = 1):
        self.row += n

    def merge(self, r1: int, c1: int, r2: int, c2: int):
        if (r1, c1) != (r2, c2):
            self.ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def box(ws, r1: int, c1: int, r2: int, c2: int, outer: Side = MED, inner: Side | None = THIN):
    """Border a rectangle: ``outer`` on its edges, ``inner`` between its cells."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            old = cell.border
            cell.border = Border(
                left=outer if c == c1 else (inner or old.left),
                right=outer if c == c2 else (inner or old.right),
                top=outer if r == r1 else (inner or old.top),
                bottom=outer if r == r2 else (inner or old.bottom))


def side(ws, r1: int, c1: int, r2: int, c2: int, **sides: Side):
    """Override named sides (left, right, top, bottom) of every cell in a rectangle."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            old = cell.border
            cell.border = Border(left=sides.get("left", old.left), right=sides.get("right", old.right),
                                 top=sides.get("top", old.top), bottom=sides.get("bottom", old.bottom))


def lines_for(text: str, width_chars: int) -> int:
    return max(1, -(-len(text or "") // width_chars))


# ---------------------------------------------------------------------------
# The generator
# ---------------------------------------------------------------------------
class Builder:
    def __init__(self, method_stamp: str | None = None):
        #: the scoring method digest the workbook names in its identity cells: the running
        #: evaluator's, or the one a committed workbook recorded when it is checked
        self.method_stamp = method_stamp or method_version()
        self.catalog = sm.catalog()
        self.methods = {m["methodKey"]: m for m in self.catalog["methods"]}
        self.by_metric = {m["metricId"]: m for m in self.catalog["methods"]}
        self.curves = sm.curve_sets()
        self.metrics = config.easi_metrics()["metrics"]
        self.meta = {m["metricId"]: m for m in self.metrics}
        self.cwa = config.cwa_mapping()          # functionId -> {physical, chemical, biological}
        self.identity = config.scoring_identity()
        self.rating_index = self.catalog.get("ratingIndex") or config.RATING_INDEX
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.regions = sorted({k for s in self.curves.values() if s["stratifier"] == "nars9"
                               for k in s["curves"] if k != "national"})
        self.slope_classes = ["lt_0.5", "0.5_to_2", "ge_2"]
        self.problems = sm.validate_catalog()
        if self.problems:
            raise SystemExit("catalog problems: " + "; ".join(self.problems))
        for m in self.metrics:
            mids = m.get("indexMidpoints") or {}
            if mids and any(abs(float(mids[k]) - float(self.rating_index[k])) > 1e-12 for k in RATINGS):
                raise SystemExit(f"{m['metricId']}: indexMidpoints differ from the catalog ratingIndex")
        for k in RATINGS:
            if abs(float(self.rating_index[k]) * config.FUNCTION_SCORE_MAX % 1 - 0.5) < 1e-6:
                raise SystemExit("a rating anchor times 15 sits on a rounding tie")
        self.scores = {k: scoring.function_score(float(self.rating_index[k])) for k in RATINGS}
        self.qspec = {row[1]: row for row in QUANTITY_INPUTS}
        self.ovspec = {row[0]: row for row in OVERRIDE_INPUTS}
        self.ov_source = {row[0]: OBSERVED_SOURCE for row in OVERRIDE_INPUTS}
        self._observed_guidance()
        self.ctxspec = {row[0]: row for row in CONTEXT_INPUTS}
        # the worksheet rows: FRONT_ROWS, then the Override Score as the last row of
        # every function (the Assessment page offers the rating select on all 20)
        self.front: dict[str, list[tuple]] = {}
        self.score_overrides: set[str] = set()
        for n, m in enumerate(self.metrics, 1):
            mkey = f"m{n:02d}"
            rows = list(FRONT_ROWS[mkey])
            name = score_override_name(mkey)
            observed = any(kind == "ov" for kind, *_ in rows)
            self.ovspec[name] = (name, SCORE_OVERRIDE_LABEL, "list_rating", None,
                                 SCORE_OVERRIDE_NOTE_OBSERVED if observed else SCORE_OVERRIDE_NOTE)
            self.ov_source[name] = SCORE_OVERRIDE_SOURCE
            self.score_overrides.add(name)
            rows.append(("ov", name))
            self.front[mkey] = rows
        self._check_front_rows()
        # cell references filled while writing
        self.inputs: dict[str, str] = {}       # key -> absolute ref of the entry cell
        self.anchor_ref: dict[str, str] = {}   # rating -> absolute ref of its index anchor
        self.score_ref: dict[str, str] = {}
        self.curve_rows: dict[str, tuple[int, int]] = {}
        self.curve_cols: dict[str, int] = {}
        self.result_cells: dict[str, dict[str, str]] = {}  # metricId -> {rating, index, score, ...}
        self.input_rating: dict[tuple[str, str], str] = {}  # (m01, key) -> rating cell of that input
        self.current_mkey = ""
        self.front_rows: list[tuple[str, str, str, int]] = []   # (mkey, kind, key/field, row)
        self.front_span: dict[str, tuple[int, int]] = {}        # mkey -> (first row, last row)
        self.rollup: dict[str, str] = {}                         # results refs used by the worksheet

    def _observed_guidance(self):
        """Guidance notes of the four observed entries, written from the catalog.

        The note is what a user reads on selecting the cell, so it has to say what Good, Fair
        and Poor mean and where the value comes from: the stage wording of the channel method's
        field reference and the class limits of the observed bank bands, the same text the
        Assessment page shows on the two cards. Excel cuts a note at 255 characters.
        """
        chan = self.methods["channel-adjustment-susceptibility"]["fieldReference"]["criteria"]
        bank = self.variant("bhr-bank-instability-susceptibility", "observed-bank-condition")

        def limits(key: str) -> str:
            bands = next(i for i in bank["inputs"] if i["key"] == key)["bands"]
            return ", ".join(f"{b['rating']} {b['label']}" for b in bands)

        notes = {
            "ov_stageClass": " ".join(f"{r}: {chan[r].rstrip('.')}." for r in RATINGS)
                             + " From a field visit or report. Needs the note below.",
            "ov_indicators": "Required for the observed class to apply. Record what was seen (headcuts, bars, width "
                             "or depth change, bank erosion, recovery features) with the source and the date.",
            "ov_erodingBankPct": f"Percent of bank length eroding, from a field visit or verified imagery. "
                                 f"{limits('erodingBankPct')}. Both bank entries are required and the worse governs.",
            "ov_armoredBankPct": f"Percent of bank length artificially armored. {limits('armoredBankPct')}. "
                                 "Both bank entries are required and the worse governs.",
        }
        for name, note in notes.items():
            assert len(note) <= 255, (name, len(note))
            spec = self.ovspec[name]
            self.ovspec[name] = (*spec[:4], note)

    def _check_front_rows(self):
        """Every entry appears once, every catalog input of a method is shown under it."""
        shown = [k for rows in self.front.values() for kind, k, *_ in rows if kind == "in"]
        missing = [k for k in self.qspec if k not in shown]
        assert not missing, f"quantities without a worksheet row: {missing}"
        unknown = [k for k in shown if k not in self.qspec]
        assert not unknown, f"worksheet rows without a quantity: {unknown}"
        ov = [k for rows in self.front.values() for kind, k, *_ in rows if kind == "ov"]
        assert sorted(ov) == sorted(self.ovspec), ov
        for n, m in enumerate(self.metrics, 1):
            mkey = f"m{n:02d}"
            method = self.by_metric[m["metricId"]]
            keys = {ALIASES.get(k, k) for kind, k, *_ in FRONT_ROWS[mkey] if kind == "in"}
            for i in method.get("inputs", []):
                if i.get("required") and not i.get("contextOnly") and i["key"] in self.qspec:
                    assert i["key"] in keys, f"{mkey}: catalog input {i['key']} is not on the worksheet"
            for v in method.get("variants") or []:
                for name, comps in ((v.get("formula") or {}).get("products") or {}).items():
                    for k in comps:
                        assert ALIASES.get(k, k) in keys, f"{mkey}: product component {k} is not on the worksheet"

    # ---- Reference -------------------------------------------------------
    def build_reference(self):
        ws = self.wb.create_sheet("Reference")
        S = Sheet(self.wb, ws, "Reference")
        self.reference_ws = ws
        S.cell(6, NO_INPUT, font=FONT_BOLD)
        S.title("Reference tables generated from the EASI scoring definitions")
        S.row += 1

        # the reference curves, drawn (owner, 2026-10-23): a chart per curve set over the rows of
        # the curve table below it, the tables the formulas read after it
        S.title("Reference curves", 11)
        S.cell(1, "Each line is the curve of one NARS-9 region or slope class. A value scores the index of its "
                  "curve: Good at 0.69 and above, Fair at 0.39 and above, Poor below. The Good, Fair and Poor "
                  "ranges of every curve are in the table under the charts.", font=FONT_NOTE)
        S.row += 1
        self.chart_rows = (S.row, S.row + CHART_ROWS)
        S.row += 2 * CHART_ROWS + 1
        self._curve_table(S)
        self._rating_lines(S)

        # rating anchors
        S.title("Rating anchors and function scores", 11)
        S.header("Rating", "Index anchor", "Function score", "Score class")
        for r in RATINGS:
            S.cell(1, r)
            S.cell(2, float(self.rating_index[r]), fmt="0.000")
            S.cell(3, self.scores[r])
            S.cell(4, "Functioning" if self.scores[r] > 10 else "Functioning-at-Risk" if self.scores[r] > 5
                   else "Non-Functioning")
            self.anchor_ref[r] = S.name_cell(f"anchor_{r.lower()}", 2, S.row)
            self.score_ref[r] = S.name_cell(f"score_{r.lower()}", 3, S.row)
            S.row += 1
        S.blank()
        S.title("Condition classes", 11)
        S.header("Index at or below", "Class", "Function score at or below")
        for edge, label, fedge in ((0.39, "Non-Functioning", 5), (0.69, "Functioning-at-Risk", 10),
                                   (1.0, "Functioning", 15)):
            S.cell(1, edge); S.cell(2, label); S.cell(3, fedge); S.row += 1
        S.blank()

        # bands per input (informational)
        S.title("Fixed bands per input (informational; the Metrics formulas carry the same tests)", 11)
        S.header("Method", "Input", "Rating", "Min", "Min inclusive", "Max", "Max inclusive", "Label")
        for m in self.catalog["methods"]:
            for i in m.get("inputs", []):
                for b in i.get("bands") or []:
                    self._band_row(S, m["methodKey"], i["key"], b)
            for b in m.get("bands") or []:
                self._band_row(S, m["methodKey"], "(combined)", b)
            for v in m.get("variants") or []:
                for i in v.get("inputs", []):
                    for b in i.get("bands") or []:
                        self._band_row(S, v.get("methodKey"), i["key"], b)
                for b in v.get("bands") or []:
                    self._band_row(S, v.get("methodKey"), "(combined)", b)
        S.blank()

        # nutrient regional bands
        nut = self.methods["regional-nutrient-condition"]
        tn = next(i for i in nut["inputs"] if i["key"] == "tn")["regionalBands"]
        tp = next(i for i in nut["inputs"] if i["key"] == "tp")["regionalBands"]
        S.title("Nutrient regional thresholds by NARS-9 region (mg/L): Good at or below the first value, "
                "Poor at or above the second, Fair between", 11)
        S.header("Region", "TN good/fair", "TN fair/poor", "TP good/fair", "TP fair/poor")
        first = S.row
        for region in self.regions:
            S.cell(1, region); S.cell(2, float(tn[region][0])); S.cell(3, float(tn[region][1]))
            S.cell(4, float(tp[region][0])); S.cell(5, float(tp[region][1])); S.row += 1
        last = S.row - 1
        self.nut_regions = S.name_range("nutrient_regions", 1, first, last)
        self.nut_cols = {k: S.name_range(f"nutrient_{k}", col, first, last)
                         for k, col in (("tn_gf", 2), ("tn_fp", 3), ("tp_gf", 4), ("tp_fp", 5))}
        S.blank()

        # categorical lookups
        att = self.methods["attains-regulatory-category"]["formula"]["lookup"]
        S.title("ATTAINS category lookup (category 3 is inconclusive and falls back to the CHEM index)", 11)
        S.header("Category", "Rating", "Label")
        first = S.row
        for key, item in att.items():
            # "none" rather than an empty cell: INDEX of an empty cell reads as 0 in Excel
            S.cell(1, str(key)); S.cell(2, item.get("rating") or "none"); S.cell(3, item.get("label")); S.row += 1
        last = S.row - 1
        self.att_keys = S.name_range("attains_keys", 1, first, last)
        self.att_ratings = S.name_range("attains_ratings", 2, first, last)
        S.blank()
        canal = self.variant("channel-adjustment-susceptibility", "channelized-fcode")["formula"]["lookup"]
        self.canal_codes = sorted(int(k) for k in canal)
        S.title("Canal and ditch feature codes (Channel evolution rates Poor)", 11)
        S.header("FCODE", "Rating")
        for code in self.canal_codes:
            S.cell(1, code); S.cell(2, canal[str(code)]["rating"]); S.row += 1
        S.blank()

        # outcome weights
        S.title("Function to outcome mapping and weights (D direct, i indirect)", 11)
        S.header("Function", "Physical", "Chemical", "Biological", "Weight D", "Weight i")
        for m in self.metrics:
            row = self.cwa[m["functionId"]]
            S.cell(1, m["functionName"]); S.cell(2, row["physical"]); S.cell(3, row["chemical"])
            S.cell(4, row["biological"]); S.cell(5, config.WEIGHTS["D"]); S.cell(6, config.WEIGHTS["i"])
            S.row += 1
        S.blank()
        S.title("Slope classes for the entrenchment curves", 11)
        S.header("Class", "Channel slope")
        for cls, txt in (("lt_0.5", "below 0.005 m/m"), ("0.5_to_2", "0.005 to below 0.02 m/m"),
                         ("ge_2", "0.02 m/m and above"), ("national", "slope unknown or negative")):
            S.cell(1, cls); S.cell(2, txt); S.row += 1
        for col, width in ((1, 34), (2, 16), (3, 16), (4, 14), (5, 12), (6, 12)):
            ws.column_dimensions[get_column_letter(col)].width = width
        for col in range(CURVE_X, CURVE_X + 10):
            ws.column_dimensions[get_column_letter(col)].width = 10
        for col in (CURVE_BINS, CURVE_BINS + 1, CURVE_BINS + 2):
            ws.column_dimensions[get_column_letter(col)].width = 13
        self._curve_charts()
        ws.freeze_panes = None
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = C_TAB_NO_INPUT
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    # -- the reference curves: one row per curve, drawn as a chart per set -------------
    def _curve_table(self, S: Sheet):
        """The 34 curves, knots padded to five, values (x1 to x5) then indices (y1 to y5) side by
        side so each row is a chart series, with the Good, Fair and Poor ranges each curve gives."""
        S.title("Reference curves (frozen), one row per curve, knots padded to five", 11)
        heads = (["Key", "Set", "Stratum", "Higher is better", "n", "Members"]
                 + [f"x{k}" for k in range(1, 6)] + [f"y{k}" for k in range(1, 6)]
                 + ["Good", "Fair", "Poor", "x39", "x69", "q25", "q50", "q75", "Tier"])
        S.header(*heads)
        for j, rating in enumerate(RATINGS):
            S.cell(CURVE_BINS + j, row=S.row - 1, fill=solid(CLASS_FILL[rating]), align=CENTER_NOWRAP)
        first = S.row
        self.curve_set_rows: dict[str, list[tuple[str, int]]] = {}
        for set_id, s in self.curves.items():
            digits, unit = CURVE_DISPLAY.get(set_id, (2, ""))
            for key, c in s["curves"].items():
                pts = [(float(x), float(y)) for x, y in c["points"]]
                assert all(pts[i][0] < pts[i + 1][0] for i in range(len(pts) - 1)), (set_id, key)
                assert 2 <= len(pts) <= 5, (set_id, key, len(pts))
                while len(pts) < 5:
                    pts.append(pts[-1])
                S.cell(1, f"{set_id}|{key}"); S.cell(2, set_id); S.cell(3, key)
                S.cell(4, "TRUE" if s["higherIsBetter"] else "FALSE")
                S.cell(5, c.get("n")); S.cell(6, c.get("nMembers"))
                for k, (x, y) in enumerate(pts):
                    S.cell(CURVE_X + k, x); S.cell(CURVE_X + 5 + k, y)
                bins = curve_bins(s["higherIsBetter"], num_text(c["x39"], digits), num_text(c["x69"], digits), unit)
                for j, rating in enumerate(RATINGS):
                    S.cell(CURVE_BINS + j, bins[rating], fill=solid(BIN_FILL[rating]), align=CENTER_NOWRAP)
                S.cell(CURVE_BINS + 3, c.get("x39")); S.cell(CURVE_BINS + 4, c.get("x69"))
                S.cell(CURVE_BINS + 5, c.get("q25")); S.cell(CURVE_BINS + 6, c.get("q50"))
                S.cell(CURVE_BINS + 7, c.get("q75")); S.cell(CURVE_BINS + 8, c.get("panelTier"))
                self.curve_set_rows.setdefault(set_id, []).append((key, S.row))
                S.row += 1
        last = S.row - 1
        self.curve_rows["all"] = (first, last)
        self.curve_key_range = S.name_range("curve_keys", 1, first, last)
        self.curve_col_range = {}
        for k in range(1, 6):
            self.curve_col_range[f"x{k}"] = S.name_range(f"curve_x{k}", CURVE_X + k - 1, first, last)
            self.curve_col_range[f"y{k}"] = S.name_range(f"curve_y{k}", CURVE_X + 4 + k, first, last)
        for k, col in (("x39", CURVE_BINS + 3), ("x69", CURVE_BINS + 4)):
            self.curve_col_range[k] = S.name_range(f"curve_{k}", col, first, last)
        S.blank()

    def _rating_lines(self, S: Sheet):
        """The two index lines every curve chart draws (0.39 and 0.69), across each set's values."""
        S.title("Rating lines of the curve charts", 11)
        S.header("Set", "Value from", "Value to", "Index 0.39", "Index 0.39", "Index 0.69", "Index 0.69")
        self.rating_line_rows = {}
        for set_id, s in self.curves.items():
            xs = [float(x) for c in s["curves"].values() for x, _y in c["points"]]
            S.cell(1, set_id); S.cell(2, min(xs)); S.cell(3, max(xs))
            for col, y in ((4, 0.39), (5, 0.39), (6, 0.69), (7, 0.69)):
                S.cell(col, y)
            self.rating_line_rows[set_id] = S.row
            S.row += 1
        S.blank()

    def _curve_charts(self):
        """A scatter chart per curve set, two across: a line per region or slope class through its
        knots, and the 0.39 and 0.69 lines dashed. They read the Reference sheet only, so the sheet
        stays one copy for every scenario of a completed workbook."""
        ws = self.reference_ws
        top, second = self.chart_rows
        places = [f"A{top}", f"F{top}", f"A{second + 1}", f"F{second + 1}"]
        for place, (set_id, rows) in zip(places, self.curve_set_rows.items()):
            ch = ScatterChart(scatterStyle="lineMarker")
            ch.title = CURVE_TITLE.get(set_id, set_id)
            ch.title.overlay = False
            ch.style = 2
            for key, r in rows:
                xs = Reference(ws, min_col=CURVE_X, max_col=CURVE_X + 4, min_row=r, max_row=r)
                ys = Reference(ws, min_col=CURVE_X + 5, max_col=CURVE_X + 9, min_row=r, max_row=r)
                label = SLOPE_LABEL.get(key) or ("National" if key == "national" else key)
                s = Series(ys, xs, title=label)
                s.marker = Marker(symbol="circle", size=3)
                s.smooth = False
                ch.series.append(s)
            lr = self.rating_line_rows[set_id]
            xs = Reference(ws, min_col=2, max_col=3, min_row=lr, max_row=lr)
            for c1, label in ((4, "Index 0.39 (Fair at or above)"), (6, "Index 0.69 (Good at or above)")):
                s = Series(Reference(ws, min_col=c1, max_col=c1 + 1, min_row=lr, max_row=lr), xs, title=label)
                s.marker = Marker(symbol="none")
                s.smooth = False
                s.graphicalProperties = GraphicalProperties(
                    ln=LineProperties(solidFill="404040", prstDash="dash", w=15875))
                ch.series.append(s)
            ch.x_axis.title = CURVE_AXIS.get(set_id, set_id)
            ch.y_axis.title = "Index"
            ch.x_axis.title.overlay = False
            ch.y_axis.title.overlay = False
            ch.x_axis.scaling.min = 0
            if CURVE_DISPLAY.get(set_id, (2, ""))[1] == "%":
                ch.x_axis.scaling.max = 100
                ch.x_axis.majorUnit = 10
            ch.y_axis.scaling.min = 0
            ch.y_axis.scaling.max = 1
            ch.y_axis.majorUnit = 0.1
            ch.y_axis.numFmt = NumFmt(formatCode="0.0", sourceLinked=False)
            ch.x_axis.delete = False
            ch.y_axis.delete = False
            ch.x_axis.majorGridlines = None
            ch.legend.position = "b"
            ch.legend.overlay = False
            ch.width, ch.height = CHART_WIDTH_CM, CHART_HEIGHT_CM
            ws.add_chart(ch, place)

    def _band_row(self, S: Sheet, method: str, key: str, b: dict):
        S.cell(1, method); S.cell(2, key); S.cell(3, b["rating"])
        S.cell(4, b.get("min")); S.cell(5, "TRUE" if b.get("minInclusive") else "FALSE")
        S.cell(6, b.get("max")); S.cell(7, "TRUE" if b.get("maxInclusive") else "FALSE")
        S.cell(8, b.get("label")); S.row += 1

    def variant(self, method_key: str, variant_key: str) -> dict:
        for v in self.methods[method_key].get("variants") or []:
            if v.get("methodKey") == variant_key:
                return v
        raise KeyError(variant_key)

    # ---- Scenario: layout and entry cells (first pass) ----------------------
    def _validation(self, ws, kind: str, bounds, *, title: str = "", note: str = "",
                    error_title: str = "", error: str = "") -> DataValidation | None:
        """One validation per entry so each cell carries its own guidance note."""
        common = dict(allow_blank=True, showErrorMessage=True, errorStyle="stop",
                      errorTitle=error_title or "Entry refused", error=error)
        if note:
            common.update(showInputMessage=True, promptTitle=(title or "Guidance")[:32], prompt=note[:255])
        if kind == "decimal":
            dv = DataValidation(type="decimal", operator="between", formula1=fnum(bounds[0]),
                                formula2=fnum(bounds[1]), **common)
        elif kind == "whole":
            dv = DataValidation(type="whole", operator="between", formula1=str(bounds[0]),
                                formula2=str(bounds[1]), **common)
        elif kind == "list":
            dv = DataValidation(type="list", formula1='"' + ",".join(bounds) + '"', **common)
        else:
            if not note:
                return None
            dv = DataValidation(type="custom", formula1="TRUE", allow_blank=True, showErrorMessage=False,
                                showInputMessage=True, promptTitle=(title or "Guidance")[:32], prompt=note[:255])
        ws.add_data_validation(dv)
        return dv

    def _entry(self, S: Sheet, row: int, c1: int, c2: int, name: str, *, kind: str = "text", bounds=None,
               title: str = "", note: str = "", error: str = "", fmt: str | None = None,
               font: Font | None = None, align: Alignment = CENTER_NOWRAP) -> str:
        """An orange, unlocked entry cell (merged c1..c2) with its validation and guidance."""
        c = S.cell(c1, None, row=row, fill=FILL_INPUT, unlocked=True, align=align, fmt=fmt, font=font)
        S.merge(row, c1, row, c2)
        dv = self._validation(S.ws, kind, bounds, title=title, note=note, error=error)
        if dv is not None:
            dv.add(c)
        return S.name_cell(name, c1, row)

    def build_score_layout(self):
        ws = self.wb.create_sheet(SCORE)
        S = Sheet(self.wb, ws, SCORE)
        self.score_ws, self.score = ws, S
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 70                # SFARI's worksheet
        # title, boxed as SFARI's
        S.cell(1, "Ecosystem Assessment Screening Index (EASI) Worksheet", row=1, fill=FILL_BAND,
               font=Font(bold=True, size=16), align=CENTER_NOWRAP)
        S.merge(1, 1, 1, 13)
        box(ws, 1, 1, 1, 13, outer=MED, inner=None)
        ws.row_dimensions[1].height = 24
        # site block: label, entry (merged), label, entry, label, entry per row; SFARI's 16 pt labels
        # and 12 pt entries
        label_font = Font(bold=True, size=12)
        site_font = Font(bold=True, size=16)
        entry_font = Font(size=12)
        entry_align = Alignment(horizontal="left", vertical="center")
        region_note = self.ctxspec["ctx_region"][4]
        fcode_note = self.ctxspec["ctx_fcode"][4]
        slope_note = self.ctxspec["ctx_slope_override"][4]
        site = [
            (2, 1, "Reach ID", 2, 5, "site_name", dict()),
            (2, 6, "Assessor(s)", 7, 9, "site_assessor", dict()),
            (2, 10, "COMID", 11, 13, "site_comid",
             dict(title="COMID", note="NHDPlus V2 reach identifier of the assessed reach (from the EASI report).")),
            (3, 1, "Lat/Long", 2, 5, "site_coords", dict()),
            (3, 6, "Date", 7, 9, "site_date", dict(fmt="yyyy-mm-dd")),
            (3, 10, "NARS-9 region", 11, 13, "ctx_region",
             dict(kind="list", bounds=self.regions, title="NARS-9 ecoregion", note=region_note,
                  error="Pick one of the nine NARS aggregate ecoregion codes, or leave the cell blank.")),
            (4, 1, "NHD FCODE", 2, 5, "ctx_fcode",
             dict(kind="whole", bounds=self.ctxspec["ctx_fcode"][3], title="NHD feature code", note=fcode_note,
                  error="Enter the whole-number NHDPlus V2 feature code of the reach, or leave the cell blank.")),
            (4, 6, "Slope class override", 7, 9, "ctx_slope_override",
             dict(kind="list", bounds=self.slope_classes + ["national"], title="Slope class override",
                  note=slope_note, error="Pick lt_0.5, 0.5_to_2, ge_2 or national, or leave the cell blank.")),
        ]
        for row, lc, label, c1, c2, name, opts in site:
            S.cell(lc, label, row=row, font=site_font, align=CENTER_NOWRAP)
            self.inputs[name] = self._entry(S, row, c1, c2, name, font=entry_font, align=entry_align, **opts)
        S.cell(10, "Slope class used", row=4, font=site_font, align=CENTER_NOWRAP)
        S.cell(11, None, row=4, align=entry_align, font=entry_font)
        S.merge(4, 11, 4, 13)
        self.inputs["ctx_slope_class"] = S.name_cell("calc_slope_class", 11, 4)
        for r in (2, 3, 4):
            ws.row_dimensions[r].height = 21
        box(ws, 2, 1, 4, 13, outer=MED, inner=THIN)
        # band headers
        for c1, c2, text in ((1, 1, "EASI"), (2, 3, "Outcomes"), (4, 5, "Functional Categories"),
                             (6, 8, "Functions"), (9, 13, "Metrics")):
            S.cell(c1, text, row=5, fill=FILL_BAND, font=label_font, align=CENTER)
            S.merge(5, c1, 5, c2)
        heads = ["Index Score", "Name", "Sub-Index Score", "Name", "Functions Rated", "Name", "Rating",
                 "Function Score", "Name", "Value", "Units", "Input Rating", "Source"]
        for c, text in enumerate(heads, 1):
            S.cell(c, text, row=6, font=label_font, align=CENTER)
            S.merge(6, c, 7, c)
        ws.row_dimensions[5].height = 18
        ws.row_dimensions[6].height = 18
        ws.row_dimensions[7].height = 18
        # body rows: the entries and linked cells (the ratings and scores come in the second pass)
        row = 8
        for n, m in enumerate(self.metrics, 1):
            mkey = f"m{n:02d}"
            first = row
            for spec in self.front[mkey]:
                kind, key = spec[0], spec[1]
                if kind == "in":
                    grp, _, label, units, qkind, bounds, source, note = self.qspec[key]
                    label = FRONT_LABELS.get(key, label)
                    S.cell(9, label, row=row, align=LEFT_MID)
                    S.cell(11, units, row=row, align=CENTER_NOWRAP)
                    S.cell(13, FRONT_SOURCES.get(key, source), row=row, align=LEFT_MID)
                    if key not in self.inputs:
                        text = FRONT_NOTES.get(key, note)
                        if qkind == "category":
                            self.inputs[key] = self._entry(
                                S, row, 10, 10, f"in_{key}", kind="list", bounds=["1", "2", "3", "4A", "4B", "4C", "5"],
                                title="ATTAINS category", note=text,
                                error="Pick 1, 2, 3, 4A, 4B, 4C or 5, or leave the cell blank.")
                        else:
                            lo, hi = bounds
                            rng_text = (f"a whole number from {lo} to {hi}" if qkind == "whole"
                                        else f"a number from {fnum(lo)} to {fnum(hi)}")
                            self.inputs[key] = self._entry(
                                S, row, 10, 10, f"in_{key}", kind=qkind, bounds=bounds, title=label[:32], note=text,
                                error=f"Enter {rng_text}, or leave the cell blank when the evidence is unavailable.")
                    else:
                        S.cell(10, link_formula(self.inputs[key]), row=row, fill=FILL_LINK, align=CENTER_NOWRAP)
                    self.front_rows.append((mkey, "in", key, row))
                elif kind == "ov":
                    name, label, okind, bounds, note = self.ovspec[key]
                    # the Override Score carries its own full label; its band, rule and
                    # height come in the second pass, after the grid borders are drawn
                    text = label if name in self.score_overrides else label + " (optional)"
                    S.cell(9, text, row=row, align=LEFT_MID)
                    S.cell(13, self.ov_source[name], row=row, align=LEFT_MID)
                    if okind == "list_rating":
                        self.inputs[name] = self._entry(S, row, 10, 10, name, kind="list", bounds=list(RATINGS),
                                                        title=label[:32], note=note,
                                                        error="Pick Good, Fair or Poor, or leave the cell blank.")
                    elif okind == "decimal":
                        S.cell(11, "%", row=row, align=CENTER_NOWRAP)
                        self.inputs[name] = self._entry(S, row, 10, 10, name, kind="decimal", bounds=bounds,
                                                        title=label[:32], note=note,
                                                        error="Enter a percentage from 0 to 100, or leave the cell blank.")
                    else:
                        self.inputs[name] = self._entry(S, row, 10, 10, name, kind="text", title=label[:32], note=note)
                    self.front_rows.append((mkey, "ov", name, row))
                else:
                    _, field, label, units = spec
                    S.cell(9, label, row=row, align=LEFT_MID)
                    S.cell(11, units, row=row, align=CENTER_NOWRAP)
                    S.cell(13, "Computed from the entries above", row=row, align=LEFT_MID, font=FONT_NOTE)
                    self.front_rows.append((mkey, "calc", field, row))
                row += 1
            self.front_span[mkey] = (first, row - 1)
        self.front_last = row - 1
        for alias, target in ALIASES.items():
            self.inputs[alias] = self.inputs[target]
        # the slope class formula now that the slope cell exists
        slope, ov = self.inputs["slope"], self.inputs["ctx_slope_override"]
        ws.cell(row=4, column=11).value = (
            f'=IF({ov}<>"",{ov},IF(NOT(ISNUMBER({slope})),"national",IF({slope}<0,"national",'
            f'IF({slope}<0.005,"lt_0.5",IF({slope}<0.02,"0.5_to_2","ge_2")))))')
        # monthly flow helper
        r = self.front_last + 2
        S.cell(1, "Monthly flow helper (optional): the twelve EROM mean monthly flows in cfs, QE_01 to QE_12. "
                  "Copy the computed variability into the Low flow row above.", row=r, font=FONT_BOLD, align=LEFT_MID)
        S.merge(r, 1, r, 13)
        S.cell(1, "Month", row=r + 1, font=FONT_BOLD, fill=FILL_BAND, align=CENTER_NOWRAP)
        S.cell(1, "Flow (cfs)", row=r + 2, font=FONT_BOLD, fill=FILL_BAND, align=CENTER_NOWRAP)
        for i, key in enumerate(MONTH_INPUTS, 1):
            col = 1 + i
            S.cell(col, f"QE_{i:02d}", row=r + 1, font=FONT_BOLD, fill=FILL_BAND, align=CENTER_NOWRAP)
            self._entry(S, r + 2, col, col, f"in_{key}", kind="decimal", bounds=(0, 1e9),
                        error="Enter the EROM mean monthly flow in cfs, or leave the cell blank.")
        rng = f"'{SCORE}'!$B${r + 2}:$M${r + 2}"
        S.cell(1, "Flow variability (CV)", row=r + 3, font=FONT_BOLD, align=CENTER)
        ws.row_dimensions[r + 3].height = 30
        # STDEV.P postdates the original file format, so the file must name it _xlfn.STDEV.P or Excel
        # reads #NAME? (Excel shows and edits it as STDEV.P). The IF hid this while the helper was
        # blank: the branch is only evaluated once all twelve flows are entered (found 2026-09-18,
        # when the application began to enter them).
        S.cell(2, f'=IF(COUNT({rng})<12,"",IF(AVERAGE({rng})<=0,"",ROUND(_xlfn.STDEV.P({rng})/AVERAGE({rng}),6)))',
               row=r + 3, fmt="0.000000", align=CENTER_NOWRAP)
        S.name_cell("flow_cv_helper", 2, r + 3)
        S.cell(3, "All twelve flows are required and the mean must be positive. Enter the result in the Monthly "
                  "flow variability cell of Low flow and baseflow dynamics.", row=r + 3, font=FONT_NOTE, align=LEFT_MID)
        S.merge(r + 3, 3, r + 3, 13)
        box(ws, r + 1, 1, r + 3, 13, outer=MED, inner=THIN)
        self.helper_rows = (r, r + 3)
        # notes box
        n0 = r + 5
        c = S.cell(1, "Other Metrics/Notes:", row=n0, fill=FILL_BAND, unlocked=True, align=LEFT_TOP)
        S.merge(n0, 1, n0 + 9, 13)
        for rr in range(n0, n0 + 10):
            for cc in range(1, 14):
                ws.cell(row=rr, column=cc).fill = FILL_BAND
        box(ws, n0, 1, n0 + 9, 13, outer=MED, inner=None)
        self.inputs["site_notes"] = S.name_cell("site_notes", 1, n0)
        self.notes_rows = (n0, n0 + 9)

    # ---- Metrics -----------------------------------------------------------
    def curve_lookup(self, set_id: str, stratum_expr: str) -> str:
        """Row number (within the curve table) for a set and a stratum expression."""
        keys = self.curve_key_range
        return (f"IFERROR(MATCH({q(set_id + '|')}&{stratum_expr},{keys},0),"
                f"MATCH({q(set_id + '|national')},{keys},0))")

    def interp(self, x: str, row_expr: str) -> str:
        """Piecewise-linear interpolation in the engine's operation order, clamped to 0..1."""
        cx = {k: f"INDEX({self.curve_col_range[f'x{k}']},{row_expr})" for k in range(1, 6)}
        cy = {k: f"INDEX({self.curve_col_range[f'y{k}']},{row_expr})" for k in range(1, 6)}

        def seg(a: int, b: int) -> str:
            return f"{cy[a]}+(({x}-{cx[a]})/({cx[b]}-{cx[a]}))*({cy[b]}-{cy[a]})"

        inner = (f"IF({x}<={cx[1]},{cy[1]},IF({x}>={cx[5]},{cy[5]},"
                 f"IF({x}<={cx[2]},{seg(1, 2)},IF({x}<={cx[3]},{seg(2, 3)},"
                 f"IF({x}<={cx[4]},{seg(3, 4)},{seg(4, 5)})))))")
        return f"IF(NOT(ISNUMBER({x})),\"\",MAX(0,MIN(1,{inner})))"

    def build_metrics(self):
        ws = self.wb.create_sheet("Metrics")
        S = Sheet(self.wb, ws, "Metrics")
        S.cell(M_GOOD, NO_INPUT, font=FONT_BOLD)
        S.title("EASI calculator: metric scoring")
        S.row += 1
        A = {r: self.anchor_ref[r] for r in RATINGS}
        n = 0
        for m in self.metrics:
            n += 1
            method = self.by_metric[m["metricId"]]
            self.current_mkey = f"m{n:02d}"
            S.cell(1, f"{n}. {m['functionName']}", font=Font(bold=True, size=11))
            S.cell(2, f"{method['title']} ({method['methodKey']}, {method['operator']})", font=FONT_NOTE)
            S.row += 1
            S.header("Item", "Value", "Good", "Fair", "Poor", "Rating", "Index", "Note")
            for col, rating in ((M_GOOD, "Good"), (M_FAIR, "Fair"), (M_POOR, "Poor")):
                S.cell(col, row=S.row - 1, fill=solid(CLASS_FILL[rating]), align=CENTER_NOWRAP)
            cells = getattr(self, f"_m_{method['methodKey'].replace('-', '_')}", self._m_generic)(S, method, A)
            cells["functionId"] = m["functionId"]
            self.result_cells[m["metricId"]] = cells
            S.blank()
        for col, width in ((M_ITEM, 52), (M_VALUE, 18), (M_GOOD, 15), (M_FAIR, 17), (M_POOR, 15),
                           (M_RATING, 12), (M_INDEX, 12), (M_NOTE, 70)):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = C_TAB_NO_INPUT
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    # -- the Good / Fair / Poor ranges beside each rated value (owner, 2026-10-23) -------
    def write_bins(self, S: Sheet, row: int, bins: dict | None):
        """The ranges a value is rated against, in the three columns after it."""
        for col, rating in ((M_GOOD, "Good"), (M_FAIR, "Fair"), (M_POOR, "Poor")):
            value = (bins or {}).get(rating)
            if value:
                S.cell(col, value, row=row, fill=solid(BIN_FILL[rating]), align=CENTER)

    def rule_bins(self, rule: dict) -> dict:
        """A catalog rule's ranges: its bands' labels, or for a curve the ranges of the curve the
        rule's lookup row (``_rule_formula``, just written) found for this scenario."""
        if rule.get("curve"):
            return self.live_curve_bins(rule["curve"]["set"], self.last_curve_row)
        return band_bins(rule.get("bands") or [])

    def live_curve_bins(self, set_id: str, row_ref: str) -> dict:
        """The ranges of the curve in row ``row_ref`` of the curve table (the scenario's region or
        slope class, else the national curve), rounded for reading; the rating itself is exact."""
        digits, unit = CURVE_DISPLAY.get(set_id, (2, ""))
        x = dict((k, f"ROUND(INDEX({self.curve_col_range[k]},{row_ref}),{digits})") for k in ("x39", "x69"))
        bins = curve_bins(self.curves[set_id]["higherIsBetter"], x["x39"], x["x69"], unit, formula=True)
        return dict((k, f'=IF(NOT(ISNUMBER({row_ref})),"",{v})') for k, v in bins.items())

    # -- shared block pieces --------------------------------------------------
    def _finish(self, S: Sheet, key: str, rating_formula: str, A: dict, route_formula: str,
                combined_formula: str | None = None, note: str = "", status_formula: str | None = None,
                blocked: str | None = None) -> dict:
        """Rows: route, combined, computed rating, override score, rating, index, score, status.

        Every function carries an Override Score entry. It replaces the computed
        rating the way ``assessment.rescore`` does, unless ``blocked`` (the complete
        observed entries of channel evolution and bank condition) holds, which is
        the engine's own order. ``computed`` in the returned refs is the automatic
        result either way, so the Input Rating column never shows an override as an
        input's rating.
        """
        cells = {}
        ov = self.inputs[score_override_name(key)]
        applies = f'OR({ov}="Good",{ov}="Fair",{ov}="Poor")'
        if blocked:
            applies = f"AND({applies},NOT({blocked}))"
        route_formula = f"=IF({applies},{q(SCORE_OVERRIDE_ROUTE)},{route_formula[1:]})"
        S.cell(1, "Route that applied"); S.cell(2, route_formula, fill=FILL_CALC)
        cells["route"] = S.name_cell(f"{key}_route", 2, S.row); S.row += 1
        if combined_formula is not None:
            S.cell(1, "Combined value"); S.cell(2, combined_formula, fill=FILL_CALC)
            cells["combined"] = S.name_cell(f"{key}_combined", 2, S.row); S.row += 1
        S.cell(1, "Computed rating (before the override score)"); S.cell(M_RATING, rating_formula, fill=FILL_CALC)
        cells["computed"] = S.name_cell(f"{key}_computed", M_RATING, S.row); S.row += 1
        S.cell(1, "Override score (optional entry on the Scenario sheet)")
        S.cell(2, link_formula(ov), fill=FILL_CALC)
        S.cell(M_NOTE, "Replaces the computed rating, as a rating changed in the web application does."
                  + (" Complete observed entries take precedence." if blocked else ""),
               font=FONT_NOTE, wrap=True)
        S.row += 1
        S.cell(1, "Function rating")
        S.cell(M_RATING, f"=IF({applies},{ov},{cells['computed']})", fill=FILL_CALC, font=FONT_BOLD)
        cells["rating"] = S.name_cell(f"{key}_rating", M_RATING, S.row)
        S.cell(M_INDEX, "=" + index_from_rating(cells["rating"], A), fill=FILL_CALC, fmt="0.000")
        cells["index"] = S.name_cell(f"{key}_index", M_INDEX, S.row)
        if note:
            S.cell(M_NOTE, note, font=FONT_NOTE, wrap=True)
        S.row += 1
        S.cell(1, "Function score (0 to 15)")
        r = cells["rating"]
        S.cell(2, f"=IF({r}={q('Good')},{self.score_ref['Good']},IF({r}={q('Fair')},{self.score_ref['Fair']},"
                  f"IF({r}={q('Poor')},{self.score_ref['Poor']},\"\")))", fill=FILL_CALC, font=FONT_BOLD)
        cells["score"] = S.name_cell(f"{key}_score", 2, S.row); S.row += 1
        S.cell(1, "Status")
        status_formula = status_formula or f'=IF({r}="","not rated","complete")'
        # the automatic status describes the computed rating, which an override replaces
        status_formula = f"=IF({applies},{q(SCORE_OVERRIDE_ROUTE)},{status_formula[1:]})"
        S.cell(2, status_formula, fill=FILL_CALC)
        cells["status"] = S.name_cell(f"{key}_status", 2, S.row); S.row += 1
        return cells

    def _input_row(self, S: Sheet, key: str, label: str, rule_formula: str | None, A: dict,
                   note: str = "", value_ref: str | None = None, bins: dict | None = None) -> tuple[str, str, str]:
        """An input row: value, its Good / Fair / Poor ranges, rating (if rule), index. Returns
        (value, rating, index) refs.

        The value returned is the entry cell itself, never the mirror in column
        B: in Excel a formula that references an empty cell evaluates to 0, so a
        rule pointed at the mirror would rate a blank input as a number.
        """
        vref = value_ref or self.inputs[key]
        S.cell(1, label); S.cell(2, f'=IF({vref}="","",{vref})', fill=FILL_CALC)
        self.write_bins(S, S.row, bins)
        v = vref
        rat = idx = ""
        if rule_formula is not None:
            S.cell(M_RATING, "=" + rule_formula, fill=FILL_CALC)
            rat = S.ref(M_RATING, S.row)
            S.cell(M_INDEX, "=" + index_from_rating(rat, A), fill=FILL_CALC, fmt="0.000")
            idx = S.ref(M_INDEX, S.row)
            self.input_rating[(self.current_mkey, key)] = rat
        if note:
            S.cell(M_NOTE, note, font=FONT_NOTE, wrap=True)
        S.row += 1
        return v, rat, idx

    def _rule_formula(self, S: Sheet, rule: dict, x: str, region_ref: str | None = None) -> str:
        """Rating formula for one catalog rule (bands or curve) on value expression x.

        A curve rule first writes two helper rows (the curve row in the
        Reference table and the interpolated index) so no formula repeats the
        lookup and every formula stays far below Excel's 8,192-character limit.
        """
        if rule.get("curve"):
            spec = rule["curve"]
            strat = region_ref if spec["stratifier"] == "nars9" else self.inputs["ctx_slope_class"]
            S.cell(1, f"Reference curve used ({spec['set']}, {spec['stratifier']}, national fallback)",
                   font=FONT_NOTE)
            S.cell(2, "=" + self.curve_lookup(spec["set"], strat), fill=FILL_CALC)
            row_ref = S.ref(2, S.row)
            self.last_curve_row = row_ref
            S.cell(M_NOTE, f'=INDEX({self.curve_key_range},{row_ref})', fill=FILL_CALC, font=FONT_NOTE)
            S.row += 1
            S.cell(1, "Curve index (interpolated, 0 to 1; Good at or above 0.69, Fair at or above 0.39)",
                   font=FONT_NOTE)
            S.cell(2, "=" + self.interp(x, row_ref), fill=FILL_CALC, fmt="0.0000")
            idx_ref = S.ref(2, S.row)
            S.row += 1
            return rating_from_index_edges(idx_ref)
        return bands_formula(rule["bands"], x)

    def _worst_best(self, S: Sheet, key: str, method: dict, A: dict, op: str, extra_note: str = "") -> dict:
        """worst_index / best_index over the method's required inputs (curve or bands per input)."""
        required = [i for i in method["inputs"] if i.get("required") and not i.get("contextOnly")]
        idx_refs, first_row = [], None
        for i in required:
            rule = sm.rule_for_input(method, i) if not i.get("regionalBands") else i
            x = self.inputs[i["key"]]
            formula = self._rule_formula(S, rule, x, self.inputs["ctx_region"])
            _, _, idx = self._input_row(S, i["key"], i["label"], formula, A, bins=self.rule_bins(rule))
            idx_refs.append(idx)
            first_row = first_row or S.row - 1
        last_row = S.row - 1
        allow = bool((method.get("formula") or {}).get("allowPartial"))
        idx_rng = S.rng(M_INDEX, first_row, last_row)
        rated = f"COUNT({idx_rng})"
        pick = "MIN" if op == "worst_index" else "MAX"
        gate = f"{rated}=0" if allow else f"{rated}<{len(required)}"
        combined = f"=IF({gate},\"\",{pick}({idx_rng}))"
        S.cell(1, "Governing input"); S.cell(2, f'=IF({S.ref(2, S.row + 1)}="","",INDEX({S.rng(1, first_row, last_row)},MATCH({S.ref(2, S.row + 1)},{idx_rng},0)))', fill=FILL_CALC)
        gov = S.name_cell(f"{key}_governing", 2, S.row); S.row += 1
        S.cell(1, "Combined rating index (the more limiting input)" if op == "worst_index" else "Combined rating index (the better pathway)")
        S.cell(2, combined, fill=FILL_CALC, fmt="0.000")
        comb = S.name_cell(f"{key}_combined", 2, S.row); S.row += 1
        status = (f'=IF({rated}=0,"not rated",IF({rated}<{len(required)},"partial","complete"))' if allow
                  else f'=IF({rated}<{len(required)},"not rated","complete")')
        cells = self._finish(S, key, "=" + rating_from_anchor(comb, A), A, "=" + q("automatic method"),
                             None, extra_note, status)
        cells["governing"] = gov
        cells["combined"] = comb
        return cells

    # -- per-method blocks ----------------------------------------------------
    def _m_generic(self, S, method, A):
        raise SystemExit(f"no block writer for {method['methodKey']}")

    def _m_catchment_land_cover_pressure(self, S, method, A):
        return self._worst_best(S, "m01", method, A, "worst_index",
                                "Partial rating allowed when one input is missing.")

    def _m_watershed_wetland_extent(self, S, method, A):
        return self._sum_capped(S, "m02", method, A)

    def _sum_capped(self, S, key, method, A, curve=False):
        required = [i for i in method["inputs"] if i.get("required") and not i.get("contextOnly")]
        refs = []
        for i in required:
            v, _, _ = self._input_row(S, i["key"], i["label"], None, A)
            refs.append(v)
        cap = fnum(method["formula"]["cap"])
        all_num = "AND(" + ",".join(f"ISNUMBER({r})" for r in refs) + ")"
        combined = f'=IF({all_num},ROUND(MIN({"+".join(refs)},{cap}),12),"")'
        S.cell(1, f"Combined value (sum capped at {method['formula']['cap']})"); S.cell(2, combined, fill=FILL_CALC)
        comb_row = S.row
        comb = S.name_cell(f"{key}_combined", 2, S.row); S.row += 1
        rule = sm.rule_for_method(method)
        rating = "=" + self._rule_formula(S, rule, comb, self.inputs["ctx_region"])
        self.write_bins(S, comb_row, self.rule_bins(rule))
        cells = self._finish(S, key, rating, A, "=" + q("automatic method"),
                             note="All classes are required. A missing class is unknown, not zero.")
        cells["combined"] = comb
        self.input_rating[(key, "combined")] = cells["computed"]
        return cells

    def _m_road_density_inflow_pressure(self, S, method, A):
        return self._threshold(S, "m03", method, A)

    def _threshold(self, S, key, method, A, note=""):
        inp = next(i for i in method["inputs"] if not i.get("contextOnly"))
        value_row = S.row
        v, _, _ = self._input_row(S, inp["key"], inp["label"], None, A)
        for ctx in (i for i in method["inputs"] if i.get("contextOnly")):
            if ctx["key"] in self.inputs or ctx["key"] == "fcodeContext":
                ref = self.inputs.get(ctx["key"]) or self.inputs["ctx_fcode"]
                self._input_row(S, ctx["key"], ctx["label"] + " (context only)", None, A, value_ref=ref)
        rule = sm.rule_for_method(method)
        rating = "=" + self._rule_formula(S, rule, v, self.inputs["ctx_region"])
        self.write_bins(S, value_row, self.rule_bins(rule))
        cells = self._finish(S, key, rating, A, "=" + q("automatic method"), note=note)
        self.input_rating[(key, inp["key"])] = cells["computed"]
        return cells

    def _m_degree_of_regulation(self, S, method, A):
        f = method["formula"]
        num = next(i for i in method["inputs"] if i["key"] == f["numerator"])
        den = next(i for i in method["inputs"] if i["key"] == f["denominator"])
        nv, _, _ = self._input_row(S, num["key"], num["label"], None, A)
        dv, _, _ = self._input_row(S, den["key"], den["label"], None, A)
        positive = [i["key"] for i in method["inputs"] if i.get("positive")]
        guards = [f"ISNUMBER({nv})", f"ISNUMBER({dv})", f"{dv}<>0"]
        for k in positive:
            guards.append(f"{self.inputs[k]}>0")
        combined = (f'=IF(AND({",".join(guards)}),ROUND(({fnum(f.get("numeratorMultiplier", 1))}*{nv})/'
                    f'({fnum(f.get("denominatorMultiplier", 1))}*{dv}),12),"")')
        S.cell(1, "Degree of regulation (%)"); S.cell(2, combined, fill=FILL_CALC)
        self.write_bins(S, S.row, band_bins(method["bands"]))
        comb = S.name_cell("m04_combined", 2, S.row); S.row += 1
        rating = "=" + bands_formula(method["bands"], comb)
        cells = self._finish(S, "m04", rating, A, "=" + q("automatic method"))
        cells["combined"] = comb
        self.input_rating[("m04", "combined")] = cells["computed"]
        return cells

    def _m_erom_flow_variability(self, S, method, A):
        return self._threshold(S, "m05", method, A,
                               "Unvalidated screening proxy: monthly flow variability, lower is better, rated on the "
                               "regional curve with the national fallback.")

    def _m_bank_height_ratio(self, S, method, A):
        return self._threshold(S, "m06", method, A, "Values below 1.0 warrant a check of the cross-section geometry.")

    def _m_entrenchment_ratio(self, S, method, A):
        return self._threshold(S, "m07", method, A, "Rated on the slope-class curve, national fallback when the slope is unknown.")

    def _m_hyporheic_exchange_potential(self, S, method, A):
        return self._worst_best(S, "m08", method, A, "best_index", "One valid pathway gives a partial rating.")

    def _m_channel_adjustment_susceptibility(self, S, method, A):
        key = "m09"
        # inputs of the automatic proxy
        bhr = next(i for i in method["inputs"] if i["key"] == "bhr")
        er = next(i for i in method["inputs"] if i["key"] == "er")
        _, _, bhr_idx = self._input_row(S, "bhr", bhr["label"], bands_formula(bhr["bands"], self.inputs["bhr"]), A,
                                        bins=band_bins(bhr["bands"]))
        first = S.row - 1
        er_rule = sm.rule_for_input(method, er)
        er_formula = self._rule_formula(S, er_rule, self.inputs["er"], None)
        _, _, er_idx = self._input_row(S, "er", er["label"], er_formula, A, bins=self.rule_bins(er_rule))
        idx_rng = S.rng(M_INDEX, first, S.row - 1)
        fcode = self.inputs["ctx_fcode"]
        canal = "OR(" + ",".join(f"{fcode}={c}" for c in self.canal_codes) + ")"
        stage, note_ref = self.inputs["ov_stageClass"], self.inputs["ov_indicators"]
        observed = f'AND(OR({stage}="Good",{stage}="Fair",{stage}="Poor"),LEN(TRIM({note_ref}&""))>0)'
        proxy_idx = f'IF(COUNT({idx_rng})<2,"",MIN({idx_rng}))'
        S.cell(1, "Automatic proxy index (worse of BHR and ER, both required)")
        S.cell(2, "=" + proxy_idx, fill=FILL_CALC, fmt="0.000")
        proxy = S.name_cell(f"{key}_proxy_index", 2, S.row); S.row += 1
        route = (f'=IF({observed},"observed channel class",IF({canal},"canal or ditch (FCODE)",'
                 f'IF({proxy}="","not rated","automatic proxy")))')
        rating = (f'=IF({observed},{stage},IF({canal},"Poor",{rating_from_anchor(proxy, A)}))')
        return self._finish(S, key, rating, A, route, None,
                            "Observed class with a note replaces the proxy; canal and ditch feature codes rate Poor.",
                            blocked=observed)

    def _m_bhr_bank_instability_susceptibility(self, S, method, A):
        key = "m10"
        v, _, _ = self._input_row(S, "bhr", method["inputs"][0]["label"], None, A, bins=band_bins(method["bands"]))
        var = self.variant("bhr-bank-instability-susceptibility", "observed-bank-condition")
        ero = next(i for i in var["inputs"] if i["key"] == "erodingBankPct")
        arm = next(i for i in var["inputs"] if i["key"] == "armoredBankPct")
        e_ref, a_ref = self.inputs["ov_erodingBankPct"], self.inputs["ov_armoredBankPct"]
        _, _, e_idx = self._input_row(S, "ov_erodingBankPct", "Observed eroding bank (%)",
                                      bands_formula(ero["bands"], e_ref), A, value_ref=e_ref,
                                      bins=band_bins(ero["bands"]))
        _, _, a_idx = self._input_row(S, "ov_armoredBankPct", "Observed armored bank (%)",
                                      bands_formula(arm["bands"], a_ref), A, value_ref=a_ref,
                                      bins=band_bins(arm["bands"]))
        observed = f"AND(ISNUMBER({e_ref}),ISNUMBER({a_ref}))"
        proxy_rating = bands_formula(method["bands"], v)
        obs_idx = f"MIN({e_idx},{a_idx})"
        route = f'=IF({observed},"observed bank condition",IF({proxy_rating}="","not rated","automatic proxy"))'
        rating = f"=IF({observed},{rating_from_anchor(obs_idx, A)},{proxy_rating})"
        return self._finish(S, key, rating, A, route, None,
                            "Both observed bank percentages replace the BHR proxy; the worse component governs.",
                            blocked=observed)

    def _m_sediment_supply_potential(self, S, method, A):
        return self._worst_best(S, "m11", method, A, "worst_index", "All three inputs are required.")

    def _m_watershed_agriculture_share(self, S, method, A):
        return self._threshold(S, "m12", method, A,
                               "Pressure proxy sharing the agriculture input of Catchment hydrology; not a bed measurement.")

    def _m_thermal_regulation_vulnerability(self, S, method, A):
        return self._worst_best(S, "m13", method, A, "worst_index", "Both inputs are required.")

    def _m_organic_matter_supply_potential(self, S, method, A):
        return self._sum_capped(S, "m14", method, A)

    def _m_regional_nutrient_condition(self, S, method, A):
        key = "m15"
        region = self.inputs["ctx_region"]
        rrow = f"MATCH({region},{self.nut_regions},0)"
        refs = {}
        for k, gf, fp in (("tn", "tn_gf", "tn_fp"), ("tp", "tp_gf", "tp_fp")):
            x = self.inputs[k]
            gfv = f"INDEX({self.nut_cols[gf]},{rrow})"
            fpv = f"INDEX({self.nut_cols[fp]},{rrow})"
            rating = (f'IF(NOT(ISNUMBER({x})),"",IFERROR(IF({x}<={gfv},"Good",IF({x}>={fpv},"Poor","Fair")),""))')
            label = next(i["label"] for i in method["inputs"] if i["key"] == k)
            # the region's thresholds, live: Good at or below the first, Poor at or above the second
            bins = {"Good": f'=IFERROR("≤"&{gfv}&" mg/L","by NARS-9 region")',
                    "Fair": f'=IFERROR(">"&{gfv}&"-<"&{fpv}&" mg/L","")',
                    "Poor": f'=IFERROR("≥"&{fpv}&" mg/L","")'}
            _, rat, idx = self._input_row(S, k, label + " (mg/L)", rating, A, bins=bins)
            refs[k] = idx
        first = S.row - 2
        idx_rng = S.rng(M_INDEX, first, S.row - 1)
        var = self.variant("regional-nutrient-condition", "streamcat-chem-integrity-nutrient")
        cat_ref, ws_ref = self.inputs["chemCatchment"], self.inputs["chemWatershed"]
        chem_val = f'IF(AND(ISNUMBER({cat_ref}),ISNUMBER({ws_ref})),ROUND(MIN({cat_ref},{ws_ref}),12),"")'
        S.cell(1, "CHEM fallback value (lower of catchment and watershed)"); S.cell(2, "=" + chem_val, fill=FILL_CALC)
        self.write_bins(S, S.row, band_bins(var["bands"]))
        chem = S.name_cell(f"{key}_chem", 2, S.row); S.row += 1
        chem_rating = bands_formula(var["bands"], chem)
        primary = f"COUNT({idx_rng})>0"
        route = f'=IF({primary},"WQP nutrients with regional thresholds",IF({chem_rating}="","not rated","CHEM integrity fallback"))'
        rating = f"=IF({primary},{rating_from_anchor(f'MIN({idx_rng})', A)},{chem_rating})"
        status = f'=IF({primary},IF(COUNT({idx_rng})<2,"partial","complete"),IF({chem_rating}="","not rated","complete"))'
        return self._finish(S, key, rating, A, route, None,
                            "One rated analyte gives a partial rating; the worse analyte governs. Without a region or "
                            "observations the CHEM index decides.", status)

    def _m_attains_regulatory_category(self, S, method, A):
        key = "m16"
        cat = self.inputs["category"]
        keyx = f'SUBSTITUTE(UPPER(TRIM({cat}&""))," ","")'
        hit = f"INDEX({self.att_ratings},MATCH({keyx},{self.att_keys},0))"
        att = f'IF({cat}="","",IFERROR(IF({hit}="none","",{hit}),""))'
        S.cell(1, "ATTAINS category"); S.cell(2, f'=IF({cat}="","",{cat})', fill=FILL_CALC)
        categories: dict = {}
        for code, item in self.methods["attains-regulatory-category"]["formula"]["lookup"].items():
            if item.get("rating"):
                categories.setdefault(item["rating"], []).append(str(code))
        self.write_bins(S, S.row, dict((r, "Category " + ", ".join(c)) for r, c in categories.items()))
        S.cell(M_RATING, "=" + att, fill=FILL_CALC); att_ref = S.ref(M_RATING, S.row)
        S.cell(M_INDEX, "=" + index_from_rating(att_ref, A), fill=FILL_CALC, fmt="0.000"); S.row += 1
        self.input_rating[(key, "category")] = att_ref
        var = self.variant("attains-regulatory-category", "streamcat-chem-integrity-regulatory")
        cat_ref, ws_ref = self.inputs["chemCatchment"], self.inputs["chemWatershed"]
        chem_val = f'IF(AND(ISNUMBER({cat_ref}),ISNUMBER({ws_ref})),ROUND(MIN({cat_ref},{ws_ref}),12),"")'
        S.cell(1, "CHEM fallback value (lower of catchment and watershed)"); S.cell(2, "=" + chem_val, fill=FILL_CALC)
        self.write_bins(S, S.row, band_bins(var["bands"]))
        chem = S.name_cell(f"{key}_chem", 2, S.row); S.row += 1
        chem_rating = bands_formula(var["bands"], chem)
        route = f'=IF({att_ref}<>"","ATTAINS category",IF({chem_rating}="","not rated","CHEM integrity fallback"))'
        rating = f'=IF({att_ref}<>"",{att_ref},{chem_rating})'
        return self._finish(S, key, rating, A, route, None,
                            "Categories 1 and 2 Good, 4A and 4B Fair, 4C and 5 Poor; category 3 and a blank fall back.")

    def _m_habitat_support_potential(self, S, method, A):
        return self._threshold(S, "m17", method, A, "Rated on the regional woody-corridor curve, national fallback.")

    def _m_streamcat_prg_bmmi(self, S, method, A):
        key = "m18"
        p = self.inputs["prGBmmi"]
        valid = f"AND(ISNUMBER({p}),{p}>=0,{p}<=1)"
        S.cell(1, method["inputs"][0]["label"]); S.cell(2, f'=IF({p}="","",{p})', fill=FILL_CALC)
        self.write_bins(S, S.row, band_bins(method["bands"]))
        model_rating = bands_formula(method["bands"], p)
        S.cell(M_RATING, f"=IF({valid},{model_rating},\"\")", fill=FILL_CALC); mr = S.ref(M_RATING, S.row)
        S.cell(M_INDEX, "=" + index_from_rating(mr, A), fill=FILL_CALC, fmt="0.000"); S.row += 1
        self.input_rating[(key, "prGBmmi")] = mr
        var = self.variant("streamcat-prg-bmmi", "streamcat-integrity-products")
        prods = var["formula"]["products"]
        prod_exprs, prod_refs = [], {}
        for name, keys in prods.items():
            refs = [self.inputs[k] for k in keys]
            all_num = "AND(" + ",".join(f"ISNUMBER({r})" for r in refs) + ")"
            expr = f'IF({all_num},PRODUCT({",".join(refs)}),"")'
            S.cell(1, f"{name} (product of six {name[1:3].lower()} components)"); S.cell(2, "=" + expr, fill=FILL_CALC, fmt="0.0000")
            ref = S.name_cell(f"{key}_{name.lower()}", 2, S.row)
            prod_exprs.append(ref); prod_refs[name.lower()] = ref; S.row += 1
        both = "AND(" + ",".join(f"ISNUMBER({r})" for r in prod_exprs) + ")"
        fb_val = f'IF({both},ROUND(MIN({",".join(prod_exprs)}),12),"")'
        S.cell(1, "Fallback value (lower of ICI and IWI)"); S.cell(2, "=" + fb_val, fill=FILL_CALC, fmt="0.0000")
        self.write_bins(S, S.row, band_bins(var["bands"]))
        fb = S.name_cell(f"{key}_products", 2, S.row); S.row += 1
        fb_rating = bands_formula(var["bands"], fb)
        route = f'=IF({mr}<>"","published benthic model",IF({fb_rating}="","not rated","ICI/IWI integrity fallback"))'
        rating = f'=IF({mr}<>"",{mr},{fb_rating})'
        cells = self._finish(S, key, rating, A, route, None,
                             "The model probability governs when present; otherwise all twelve integrity components.")
        cells.update(prod_refs)
        return cells

    def _m_nas_established_taxa_count(self, S, method, A):
        return self._threshold(S, "m19", method, A, "Whole counts only; a fraction or a negative count is unrated.")

    def _m_nearby_dam_proximity(self, S, method, A):
        return self._threshold(S, "m20", method, A, "Whole counts only.")

    # ---- Results (hidden engine sheet) --------------------------------------
    def build_results(self):
        ws = self.wb.create_sheet("Results")
        S = Sheet(self.wb, ws, "Results")
        S.title("EASI calculator: results")
        S.cell(1, "Engine sheet (hidden). Function scores, outcome sub-indices and the Ecosystem Condition Index, "
                  "computed with the STAF rollup: unrated functions leave both the numerator and the denominator. "
                  "The Scenario sheet displays these cells.",
               font=FONT_NOTE)
        S.row += 2
        S.header("Function", "Discipline", "Metric", "Rating", "Score", "Class", "Physical", "Chemical",
                 "Biological", "Score or 0", "Rated", "Route")
        first = S.row
        for m in self.metrics:
            cells = self.result_cells[m["metricId"]]
            row = self.cwa[m["functionId"]]
            S.cell(1, m["functionName"]); S.cell(2, m["discipline"]); S.cell(3, m["name"])
            S.cell(4, f"={cells['rating']}", fill=FILL_CALC)
            S.cell(5, f"={cells['score']}", fill=FILL_CALC)
            sc = S.ref(5, S.row)
            S.cell(6, f'=IF(NOT(ISNUMBER({sc})),"Not assessed",IF({sc}<=5,"Non-Functioning",IF({sc}<=10,"Functioning-at-Risk","Functioning")))', fill=FILL_CALC)
            S.cell(7, config.WEIGHTS.get(row["physical"], 0.0)); S.cell(8, config.WEIGHTS.get(row["chemical"], 0.0))
            S.cell(9, config.WEIGHTS.get(row["biological"], 0.0))
            S.cell(10, f"=IF(ISNUMBER({sc}),{sc},0)", fill=FILL_CALC)
            S.cell(11, f"=IF(ISNUMBER({sc}),1,0)", fill=FILL_CALC)
            S.cell(12, f"={cells['route']}", fill=FILL_CALC)
            S.name_cell(f"result_score_{m['functionId'].replace('-', '_')}", 5, S.row)
            S.row += 1
        last = S.row - 1
        score0, mask = S.rng(10, first, last), S.rng(11, first, last)
        wcols = {"physical": S.rng(7, first, last), "chemical": S.rng(8, first, last), "biological": S.rng(9, first, last)}
        S.blank()
        S.header("Outcome", "Weighted score", "Weighted maximum", "Sub-index", "Displayed", "Class", "Coverage")
        sub_refs, cov_refs = [], []
        for outcome in config.OUTCOMES:
            w = wcols[outcome]
            S.cell(1, outcome.capitalize())
            S.cell(2, f"=SUMPRODUCT({score0},{w},{mask})", fill=FILL_CALC)
            S.cell(3, f"=SUMPRODUCT({w},{mask})*{config.FUNCTION_SCORE_MAX}", fill=FILL_CALC)
            wsum, wmax = S.ref(2, S.row), S.ref(3, S.row)
            S.cell(4, f'=IF({wmax}=0,"",{wsum}/{wmax})', fill=FILL_CALC, fmt="0.000000")
            sub = S.name_cell(f"sub_index_{outcome}", 4, S.row)
            S.cell(5, f'=IF({sub}="","",ROUND({sub},2))', fill=FILL_CALC, fmt="0.00")
            disp = S.name_cell(f"sub_index_{outcome}_display", 5, S.row)
            self.rollup[f"sub_index_{outcome}_display"] = disp
            S.cell(6, f'=IF({disp}="","Not assessed",IF({disp}<=0.39,"Non-Functioning",IF({disp}<=0.69,"Functioning-at-Risk","Functioning")))', fill=FILL_CALC)
            S.cell(7, f"=IF(SUMPRODUCT({w})=0,\"\",SUMPRODUCT({w},{mask})/SUMPRODUCT({w}))", fill=FILL_CALC, fmt="0.00")
            cov_refs.append(S.name_cell(f"coverage_{outcome}", 7, S.row))
            sub_refs.append(sub)
            S.row += 1
        S.cell(1, "Ecosystem Condition Index", font=FONT_BOLD)
        eci = f"IF(COUNT({','.join(sub_refs)})=0,\"\",AVERAGE({','.join(sub_refs)}))"
        S.cell(4, "=" + eci, fill=FILL_CALC, fmt="0.000000", font=FONT_BOLD)
        eci_ref = S.name_cell("eci", 4, S.row)
        S.cell(5, f'=IF({eci_ref}="","",ROUND({eci_ref},2))', fill=FILL_CALC, fmt="0.00", font=FONT_BOLD)
        disp = S.name_cell("eci_display", 5, S.row)
        self.rollup["eci_display"] = disp
        S.cell(6, f'=IF({disp}="","Not assessed",IF({disp}<=0.39,"Non-Functioning",IF({disp}<=0.69,"Functioning-at-Risk","Functioning")))', fill=FILL_CALC, font=FONT_BOLD)
        self.rollup["eci_class"] = S.name_cell("eci_class", 6, S.row)
        S.row += 2
        S.header("Coverage", "Value", "", "", "", "", "")
        S.cell(1, "Functions rated"); S.cell(2, f"=SUM({mask})", fill=FILL_CALC)
        rated = S.name_cell("functions_rated", 2, S.row); S.row += 1
        self.rollup["functions_rated"] = rated
        S.cell(1, "Functions selected"); S.cell(2, len(self.metrics), fill=FILL_CALC)
        selected = S.name_cell("functions_selected", 2, S.row); S.row += 1
        S.cell(1, "Overall coverage"); S.cell(2, f"={rated}/{selected}", fill=FILL_CALC, fmt="0.00")
        overall = S.name_cell("coverage_overall", 2, S.row); S.row += 1
        self.rollup["coverage_overall"] = overall
        covs = [overall] + cov_refs
        limited = "OR(" + ",".join(f"AND(ISNUMBER({c}),{c}<{fnum(COVERAGE_THRESHOLD)})" for c in covs) + ")"
        S.cell(1, "Provisional (coverage below 0.70 overall or for an outcome)")
        S.cell(2, f'=IF({limited},"yes","no")', fill=FILL_CALC)
        self.rollup["provisional"] = S.name_cell("provisional", 2, S.row); S.row += 1
        S.cell(1, "Status")
        S.cell(2, f'=IF({overall}=1,"Complete screening coverage","Partial screening coverage")', fill=FILL_CALC)
        self.rollup["coverage_status"] = S.name_cell("coverage_status", 2, S.row); S.row += 1
        for col, width in ((1, 34), (2, 14), (3, 44), (4, 12), (5, 10), (6, 20), (7, 10), (8, 10), (9, 10),
                           (10, 10), (11, 8), (12, 40)):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A5"
        ws.sheet_state = "hidden"
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    # ---- Scenario: ratings, scores, rollup, summary, legends, charts (second pass)
    def build_score_results(self):
        ws, S = self.score_ws, self.score
        first_body, last_body = 8, self.front_last
        score_font = Font(size=12)          # SFARI's worksheet: 12 pt, the class colours mark them
        # function columns and per-row ratings
        metric_rows = []   # (mkey, metric, first, last)
        for n, m in enumerate(self.metrics, 1):
            mkey = f"m{n:02d}"
            first, last = self.front_span[mkey]
            cells = self.result_cells[m["metricId"]]
            outcome = OUTCOME_OF[m["discipline"]]
            assert self.cwa[m["functionId"]][outcome] == "D", (m["functionId"], outcome)
            fill = solid(OUTCOME_FILL[outcome])
            S.cell(6, m["functionName"], row=first, fill=fill, align=CENTER, font=Font(size=12))
            S.merge(first, 6, last, 6)
            S.cell(7, f"={cells['rating']}", row=first, align=CENTER_NOWRAP, font=score_font)
            S.merge(first, 7, last, 7)
            S.cell(8, f"={cells['score']}", row=first, align=CENTER_NOWRAP, font=score_font, fmt="0")
            S.merge(first, 8, last, 8)
            metric_rows.append((mkey, m, first, last))
            for r in range(first, last + 1):
                ws.cell(row=r, column=9).fill = fill
        for mkey, kind, key, row in self.front_rows:
            if kind == "calc":
                cells = self.result_cells[next(m["metricId"] for n, m in enumerate(self.metrics, 1) if f"m{n:02d}" == mkey)]
                S.cell(10, link_formula(cells[key]), row=row, align=CENTER_NOWRAP,
                       fmt="0.0000" if key in ("ici", "iwi") else "0.00")
            rat = self.input_rating.get((mkey, key)) or (self.input_rating.get((mkey, "combined"))
                                                          if kind == "calc" and key == "combined" else None)
            if rat:
                S.cell(12, f"={rat}", row=row, align=CENTER_NOWRAP)
        # categories and outcomes: merged label cells with the SFARI greys
        cat_rows = []      # (discipline, first, last)
        for mkey, m, first, last in metric_rows:
            if cat_rows and cat_rows[-1][0] == m["discipline"]:
                cat_rows[-1] = (m["discipline"], cat_rows[-1][1], last)
            else:
                cat_rows.append((m["discipline"], first, last))
        out_rows = []      # (outcome, first, last)
        for disc, first, last in cat_rows:
            outcome = OUTCOME_OF[disc]
            if out_rows and out_rows[-1][0] == outcome:
                out_rows[-1] = (outcome, out_rows[-1][1], last)
            else:
                out_rows.append((outcome, first, last))
        for disc, first, last in cat_rows:
            fill = solid(OUTCOME_FILL[OUTCOME_OF[disc]])
            S.cell(4, disc, row=first, fill=fill, align=CENTER, font=Font(size=12))
            S.merge(first, 4, last, 4)
            n_funcs = sum(1 for _, m, *_ in metric_rows if m["discipline"] == disc)
            S.cell(5, f"=COUNT('{SCORE}'!$H${first}:$H${last})", row=first, align=CENTER_NOWRAP,
                   fmt=f'0" of {n_funcs}"', font=Font(size=12))
            S.merge(first, 5, last, 5)
        for outcome, first, last in out_rows:
            fill = solid(OUTCOME_FILL[outcome])
            S.cell(2, OUTCOME_LABEL[outcome], row=first, fill=fill, align=CENTER, font=Font(size=12))
            S.merge(first, 2, last, 2)
            S.cell(3, link_formula(self.rollup[f"sub_index_{outcome}_display"]), row=first, align=CENTER_NOWRAP,
                   fmt="0.00", font=score_font)
            S.merge(first, 3, last, 3)
        S.cell(1, link_formula(self.rollup["eci_display"]), row=first_body, align=CENTER_NOWRAP, fmt="0.00",
               font=Font(size=12))
        S.merge(first_body, 1, last_body, 1)
        # borders: thin grid, medium around the groups of columns and the outcomes
        box(ws, 5, 1, 7, 13, outer=MED, inner=THIN)
        box(ws, first_body, 1, last_body, 13, outer=MED, inner=THIN)
        for col in (2, 4, 6, 9):
            side(ws, 5, col, last_body, col, left=MED)
        for _, first, last in out_rows:
            side(ws, last, 1, last, 13, bottom=MED)
        for _, first, last in cat_rows:
            side(ws, last, 4, last, 13, bottom=MED)
        for mkey, m, first, last in metric_rows:
            side(ws, last, 6, last, 13, bottom=MED)
        # the Override Score row, the last of every function: set apart from the function's
        # metrics as a user-defined input. A double rule above it, a band of its own across
        # the Metrics columns (the entry keeps the entry orange), a bold label and a taller
        # row. After the fills and borders above, which would otherwise paint over it.
        override_rows = {row for _mkey, kind, key, row in self.front_rows
                         if kind == "ov" and key in self.score_overrides}
        assert len(override_rows) == len(self.metrics), "one Override Score row per function"
        for row in override_rows:
            for col in (9, 11, 12, 13):
                ws.cell(row=row, column=col).fill = FILL_OVERRIDE
            label = ws.cell(row=row, column=9)
            label.font = Font(bold=True, italic=True)
            # flush against its entry cell, like a form prompt: the metric names above are
            # left-aligned descriptions, this is a field the user may fill
            label.alignment = Alignment(horizontal="right", vertical="center", indent=1)
            ws.cell(row=row, column=13).font = Font(italic=True)
            side(ws, row - 1, 9, row - 1, 13, bottom=DOUBLE)     # both edges: Excel draws the heavier one
            side(ws, row, 9, row, 13, top=DOUBLE)
        # row heights follow the longest wrapped text in the row
        for r in range(first_body, last_body + 1):
            lines = max(lines_for(str(ws.cell(row=r, column=9).value or ""), 46),
                        lines_for(str(ws.cell(row=r, column=13).value or ""), 40))
            ws.row_dimensions[r].height = max(15.75 * lines, OVERRIDE_ROW_HEIGHT if r in override_rows else 0)
        for mkey, m, first, last in metric_rows:
            need = lines_for(m["functionName"], 24)
            have = sum((ws.row_dimensions[r].height or 15.75) for r in range(first, last + 1))
            if have < 15.75 * need:
                ws.row_dimensions[last].height = (ws.row_dimensions[last].height or 15.75) + 15.75 * need - have
        # conditional formatting: the three classes, as in the legends and charts
        good, fair, poor = cf_fill(C_GOOD), cf_fill(C_FAIR), cf_fill(C_POOR)
        text_rules = lambda col: [  # noqa: E731
            FormulaRule(formula=[f'${col}8="Good"'], fill=good),
            FormulaRule(formula=[f'${col}8="Fair"'], fill=fair),
            FormulaRule(formula=[f'${col}8="Poor"'], fill=poor)]
        score_rules = lambda col, r: [  # noqa: E731
            FormulaRule(formula=[f'AND(ISNUMBER(${col}{r}),${col}{r}>10)'], fill=good),
            FormulaRule(formula=[f'AND(ISNUMBER(${col}{r}),${col}{r}>5,${col}{r}<=10)'], fill=fair),
            FormulaRule(formula=[f'AND(ISNUMBER(${col}{r}),${col}{r}<=5)'], fill=poor)]
        index_rules = lambda col, r: [  # noqa: E731
            FormulaRule(formula=[f'AND(ISNUMBER(${col}{r}),${col}{r}>0.69)'], fill=good),
            FormulaRule(formula=[f'AND(ISNUMBER(${col}{r}),${col}{r}>0.39,${col}{r}<=0.69)'], fill=fair),
            FormulaRule(formula=[f'AND(ISNUMBER(${col}{r}),${col}{r}<=0.39)'], fill=poor)]
        for col in ("G", "L"):
            for rule in text_rules(col):
                ws.conditional_formatting.add(f"{col}{first_body}:{col}{last_body}", rule)
        for rule in score_rules("H", first_body):
            ws.conditional_formatting.add(f"H{first_body}:H{last_body}", rule)
        for col in ("A", "C"):
            for rule in index_rules(col, first_body):
                ws.conditional_formatting.add(f"{col}{first_body}:{col}{last_body}", rule)
        # scoring summary table (columns O to T)
        self._summary_table(metric_rows, cat_rows, score_rules, index_rules)
        self._legends()
        self._charts(metric_rows)
        # widths, panes, print setup, protection
        for col, width in ((1, 17), (2, 13), (3, 13), (4, 17), (5, 13), (6, 31), (7, 10), (8, 12), (9, 46),
                           (10, 25), (11, 10), (12, 12), (13, 40), (14, 3), (15, 17), (16, 32), (17, 8),
                           (18, 10), (19, 10), (20, 11), (21, 3), (22, 8), (23, 20), (24, 3), (25, 8), (26, 20)):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A8"
        ws.print_area = f"A1:Z{self.notes_rows[1]}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    def _summary_table(self, metric_rows, cat_rows, score_rules, index_rules):
        ws, S = self.score_ws, self.score
        c0 = 15   # column O
        head_font = Font(bold=True, size=12)
        S.cell(c0 + 3, "Outcomes", row=5, font=head_font, align=CENTER_NOWRAP)
        S.merge(5, c0 + 3, 5, c0 + 5)
        S.cell(c0, "Scoring Summary", row=5, font=head_font, align=CENTER_NOWRAP)
        S.merge(5, c0, 5, c0 + 2)
        for i, text in enumerate(("Functional Categories", "Functions", "Score", "Physical", "Chemical", "Biological")):
            S.cell(c0 + i, text, row=6, font=head_font, align=CENTER)
            S.merge(6, c0 + i, 7, c0 + i)
        row = 8
        first = row
        for mkey, m, gfirst, glast in metric_rows:
            fill = solid(CATEGORY_FILL[m["discipline"]])
            S.cell(c0 + 1, m["functionName"], row=row, fill=fill, align=LEFT_MID, font=Font(size=10))
            S.cell(c0 + 2, f"=IF('{SCORE}'!$H${gfirst}=\"\",\"\",'{SCORE}'!$H${gfirst})", row=row,
                   align=CENTER_NOWRAP, font=Font(bold=True, size=10), fmt="0")
            weights = self.cwa[m["functionId"]]
            for j, outcome in enumerate(("physical", "chemical", "biological")):
                mark = weights[outcome]
                S.cell(c0 + 3 + j, "" if mark == "-" else mark, row=row, fill=fill, align=CENTER_NOWRAP,
                       font=Font(size=10))
            row += 1
        last = row - 1
        r = first
        for disc, *_ in cat_rows:
            n = sum(1 for _, m, *_ in metric_rows if m["discipline"] == disc)
            S.cell(c0, disc, row=r, fill=solid(CATEGORY_FILL[disc]), align=CENTER, font=Font(size=10))
            S.merge(r, c0, r + n - 1, c0)
            for rr in range(r, r + n):
                ws.cell(row=rr, column=c0).fill = solid(CATEGORY_FILL[disc])
            r += n
        box(ws, 5, c0, 7, c0 + 5, outer=THIN, inner=THIN)
        box(ws, first, c0, last, c0 + 5, outer=THIN, inner=THIN)
        for rule in score_rules(get_column_letter(c0 + 2), first):
            ws.conditional_formatting.add(f"{get_column_letter(c0 + 2)}{first}:{get_column_letter(c0 + 2)}{last}", rule)
        # sub-indices, index, coverage
        r = last + 1
        bold = Font(bold=True, size=11)
        S.cell(c0, "Sub-index Scores", row=r, font=bold, align=RIGHT_MID); S.merge(r, c0, r, c0 + 2)
        for j, outcome in enumerate(("physical", "chemical", "biological")):
            S.cell(c0 + 3 + j, link_formula(self.rollup[f"sub_index_{outcome}_display"]), row=r, font=bold,
                   align=CENTER_NOWRAP, fmt="0.00")
        S.cell(c0, "EASI Index", row=r + 1, font=bold, align=RIGHT_MID); S.merge(r + 1, c0, r + 1, c0 + 2)
        S.cell(c0 + 3, link_formula(self.rollup["eci_display"]), row=r + 1, font=Font(bold=True, size=12),
               align=CENTER_NOWRAP, fmt="0.00")
        S.merge(r + 1, c0 + 3, r + 1, c0 + 5)
        S.cell(c0, "Functions rated", row=r + 2, font=bold, align=RIGHT_MID); S.merge(r + 2, c0, r + 2, c0 + 2)
        S.cell(c0 + 3, f"={self.rollup['functions_rated']}", row=r + 2, align=CENTER_NOWRAP,
               fmt=f'0" of {len(self.metrics)}"')
        S.merge(r + 2, c0 + 3, r + 2, c0 + 5)
        S.cell(c0, "Coverage", row=r + 3, font=bold, align=RIGHT_MID); S.merge(r + 3, c0, r + 3, c0 + 2)
        S.cell(c0 + 3, f"={self.rollup['coverage_status']}", row=r + 3, align=CENTER_NOWRAP)
        S.merge(r + 3, c0 + 3, r + 3, c0 + 5)
        S.cell(c0, "Provisional result", row=r + 4, font=bold, align=RIGHT_MID); S.merge(r + 4, c0, r + 4, c0 + 2)
        S.cell(c0 + 3, f'=IF({self.rollup["provisional"]}="yes","yes (coverage below 0.70)","no")', row=r + 4,
               align=CENTER_NOWRAP)
        S.merge(r + 4, c0 + 3, r + 4, c0 + 5)
        box(ws, r, c0, r + 4, c0 + 5, outer=THIN, inner=THIN)
        col = get_column_letter(c0 + 3)
        for rule in index_rules(col, r):
            ws.conditional_formatting.add(f"{col}{r}:{get_column_letter(c0 + 5)}{r + 1}", rule)
        self.summary_rows = (5, r + 4)

    def _legends(self):
        """The SFARI legends drawn in cells: index scoring, function scoring, rating to score."""
        ws, S = self.score_ws, self.score
        v, w = 22, 23    # columns V, W
        y, z = 25, 26    # columns Y, Z
        head = Font(bold=True, size=11)
        S.cell(v, "EASI Index Scoring", row=5, font=head, align=CENTER_NOWRAP); S.merge(5, v, 5, w)
        S.cell(y, "Function Scoring", row=5, font=head, align=CENTER_NOWRAP); S.merge(5, y, 5, z)
        blocks = ((C_GOOD, "Functioning", "1.00", "0.70", "15", "11"),
                  (C_FAIR, "Functioning At-Risk", None, "0.40", "10", "6"),
                  (C_POOR, "Non-Functioning", None, "0.00", "5", "0"))
        r = 6
        for color, label, top, bottom, ftop, fbottom in blocks:
            for col, t, b in ((w, top, bottom), (z, ftop, fbottom)):
                S.cell(col, label, row=r, fill=solid(color), font=Font(bold=True, size=10), align=CENTER)
                S.merge(r, col, r + 2, col)
                for rr in range(r, r + 3):
                    ws.cell(row=rr, column=col).fill = solid(color)
                edge = col - 1
                if t is not None:
                    S.cell(edge, t, row=r, align=Alignment(horizontal="right", vertical="top"), font=Font(size=9))
                S.cell(edge, b, row=r + 2, align=Alignment(horizontal="right", vertical="bottom"), font=Font(size=9))
            r += 3
        box(ws, 6, w, 14, w, outer=THIN, inner=THIN)
        box(ws, 6, z, 14, z, outer=THIN, inner=THIN)
        # rating to function score
        r = 17
        S.cell(v, "Rating to Function Score", row=r, font=head, align=CENTER_NOWRAP); S.merge(r, v, r, z)
        S.cell(v, "Rating", row=r + 1, font=Font(bold=True, size=10), align=CENTER_NOWRAP); S.merge(r + 1, v, r + 1, w)
        S.cell(y, "Score", row=r + 1, font=Font(bold=True, size=10), align=CENTER_NOWRAP); S.merge(r + 1, y, r + 1, z)
        rr = r + 2
        for rating, color in (("Good", C_GOOD), ("Fair", C_FAIR), ("Poor", C_POOR)):
            S.cell(v, f"{rating} ({CLASS_LABELS[rating]})", row=rr, fill=solid(color), align=CENTER_NOWRAP,
                   font=Font(size=10))
            S.merge(rr, v, rr, w)
            ws.cell(row=rr, column=w).fill = solid(color)
            S.cell(y, f"={self.score_ref[rating]}", row=rr, fill=solid(color), align=CENTER_NOWRAP,
                   font=Font(bold=True, size=10), fmt="0")
            S.merge(rr, y, rr, z)
            ws.cell(row=rr, column=z).fill = solid(color)
            rr += 1
        S.cell(v, "Not rated (blank input)", row=rr, align=CENTER_NOWRAP, font=Font(size=10)); S.merge(rr, v, rr, w)
        S.cell(y, "excluded from the rollup", row=rr, align=CENTER_NOWRAP, font=Font(size=10)); S.merge(rr, y, rr, z)
        box(ws, r + 1, v, rr, z, outer=THIN, inner=THIN)

    def _charts(self, metric_rows):
        """ChartData (hidden) and the two bar charts of the SFARI worksheet."""
        cd = self.wb.create_sheet("ChartData")
        C = Sheet(self.wb, cd, "ChartData")
        C.cell(1, "Chart series (hidden). One column per class so each bar takes the class colour.", font=FONT_NOTE)
        heads = ["Function", "Score", "Non-Functioning (0-5)", "Functioning At-Risk (6-10)", "Functioning (11-15)"]
        for c, text in enumerate(heads, 1):
            C.cell(c, text, row=2, font=FONT_BOLD)
        r = 3
        for mkey, m, gfirst, glast in metric_rows:
            C.cell(1, m["functionName"], row=r)
            C.cell(2, link_formula(f"'{SCORE}'!$H${gfirst}"), row=r)
            s = C.ref(2, r)
            C.cell(3, f"=IF(ISNUMBER({s}),IF({s}<=5,{s},NA()),NA())", row=r)
            C.cell(4, f"=IF(ISNUMBER({s}),IF(AND({s}>5,{s}<=10),{s},NA()),NA())", row=r)
            C.cell(5, f"=IF(ISNUMBER({s}),IF({s}>10,{s},NA()),NA())", row=r)
            r += 1
        f_first, f_last = 3, r - 1
        r += 1
        heads = ["Outcome", "Score", "Non-Functioning (0.00-0.39)", "Functioning At-Risk (0.40-0.69)",
                 "Functioning (0.70-1.00)"]
        for c, text in enumerate(heads, 1):
            C.cell(c, text, row=r, font=FONT_BOLD)
        o_head = r
        r += 1
        for outcome in ("physical", "chemical", "biological"):
            C.cell(1, OUTCOME_LABEL[outcome], row=r)
            C.cell(2, link_formula(self.rollup[f"sub_index_{outcome}_display"]), row=r)
            s = C.ref(2, r)
            C.cell(3, f"=IF(ISNUMBER({s}),IF({s}<=0.39,{s},NA()),NA())", row=r)
            C.cell(4, f"=IF(ISNUMBER({s}),IF(AND({s}>0.39,{s}<=0.69),{s},NA()),NA())", row=r)
            C.cell(5, f"=IF(ISNUMBER({s}),IF({s}>0.69,{s},NA()),NA())", row=r)
            r += 1
        o_first, o_last = o_head + 1, r - 1
        cd.column_dimensions["A"].width = 34
        cd.sheet_state = "hidden"
        cd.protection.sheet = True
        cd.protection.password = PASSWORD

        def bar(title, head_row, first, last, vmax, unit, fmt, height):
            ch = BarChart()
            ch.type = "bar"
            ch.grouping = "clustered"
            ch.overlap = 100
            ch.gapWidth = 40
            ch.title = title
            ch.title.overlay = False
            ch.style = 2
            ch.add_data(Reference(cd, min_col=3, max_col=5, min_row=head_row, max_row=last), titles_from_data=True)
            ch.set_categories(Reference(cd, min_col=1, min_row=first, max_row=last))
            for s, color in zip(ch.series, (C_POOR, C_FAIR, C_GOOD)):
                s.graphicalProperties = GraphicalProperties(solidFill=color)
                s.graphicalProperties.line = LineProperties(solidFill="000000")
            ch.x_axis.scaling.orientation = "maxMin"
            ch.x_axis.delete = False
            ch.y_axis.delete = False
            ch.y_axis.scaling.min = 0
            ch.y_axis.scaling.max = vmax
            ch.y_axis.majorUnit = unit
            ch.y_axis.numFmt = NumFmt(formatCode=fmt, sourceLinked=False)
            ch.y_axis.majorGridlines = None
            ch.dataLabels = DataLabelList()
            ch.dataLabels.showVal = True
            ch.dataLabels.showSerName = False
            ch.dataLabels.showCatName = False
            ch.dataLabels.showLegendKey = False
            ch.legend.position = "b"
            ch.legend.overlay = False
            ch.width, ch.height = 24, height
            return ch

        chart_row = self.summary_rows[1] + 2
        self.score_ws.add_chart(bar("Function Score", 2, f_first, f_last, 15, 1, "0", 13), f"O{chart_row}")
        self.score_ws.add_chart(bar("Outcome Score", o_head, o_first, o_last, 1, 0.1, "0.00", 6.5),
                                f"O{chart_row + 27}")

    # ---- Instructions -------------------------------------------------------
    def build_instructions(self):
        """The Instructions sheet laid out exactly as SFARI's (owner, 2026-10-23): the model block
        and the colour key on a white ground, then the steps and the comments in ruled boxes."""
        ws = self.wb.create_sheet("Instructions", 0)
        S = Sheet(self.wb, ws, "Instructions")
        label = Font(name="Arial", size=10, bold=True)
        body = Font(name="Arial", size=10)
        # SFARI's A1:I13: a white fill and white rules hide the grid behind the model block
        ground = Side(style="thin", color="FFFFFF")
        for r in range(1, 14):
            for c in range(1, 10):
                cell = ws.cell(row=r, column=c)
                cell.fill = solid("FFFFFF")
                cell.border = Border(left=ground, right=ground, top=ground, bottom=ground)
        rows = [
            (1, "Model Name", "Ecosystem Assessment Screening Index (EASI)", label),
            (2, "Developers", DEVELOPERS, None),
            (3, None, "U.S. Army Engineer Research and Development Center, Environmental Laboratory", None),
            (4, None, "Vicksburg, MS", None),
            (5, None, EMAILS, None),
            (7, "Model Version", MODEL_VERSION, body),
            (8, "Date of Last Update", _dt.datetime.strptime(TEMPLATE_DATE, "%Y-%m-%d"), None),
            (10, "Waiver", WAIVER, body),
            (11, None, NO_LIABILITY, body),
        ]
        for r, a, b, font in rows:
            if a:
                S.cell(1, a, row=r, font=label)
            c = S.cell(2, b, row=r, font=font)
            if r == 8:
                c.number_format = "[$-409]mmmm d, yyyy;@"
                c.alignment = Alignment(horizontal="left")
            elif r == 10:
                c.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[10].height = 39.6
        # the colour key, column I
        S.cell(9, "Color Coding Key:", row=1)
        key = ((2, "User Input", FILL_INPUT, THIN), (3, "Calculated", None, THIN),
               (4, "Linked to Another Cell", FILL_LINK, Side(style="thin", color="7F7F7F")),
               (5, "Override Score Row", FILL_OVERRIDE, THIN))
        for r, text, fill, rule in key:
            S.cell(9, text, row=r, fill=fill or solid("FFFFFF"), align=Alignment(horizontal="center"),
                   border=Border(left=rule, right=rule, top=rule, bottom=rule))
        S.cell(9, "Protect/unprotect password:", row=6, font=Font(size=10), align=Alignment(horizontal="center"),
               border=Border(left=ground, right=ground))
        S.cell(9, PASSWORD, row=7, align=Alignment(horizontal="center"), border=BOX)
        # the shared edges stay the key's own rules
        side(ws, 1, 9, 1, 9, bottom=Side())
        side(ws, 8, 9, 8, 9, top=Side())
        for r in (2, 3, 4, 5, 7):
            side(ws, r, 8, r, 8, right=Side())
        # the steps and the comments, ruled like SFARI's
        r = 15
        S.cell(1, "Instructions", row=r, font=label, align=Alignment(horizontal="center"), border=BOX)
        S.cell(2, None, row=r, border=BOX)
        S.merge(r, 1, r, 2)
        for i, text in enumerate(INSTRUCTIONS, 1):
            r += 1
            self._numbered(S, r, i, text)
        self._instruction_legends(S, 15, r)
        r += 2
        S.cell(1, "Comments/Assumptions", row=r, font=Font(bold=True), align=Alignment(horizontal="center"),
               border=BOX)
        S.cell(2, None, row=r, border=BOX)
        S.merge(r, 1, r, 2)
        for i, text in enumerate(COMMENTS, 1):
            r += 1
            self._numbered(S, r, i, text)
        for col, width in (("A", 18.5703125), ("B", 67.5703125), ("C", 8.85546875), ("D", 8.85546875),
                           ("E", 8.85546875), ("F", 8.85546875), ("G", 8.85546875), ("H", 18.85546875),
                           ("I", 26.7109375)):
            ws.column_dimensions[col].width = width
        ws.protection.sheet = True
        ws.protection.password = PASSWORD

    def _instruction_legends(self, S: Sheet, head: int, last: int):
        """The scoring legends beside the steps, where SFARI's Instructions shows its own: the EASI
        Index and Function Score classes as colour bars, and the score each rating gives."""
        ws = S.ws
        rows = list(range(head + 1, last + 1))
        thirds = [rows[: len(rows) // 3], rows[len(rows) // 3: 2 * len(rows) // 3], rows[2 * len(rows) // 3:]]
        classes = ((C_GOOD, "Functioning"), (C_FAIR, "Functioning At-Risk"), (C_POOR, "Non-Functioning"))
        title = Font(size=11)
        tick_top = Alignment(horizontal="right", vertical="top")
        tick_bottom = Alignment(horizontal="right", vertical="bottom")
        for tick_col, c1, c2, text, edges in ((4, 5, 6, "EASI Index Scoring", ("1.0", "0.7", "0.4", "0.0")),
                                              (7, 8, 8, "Function Scoring", ("15", "11", "10", "6", "5", "0"))):
            S.cell(tick_col, text, row=head, font=title, align=Alignment(horizontal="center"))
            S.merge(head, tick_col, head, c2)
            for k, (part, (color, label)) in enumerate(zip(thirds, classes)):
                S.cell(c1, label, row=part[0], fill=solid(color), font=Font(bold=True), align=CENTER)
                for rr in part:
                    for cc in range(c1, c2 + 1):
                        ws.cell(row=rr, column=cc).fill = solid(color)
                S.merge(part[0], c1, part[-1], c2)
                box(ws, part[0], c1, part[-1], c2, outer=THIN, inner=None)
                if len(edges) == 4:        # index: the top, then each class's lower edge
                    if k == 0:
                        S.cell(tick_col, edges[0], row=part[0], align=tick_top)
                    S.cell(tick_col, edges[k + 1], row=part[-1], align=tick_bottom)
                else:                      # function score: each class's top and bottom score
                    S.cell(tick_col, edges[2 * k], row=part[0], align=tick_top)
                    S.cell(tick_col, edges[2 * k + 1], row=part[-1], align=tick_bottom)
        # the score each rating gives (EASI's counterpart of SFARI's Likert to function score check)
        S.cell(9, "Rating to Function Score", row=head, font=title, align=Alignment(horizontal="center"))
        lines = [(f"{self.scores[r]}  {r} ({CLASS_LABELS[r]})", CLASS_FILL[r]) for r in RATINGS]
        lines.append(("Not rated: left out of the sub-indices and the index", None))
        for k, (text, color) in enumerate(lines, 1):
            S.cell(9, text, row=head + k, fill=solid(color) if color else None,
                   align=Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True), border=BOX)

    def _numbered(self, S: Sheet, row: int, n: int, text: str):
        """One numbered line of the steps or the comments: the number centred, the text wrapped."""
        S.cell(1, f"{n})", row=row, align=Alignment(horizontal="center", vertical="center"), border=BOX)
        S.cell(2, text, row=row, align=LEFT_MID, border=BOX)
        S.ws.row_dimensions[row].height = 15.0 * lines_for(text, 62)

    def stamp(self, generator_sha: str):
        """The workbook's identity, out of sight in its document properties: the scoring method
        digest it was generated under, the package digest of the method files it was generated
        from (what the apps match before they serve it) and the generator's sha256."""
        for name, value in ((STAMP_METHOD, self.method_stamp), (STAMP_PACKAGE, method_files_digest()),
                            (STAMP_GENERATOR, generator_sha)):
            self.wb.custom_doc_props.append(StringProperty(name=name, value=value))

    # ---- assemble -----------------------------------------------------------
    def build(self, generator_sha: str) -> bytes:
        self.build_reference()
        self.build_score_layout()
        self.build_metrics()
        self.build_results()
        self.build_score_results()
        self.build_instructions()
        self.stamp(generator_sha)
        order = ["Instructions", SCORE, "Metrics", "Results", "Reference", "ChartData"]
        self.wb._sheets = [self.wb[name] for name in order]
        self.wb.active = 1
        for ws in self.wb.worksheets:
            ws.sheet_view.tabSelected = ws.title == SCORE
        self.wb.properties.creator = "EASI"
        self.wb.properties.lastModifiedBy = "EASI"
        self.wb.properties.title = "EASI calculator"
        self.wb.properties.created = FIXED_STAMP
        self.wb.properties.modified = FIXED_STAMP
        self.wb.calculation.fullCalcOnLoad = True
        buf = io.BytesIO()
        with openpyxl_namespaces():
            self.wb.save(buf)
        return repack(buf.getvalue())


#: Excel 2016+ "show #N/A as an empty cell" for a chart, the option the SFARI
#: charts carry. openpyxl does not write it, and without it every unrated bar
#: prints a "#N/A" data label.
NA_AS_BLANK = ('<extLst><ext uri="{56B9EC1D-385E-4148-901F-78D8002777C0}" '
               'xmlns:c16r3="http://schemas.microsoft.com/office/drawing/2017/03/chart">'
               '<c16r3:dataDisplayOptions16><c16r3:dispNaAsBlank val="1"/></c16r3:dataDisplayOptions16>'
               '</ext></extLst>')


def repack(source: bytes) -> bytes:
    """Rewrite the package with fixed entry timestamps so the bytes are reproducible.

    openpyxl stamps ``dcterms:modified`` with the save time regardless of the
    workbook properties, so that element is pinned here too (and the build stops
    if either stamp is not where it is pinned: the bytes would carry the clock).
    The chart parts gain the "#N/A as blank" option on the way through.
    """
    stamp = FIXED_STAMP.strftime("%Y-%m-%dT%H:%M:%SZ")
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                for name in ("modified", "created"):
                    text, n = re.subn(rf"(<dcterms:{name}[^>]*>)[^<]*(</dcterms:{name}>)", rf"\g<1>{stamp}\g<2>", text)
                    if n != 1:
                        raise ValueError(f"docProps/core.xml has {n} dcterms:{name} elements, not one")
                data = text.encode("utf-8")
            elif info.filename.startswith("xl/charts/chart") and info.filename.endswith(".xml"):
                text = data.decode("utf-8")
                assert text.count("</chart>") == 1 and "dispNaAsBlank" not in text, info.filename
                data = text.replace("</chart>", NA_AS_BLANK + "</chart>").encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0
            zout.writestr(zi, data)
    return out.getvalue()


def generate(*, method_stamp: str | None = None, generator_sha: str | None = None) -> bytes:
    """The workbook bytes. ``method_stamp`` and ``generator_sha`` default to the running
    evaluator's method version and this file's sha256 (a rebuild); a check passes the ones
    a committed workbook recorded, so the same arithmetic gives the same bytes."""
    if generator_sha is None:
        with open(os.path.abspath(__file__), "rb") as f:
            generator_sha = hashlib.sha256(f.read()).hexdigest()
    return Builder(method_stamp=method_stamp).build(generator_sha)


def method_files_digest() -> str:
    """``method_package.package_digest`` of the active method files: the digest a method package
    of exactly these files carries, and the adopted method's pin names."""
    from easi import method_package as mp
    data_dir = Path(config.DATA_DIR)
    return mp.package_digest({n: hashlib.sha256((data_dir / n).read_bytes()).hexdigest() for n in mp.METHOD_FILES})


def committed_stamps(path: str) -> dict:
    """The identity a committed workbook records in its document properties: the scoring
    method digest it was generated under and the sha256 of the generator that built it."""
    with open(path, "rb") as f:
        stamps = served.workbook_stamps(f.read())
    return {"method": stamps.get(STAMP_METHOD), "generator": stamps.get(STAMP_GENERATOR)}


def accepted_method_stamps() -> list[str]:
    """The method versions a committed workbook or the parity fixture may name and still be
    the active method's: the running evaluator's, and every version the assessment library
    records for exactly the active method files under an evaluator it validated them with
    (``method_authority.validated_under``, the version whose package digest is the active
    files'). An evaluator move that scores identically never changes a cell of the
    workbook's arithmetic, so the identity cells are compared with the method files'
    identity instead of forcing a rebuild."""
    from easi import method_authority as ma
    from easi import method_package as mp
    stamps = [method_version()]
    try:
        entry = ma.library_entry()
        if entry is not None:
            active = method_files_digest()
            version = ma.default_version(entry)
            envelope = json.loads((ma.version_dir(entry, version) / mp.ENVELOPE)
                                  .read_text(encoding="utf-8"))
            if (envelope.get("identity") or {}).get("packageDigest") == active:
                stamps += [str(r["methodVersion"]) for r in ma.validated_under(entry, version)
                           if r.get("methodVersion")]
    except (OSError, ValueError, KeyError, TypeError):
        pass
    return list(dict.fromkeys(stamps))


def check(path: str) -> tuple[bool, str]:
    """Whether the committed workbook is what this generator produces from the active
    method files. Its recorded identity cells are inputs: the workbook is regenerated with
    the method digest and the generator sha256 it recorded and compared byte for byte, and
    the recorded method digest must be one of ``accepted_method_stamps()`` (the method
    files' identity), so an evaluator the library validated never forces a rebuild."""
    with open(path, "rb") as f:
        committed = f.read()
    stamps = committed_stamps(path)
    accepted = accepted_method_stamps()
    if stamps["method"] not in accepted:
        return False, (f"DIFFERS: the committed workbook names method {stamps['method']}, not one "
                       f"the library records for the active method files ({', '.join(accepted)})")
    if generate(method_stamp=stamps["method"], generator_sha=stamps["generator"]) != committed:
        return False, "DIFFERS: the generator's output changed beyond the identity cells; rebuild"
    note = ("" if stamps["method"] == method_version() else
            f" (the workbook records method {stamps['method']}; this evaluator computes "
            f"{method_version()} for the same method files)")
    return True, "unchanged" + note


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(OUT_DIR, OUT_NAME))
    ap.add_argument("--check", action="store_true",
                    help="fail if the committed file differs from what the generator produces "
                         "with the identity the file records")
    args = ap.parse_args()
    if args.check:
        ok, why = check(args.out)
        print(why)
        raise SystemExit(0 if ok else 1)
    data = generate()
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with open(args.out, "wb") as f:
        f.write(data)
    print(f"wrote {args.out} ({len(data):,} bytes, sha256 {hashlib.sha256(data).hexdigest()[:12]}, "
          f"method {method_version()})")


if __name__ == "__main__":
    main()
